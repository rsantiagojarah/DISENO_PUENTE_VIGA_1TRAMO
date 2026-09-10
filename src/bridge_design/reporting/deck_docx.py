"""Detailed A4 Word calculation report for ``diseno-tablero``."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from bridge_design.domain.diaphragm import (
    combine_diaphragm_moments,
    combine_diaphragm_shears,
)
from bridge_design.domain.exterior_girder import (
    combine_exterior_girder_moments,
    combine_exterior_girder_shears,
)
from bridge_design.domain.interior_girder import (
    combine_interior_girder_moments,
    combine_interior_girder_shears,
)
from bridge_design.domain.load_combinations import combine_transverse_slab_moments
from bridge_design.reporting.deck_docx_charts import save_strength_chart
from bridge_design.reporting import deck_docx_detail as detail
from bridge_design.reporting.models import DeckReportData, selected_option

NAVY = "173746"
TEAL = "155E75"
PALE = "EAF2F4"
LIGHT = "F5F7F8"
GRAY = "61737D"
RULE = "B7C6CC"
WHITE = "FFFFFF"
GREEN = "2E6B52"
RED = "8D3D32"
FONT = "Arial Narrow"
UNIFORM_FONT_SIZE_PT = 11
COVER_TITLE_FONT_SIZE_PT = 22
TEXT_COLOR = "000000"
NORMAL_MARGIN_MM = 25.4
MAX_CONTENT_WIDTH_MM = 158.0

REF_COMB = (
    "Manual de Puentes MTC 2018, Art. 2.4.5.3.1 y Tabla 2.4.5.3.1-1."
)
REF_STRIP = (
    "Manual de Puentes MTC 2018, Art. 2.6.4.2.1.3 y Tabla 2.6.4.2.1.3-1."
)
REF_FLEX = (
    "Manual de Puentes MTC 2018, Art. 2.9.1.4.4.2."
)
REF_CRACK = (
    "Manual de Puentes MTC 2018, Art. 2.9.1.4.4.3."
)
REF_TEMP = (
    "Manual de Puentes MTC 2018, Art. 2.9.1.4.5.8."
)
REF_DISTRIBUTION = (
    "Manual de Puentes MTC 2018, Art. 2.9.1.4.6.3.2."
)
REF_SHEAR = (
    "Manual de Puentes MTC 2018, Art. 2.9.1.5.6.3."
)
REF_DEVELOPMENT = (
    "Manual de Puentes MTC 2018, Art. 2.6.5.6.2.1."
)
REF_FATIGUE = (
    "Manual de Puentes MTC 2018, Art. 2.7.1.1.3."
)
REF_HL93 = (
    "Manual de Puentes MTC 2018, Art. 2.4.3.2.2; presencia múltiple e "
    "incremento dinámico: Arts. 2.4.3.2.2.6 y 2.4.3.3."
)
REF_DEAD_LOAD = (
    "Manual de Puentes MTC 2018, Art. 2.4.2.1 y Tabla 2.4.2.1-1."
)
REF_PEDESTRIAN = (
    "Manual de Puentes MTC 2018, Art. 2.4.3.6.1."
)
REF_DISTRIBUTION_FACTORS = (
    "Manual de Puentes MTC 2018, Arts. 2.6.4.2.2.2b, 2.6.4.2.2.2d, "
    "2.6.4.2.2.3a y 2.6.4.2.2.3b."
)
REF_BARRIER = (
    "Rodríguez Serquén, Puentes con AASHTO LRFD 2020, desarrollo A13.3.1; "
    "Manual de Puentes MTC 2018, criterios de barreras y evento extremo aplicables."
)


def generate_deck_docx(data: DeckReportData, output_path: str | Path) -> Path:
    """Create the detailed Word calculation report and return its absolute path."""
    path = Path(output_path).expanduser().resolve()
    if path.suffix.lower() != ".docx":
        path = path.with_suffix(".docx")
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_document(document)
    with TemporaryDirectory(prefix="deck_report_") as chart_dir:
        _cover(document, data)
        _contents(document)
        _general_basis(document, data)
        _transverse_slab(document, data, Path(chart_dir))
        _girder(document, data, exterior=False, chart_dir=Path(chart_dir))
        _girder(document, data, exterior=True, chart_dir=Path(chart_dir))
        _barrier(document, data)
        _cantilever(document, data)
        _diaphragm(document, data, Path(chart_dir))
        _reactions(document, data)
        _conclusions(document, data)
        _references(document)
        _enforce_uniform_typography(document)
        document.save(path)
    return path


def select_deck_docx_save_path() -> Path | None:
    """Open a native save dialog for the final Word memory."""
    try:
        import tkinter as tk
        from tkinter import filedialog
    except ImportError as exc:
        raise RuntimeError("Tkinter no está disponible para elegir el destino Word.") from exc
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    default_name = f"MEMORIA_CALCULO_TABLERO_{datetime.now():%Y%m%d_%H%M}.docx"
    try:
        selected = filedialog.asksaveasfilename(
            parent=root,
            title="Guardar memoria de cálculo detallada del tablero",
            defaultextension=".docx",
            initialfile=default_name,
            filetypes=(("Documento de Word", "*.docx"),),
        )
    finally:
        root.destroy()
    return Path(selected) if selected else None


def generate_deck_docx_with_dialog(data: DeckReportData) -> Path | None:
    """Ask for a destination and create the Word memory when accepted."""
    path = select_deck_docx_save_path()
    if path is None:
        print("Generación de la memoria Word cancelada por el usuario.")
        return None
    generated = generate_deck_docx(data, path)
    print(f"Memoria Word guardada en: {generated}")
    return generated


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(NORMAL_MARGIN_MM)
    section.bottom_margin = Mm(NORMAL_MARGIN_MM)
    section.left_margin = Mm(NORMAL_MARGIN_MM)
    section.right_margin = Mm(NORMAL_MARGIN_MM)
    section.header_distance = Mm(12.7)
    section.footer_distance = Mm(12.7)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    _font_style(normal, UNIFORM_FONT_SIZE_PT, color=TEXT_COLOR)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color, before, after in (
        ("Title", 11, TEXT_COLOR, 0, 10),
        ("Subtitle", 11, TEXT_COLOR, 0, 8),
        ("Heading 1", 11, TEXT_COLOR, 14, 7),
        ("Heading 2", 11, TEXT_COLOR, 10, 4),
        ("Heading 3", 11, TEXT_COLOR, 7, 3),
    ):
        style = document.styles[name]
        _font_style(style, size, bold=name != "Subtitle", color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        style.paragraph_format.keep_with_next = True

    for name, size, italic, color in (
        ("Caption", 11, True, TEXT_COLOR),
        ("Intense Quote", 11, False, TEXT_COLOR),
    ):
        style = document.styles[name]
        _font_style(style, size, italic=italic, color=color)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if "Formula" not in document.styles:
        style = document.styles.add_style("Formula", WD_STYLE_TYPE.PARAGRAPH)
        _font_style(style, 11, color=TEXT_COLOR)
        style.paragraph_format.left_indent = Mm(4)
        style.paragraph_format.right_indent = Mm(3)
        style.paragraph_format.space_before = Pt(1)
        style.paragraph_format.space_after = Pt(1)
    if "Equation" not in document.styles:
        style = document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
        _font_style(style, 11, color=TEXT_COLOR)
        style.paragraph_format.left_indent = Mm(7)
        style.paragraph_format.right_indent = Mm(5)
        style.paragraph_format.space_before = Pt(2)
        style.paragraph_format.space_after = Pt(4)
    bullet = document.styles["List Bullet"]
    _font_style(bullet, 11, color=TEXT_COLOR)
    bullet.paragraph_format.left_indent = Mm(11)
    bullet.paragraph_format.first_line_indent = Mm(-5)
    bullet.paragraph_format.space_after = Pt(2)
    bullet.paragraph_format.line_spacing = 1.5
    bullet.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _header_footer(section)


def _font_style(style, size: float, *, bold: bool = False, italic: bool = False, color: str) -> None:
    style.font.name = FONT
    style.font.size = Pt(UNIFORM_FONT_SIZE_PT)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _header_footer(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("MEMORIA DE CÁLCULO · TABLERO DE PUENTE")
    _format_run(run, 7.8, bold=True, color=TEAL)
    _bottom_border(paragraph, RULE, 5)
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("DISEÑO ESTRUCTURAL  ·  ")
    _format_run(run, 7.5, color=GRAY)
    _field(paragraph, "PAGE")


def _cover(document: Document, data: DeckReportData) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(34)
    run = p.add_run("INGENIERÍA ESTRUCTURAL")
    _format_run(run, 9, bold=True, color=TEAL)
    _bottom_border(p, TEAL, 16)
    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(42)
    p.add_run("Memoria de cálculo\ndel tablero de puente")
    p = document.add_paragraph(style="Subtitle")
    p.add_run(
        "Losa transversal · vigas principales · barrera · voladizo · "
        "diafragmas · reacciones"
    )
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(32)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("DOCUMENTO TÉCNICO DE DISEÑO")
    _format_run(run, 9, bold=True, color=TEAL)
    geom = data.project_inputs.transverse_slab.geometry
    rows = (
        ("Sistema", "Puente de concreto tipo viga–losa, un tramo"),
        ("Luz longitudinal", f"{data.project_inputs.interior_girder.span_length_m:.2f} m"),
        ("Ancho total del tablero", f"{geom.total_width_m:.2f} m"),
        ("Número de vigas", str(geom.girder_count)),
        ("Norma principal", "Manual de Puentes MTC 2018"),
        ("Fecha de emisión", datetime.now().strftime("%d/%m/%Y")),
    )
    _table(document, ("Dato", "Descripción"), rows, widths=(42, 125), accent=True)
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(
        "La presente memoria desarrolla las expresiones, leyendas, sustituciones "
        "numéricas, resultados y criterios técnicos correspondientes a la opción "
        "de refuerzo adoptada."
    )
    _format_run(run, 9.5, italic=True, color=GRAY)
    document.add_page_break()


def _contents(document: Document) -> None:
    document.add_heading("Contenido", level=1)
    _body(
        document,
        "La memoria se organiza por componente estructural. Los títulos conservan estilos "
        "jerárquicos de Word para permitir navegación directa desde el panel de títulos."
    )
    _table(
        document,
        ("Sección", "Contenido desarrollado"),
        (
            ("1", "Bases de diseño, materiales, acciones y combinaciones"),
            ("2", "Losa transversal: análisis, flexión, fisuración y envolventes"),
            ("3", "Viga principal interior: resistencia, servicio, fatiga y detalle"),
            ("4", "Viga principal exterior: resistencia, servicio, fatiga y detalle"),
            ("5", "Barrera de concreto: línea de fluencia, interfaz y anclaje"),
            ("6", "Losa en voladizo: Resistencia I, colisión, acero y desarrollo"),
            ("7", "Viga diafragma: envolventes, flexión y cortante"),
            ("8", "Reacciones para apoyos y estribos"),
            ("9", "Resumen del diseño adoptado"),
            ("10", "Referencias normativas"),
        ),
        widths=(25, 142),
        accent=True,
    )
    document.add_page_break()


def _general_basis(document: Document, data: DeckReportData) -> None:
    inputs = data.project_inputs
    concrete = inputs.materials.concrete
    steel = inputs.materials.steel
    vehicle = inputs.live_loads.vehicular
    layout = inputs.transverse_slab.load_layout
    document.add_heading("1. Bases de diseño", level=1)
    _body(
        document,
        "El cálculo se desarrolla con estados límite LRFD. Las acciones permanentes "
        "se evalúan en su ubicación física; las cargas móviles se desplazan por el "
        "dominio permitido y se conservan, estación por estación, los valores máximo "
        "y mínimo que forman las envolventes de diseño."
    )
    document.add_heading("1.1 Unidades y materiales", level=2)
    _table(
        document,
        ("Parámetro", "Símbolo", "Valor adoptado"),
        (
            ("Peso específico del concreto", "γc", f"{concrete.specific_weight_tn_m3:.3f} Tn/m³"),
            ("Resistencia del concreto", "f'c", f"{concrete.compressive_strength_kg_cm2:.1f} kg/cm²"),
            ("Módulo del concreto", "Ec", f"{concrete.elastic_modulus_kg_cm2:,.0f} kg/cm²"),
            ("Especificación del acero", "—", steel.specification),
            ("Fluencia del acero", "fy", f"{steel.yield_strength_kg_cm2:,.0f} kg/cm²"),
            ("Módulo del acero", "Es", f"{steel.elastic_modulus_kg_cm2:,.0f} kg/cm²"),
        ),
        widths=(83, 24, 60),
    )
    wc_kcf = concrete.specific_weight_tn_m3 * 0.06242796
    fc_ksi = concrete.compressive_strength_kg_cm2 * 0.01422334
    ec_ksi = concrete.elastic_modulus_kg_cm2 * 0.01422334
    _calc(
        document,
        "Módulo de elasticidad del concreto",
        "Ec = 120 000·K₁·wc²·(f'c)^0.33",
        "Ec: módulo elástico (ksi); K₁: factor del agregado; wc: densidad (kcf); f'c: resistencia (ksi).",
        f"Ec = 120 000·{concrete.aggregate_correction_factor:.3f}·{wc_kcf:.5f}²·{fc_ksi:.4f}^0.33 = {ec_ksi:,.1f} ksi",
        f"Ec = {concrete.elastic_modulus_kg_cm2:,.0f} kg/cm².",
        "El valor se emplea en las rigideces de los modelos matriciales y en las propiedades de sección fisurada.",
        "Manual de Puentes MTC 2018, Art. 2.5.4.4.",
    )
    document.add_heading("1.2 Acciones y combinaciones", level=2)
    _table(
        document,
        ("Acción", "Dato utilizado", "Uso"),
        (
            ("DC", "Peso propio de losa, vigas, diafragmas, barrera y barandas", "Permanente estructural"),
            ("DW", f"{inputs.materials.asphalt.thickness_m:.3f} m de {inputs.materials.asphalt.name}", "Superficie de rodadura"),
            ("PL", f"{inputs.live_loads.pedestrian.load_tn_m2:.4f} Tn/m²", "Veredas"),
            ("LL+IM", f"{vehicle.name}; IM = {data.transverse_result.dynamic_load_allowance:.0%}", "Envolvente móvil"),
        ),
        widths=(25, 82, 60),
    )
    _calc(
        document,
        "Combinación de Resistencia I",
        "U = 1.25·DC + 1.50·DW + 1.75·PL + 1.75·(LL+IM)",
        "U: efecto factorizado; DC y DW: permanentes; PL: peatonal; LL+IM: vehicular con incremento dinámico.",
        "En cada estación se sustituyen los efectos con su signo. Los factores mínimos u omisión de acciones favorables se aplican al formar cada extremo de la envolvente.",
        "La combinación se evalúa para momento y cortante y se conserva el extremo que gobierna cada diseño.",
        "Este criterio evita sumar máximos que no ocurren en una misma estación y conserva el signo físico de cada acción.",
        REF_COMB,
    )
    _calc(
        document,
        "Modelo de carga móvil",
        "E⁺(x) = max[Eᵢ(x)]  ;  E⁻(x) = min[Eᵢ(x)]",
        "Eᵢ(x): efecto en la estación x para la posición i del vehículo; E⁺ y E⁻: envolventes superior e inferior.",
        f"Recorrido transversal: {layout.vehicle_move_start_m:.2f} a {layout.vehicle_move_end_m:.2f} m, paso {layout.vehicle_step_m:.3f} m. En vigas se desplazan camión y tándem a lo largo de la luz.",
        "Se obtienen envolventes físicas de máximos positivos y máximos negativos; el orden de recorrido no altera el resultado.",
        "Los diagramas muestran únicamente las envolventes que intervienen en el diseño, no todas las posiciones individuales.",
        REF_HL93,
    )
    document.add_heading("1.3 Criterio de presentación", level=2)
    _body(
        document,
        "Se desarrollan numéricamente las estaciones críticas y la alternativa de acero "
        "adoptada. El análisis matricial y la generación de líneas de influencia se "
        "describen como método; no se imprimen matrices completas ni barridos redundantes."
    )


def _transverse_slab(document: Document, data: DeckReportData, chart_dir: Path) -> None:
    inputs = data.project_inputs
    geom = inputs.transverse_slab.geometry
    layout = inputs.transverse_slab.load_layout
    reinf = data.slab_reinforcement
    concrete = inputs.materials.concrete
    steel = inputs.materials.steel
    vehicle = inputs.live_loads.vehicular
    rows = combine_transverse_slab_moments(data.transverse_result)
    strength = [row for row in rows if row.limit_state == "Resistencia"]
    positive = max(strength, key=lambda row: row.combined_moment_tn_m)
    negative = min(strength, key=lambda row: row.combined_moment_tn_m)
    document.add_heading("2. Losa transversal", level=1)
    _body(
        document,
        "La franja longitudinal de un metro se modela como una viga continua transversal "
        "apoyada en los ejes de las vigas principales, con voladizos extremos. La matriz "
        "global se ensambla con elementos de flexión; las cargas distribuidas y puntuales "
        "se convierten en acciones nodales consistentes."
    )
    document.add_heading("2.1 Geometría y modelo estructural", level=2)
    _table(
        document,
        ("Parámetro", "Valor"),
        (
            ("Separación entre vigas S", f"{geom.girder_spacing_m:.3f} m"),
            ("Voladizo por borde", f"{geom.overhang_m:.3f} m"),
            ("Espesor de losa", f"{geom.slab_thickness_m:.3f} m"),
            ("Ancho total", f"{geom.total_width_m:.3f} m"),
            ("Franja de cálculo", f"{geom.strip_length_m:.3f} m"),
        ),
        widths=(105, 62),
    )
    inertia = geom.slab_inertia_m4
    _calc(
        document,
        "Propiedad geométrica de la franja de losa",
        "I = b·h³/12",
        "I: momento de inercia de la franja; b: ancho longitudinal de cálculo; h: espesor de la losa",
        f"I = {geom.strip_length_m:.3f}·({geom.slab_thickness_m:.3f})³/12 = {inertia:.7f} m⁴",
        f"la franja de un metro se analiza con I = {inertia:.7f} m⁴ y rigidez flexional Ec·I.",
        "Esta propiedad se utiliza en cada elemento de flexión del modelo transversal; la continuidad se impone en los apoyos ubicados sobre los ejes de las vigas.",
        "Manual de Puentes MTC 2018, Art. 2.6.3, criterios de análisis estructural.",
    )
    document.add_heading("2.2 Determinación detallada de cargas", level=2)
    materials = inputs.materials
    live_loads = inputs.live_loads
    slab_q = concrete.specific_weight_tn_m3 * geom.slab_thickness_m * geom.strip_length_m
    sidewalk_q = (
        materials.sidewalk.specific_weight_tn_m3
        * materials.sidewalk.thickness_m
        * geom.strip_length_m
    )
    asphalt_q = (
        materials.asphalt.specific_weight_tn_m3
        * materials.asphalt.thickness_m
        * geom.strip_length_m
    )
    pedestrian_q = live_loads.pedestrian.load_tn_m2 * geom.strip_length_m
    _calc(
        document,
        "Peso propio de la losa por metro transversal",
        "qDC,losa = γc·h·b",
        "qDC,losa: carga lineal de la losa; γc: peso específico del concreto; h: espesor de losa; b: ancho longitudinal de la franja",
        f"qDC,losa = {concrete.specific_weight_tn_m3:.3f}·{geom.slab_thickness_m:.3f}·{geom.strip_length_m:.3f} = {slab_q:.3f} Tn/m",
        f"la carga uniforme de peso propio aplicada en todo el ancho del tablero es qDC,losa = {slab_q:.3f} Tn/m.",
        "La carga se mantiene en su posición física y participa como DC en todas las combinaciones aplicables.",
        REF_DEAD_LOAD,
    )
    _calc(
        document,
        "Peso propio de las veredas",
        "qDC,ver = γver·tver·b",
        "qDC,ver: carga lineal de vereda; γver: peso específico del material; tver: espesor; b: ancho longitudinal de la franja",
        f"qDC,ver = {materials.sidewalk.specific_weight_tn_m3:.3f}·{materials.sidewalk.thickness_m:.3f}·{geom.strip_length_m:.3f} = {sidewalk_q:.3f} Tn/m",
        f"en cada franja lateral de {layout.sidewalk_width_m:.3f} m se aplica qDC,ver = {sidewalk_q:.3f} Tn/m.",
        "La ubicación lateral de las veredas se conserva; no se sustituye por una carga uniforme sobre todo el tablero.",
        REF_DEAD_LOAD,
    )
    _calc(
        document,
        "Cargas lineales de baranda y barrera",
        "P = w·b/1000",
        "P: carga puntual sobre la franja transversal; w: peso lineal del accesorio en kg/m; b: ancho longitudinal analizado",
        f"Pbaranda = {materials.railing.weight_kg_m:.3f}·{geom.strip_length_m:.3f}/1000 = {materials.railing.weight_kg_m*geom.strip_length_m/1000:.3f} Tn\n"
        f"Pbarrera = {materials.barrier.weight_kg_m:.3f}·{geom.strip_length_m:.3f}/1000 = {materials.barrier.weight_kg_m*geom.strip_length_m/1000:.3f} Tn",
        "las cargas se aplican puntualmente en los ejes físicos de las barandas y barreras de ambos bordes.",
        "La posición real de cada accesorio determina su contribución positiva o negativa en cada estación de la losa.",
        REF_DEAD_LOAD,
    )
    _calc(
        document,
        "Carga de la superficie de rodadura",
        "qDW = γasf·tasf·b",
        "qDW: carga lineal de superficie de rodadura; γasf: peso específico del asfalto; tasf: espesor; b: ancho longitudinal de la franja",
        f"qDW = {materials.asphalt.specific_weight_tn_m3:.3f}·{materials.asphalt.thickness_m:.3f}·{geom.strip_length_m:.3f} = {asphalt_q:.3f} Tn/m",
        f"la carga DW aplicada entre x = {layout.asphalt_start_m:.3f} m y x = {layout.asphalt_end_m:.3f} m es {asphalt_q:.3f} Tn/m.",
        "La carpeta asfáltica se mantiene restringida al ancho de calzada definido en los datos del proyecto.",
        REF_DEAD_LOAD,
    )
    _calc(
        document,
        "Carga peatonal sobre veredas",
        "qPL = pPL·b",
        "qPL: carga lineal peatonal; pPL: carga peatonal superficial; b: ancho longitudinal de la franja",
        f"qPL = {live_loads.pedestrian.load_tn_m2:.4f}·{geom.strip_length_m:.3f} = {pedestrian_q:.4f} Tn/m",
        f"en cada vereda se aplica qPL = {pedestrian_q:.4f} Tn/m sobre un ancho lateral de {layout.sidewalk_width_m:.3f} m.",
        "La carga peatonal se considera simultáneamente con la sobrecarga vehicular cuando la combinación aplicable así lo requiere.",
        REF_PEDESTRIAN,
    )
    document.add_heading("2.3 Anchos de franja equivalente para carga vehicular", level=2)
    _calc(
        document,
        "Anchos equivalentes para carga vehicular",
        "E₊ = 0.660 + 0.55·S  ;  E₋ = 1.220 + 0.25·S",
        "E₊ y E₋: anchos equivalentes para momento positivo y negativo; S: separación de vigas, en m.",
        f"E₊ = 0.660 + 0.55·{geom.girder_spacing_m:.3f} = {data.transverse_result.equivalent_strip_width_positive_m:.3f} m\n"
        f"E₋ = 1.220 + 0.25·{geom.girder_spacing_m:.3f} = {data.transverse_result.equivalent_strip_width_negative_m:.3f} m.",
        "Las cargas de rueda se distribuyen con el ancho correspondiente al signo del momento evaluado.",
        "Se aplican los anchos normativos sin interpolaciones adicionales.",
        REF_STRIP,
    )
    heavy_axle = max(vehicle.design_truck_axles_tn)
    presence = data.transverse_result.ll_im_one_truck.multiple_presence_factor or 1.0
    impact = data.transverse_result.dynamic_load_allowance
    e_positive = data.transverse_result.equivalent_strip_width_positive_m
    e_negative = data.transverse_result.equivalent_strip_width_negative_m
    wheel_positive = heavy_axle / 2.0 * (1.0 + impact) * presence / e_positive
    wheel_negative = heavy_axle / 2.0 * (1.0 + impact) * presence / e_negative
    _calc(
        document,
        "Carga de rueda distribuida en la franja equivalente",
        "w(LL+IM) = (Pax/2)·(1+IM)·m/E",
        "w(LL+IM): carga lineal por línea de rueda; Pax: carga del eje considerado; IM: incremento dinámico; m: factor de presencia múltiple; E: ancho equivalente",
        f"w+ = ({heavy_axle:.3f}/2)·(1+{impact:.2f})·{presence:.2f}/{e_positive:.3f} = {wheel_positive:.3f} Tn/m\n"
        f"w− = ({heavy_axle:.3f}/2)·(1+{impact:.2f})·{presence:.2f}/{e_negative:.3f} = {wheel_negative:.3f} Tn/m",
        "las líneas de rueda se desplazan por el dominio permitido empleando el ancho equivalente correspondiente al signo del momento evaluado.",
        "El barrido conserva por separado los máximos positivos y negativos; por ello la envolvente no depende del sentido en que se enumeran las posiciones.",
        REF_HL93,
    )
    document.add_heading("2.4 Efectos críticos factorizados", level=2)
    _table(
        document,
        ("Dirección", "Combinación", "x crítica", "DC", "DW", "PL", "LL+IM", "Mu"),
        tuple(_moment_row_values(row, include_pl=True) for row in (positive, negative)),
        widths=(20, 30, 18, 17, 17, 17, 19, 20),
        font_size=7.8,
    )
    for row, title in ((positive, "Momento positivo crítico"), (negative, "Momento negativo crítico")):
        _moment_combination_calc(document, title, row, include_pl=True)
    document.add_heading("2.5 Diseño del acero principal", level=2)
    for design, prefix in ((reinf.positive, "B."), (reinf.negative, "A.")):
        option = selected_option(data.slab_selected, prefix)
        _flexural_calc(document, design.label, design, option, concrete, steel, strip=True)
    temp = selected_option(data.slab_selected, "C.")
    _calc(
        document,
        "Acero por contracción y temperatura",
        "As,temp = ρtemp·Ag",
        "ρtemp: cuantía adoptada; Ag: área bruta de la franja de 1 m.",
        f"As,temp = {reinf.temperature.ratio:.6f}·(100·{geom.slab_thickness_m*100:.2f}) = {reinf.temperature.required_area_cm2_m:.3f} cm²/m.",
        f"Se adopta {_option_text(temp)}.",
        _compliance_comment(temp, "El acero transversal mínimo controla la respuesta frente a cambios volumétricos."),
        REF_TEMP,
    )
    distribution = selected_option(data.slab_selected, "D.")
    _calc(
        document,
        "Acero longitudinal de distribución",
        "% = min[67; 3840/√S]  ;  As,dist = (%/100)·As,req,+",
        "S: luz efectiva en mm; As,req,+: acero principal positivo requerido (no el adoptado); As,dist: acero longitudinal requerido.",
        f"% = min[67; 3840/√({reinf.distribution.effective_span_m*1000:.1f})] = {reinf.distribution.percent_of_positive_steel:.2f}%\n"
        f"As,req,+ = {reinf.distribution.positive_main_area_cm2_m:.3f} cm²/m\n"
        f"As,dist = {reinf.distribution.percent_of_positive_steel/100:.4f}·{reinf.distribution.positive_main_area_cm2_m:.3f} = {reinf.distribution.required_area_cm2_m:.3f} cm²/m",
        f"Se adopta {_option_text(distribution)}.",
        _compliance_comment(
            distribution,
            "La distribución longitudinal se calcula sobre As,req del acero positivo, conforme al criterio normativo de acero requerido.",
        ),
        REF_DISTRIBUTION,
    )
    document.add_heading("2.6 Control de fisuración", level=2)
    for crack in (data.slab_crack.positive_main, data.slab_crack.negative_main):
        _crack_calc(document, crack, slab_thickness_m=geom.slab_thickness_m)
    document.add_heading("2.7 Diagramas de diseño", level=2)
    moment_path = save_strength_chart(
        data.transverse_result,
        chart_dir / "losa_momentos.png",
        title="Losa transversal · envolventes de momento Resistencia I",
        kind="moment",
        include_pl=True,
        transverse=True,
    )
    _picture(document, moment_path, "Figura 2.1. Envolventes factorizadas de momento empleadas en el diseño de la losa.")
    _body(
        document,
        "Para la losa transversal el diseño principal se basa en flexión. Las reacciones "
        "de apoyo del modelo se conservan para la transferencia a las vigas; el cortante "
        "local no define un refuerzo independiente en este módulo."
    )


def _girder(document: Document, data: DeckReportData, *, exterior: bool, chart_dir: Path) -> None:
    number = 4 if exterior else 3
    label = "Viga principal exterior" if exterior else "Viga principal interior"
    inputs = data.project_inputs
    geometry = inputs.exterior_girder if exterior else inputs.interior_girder
    analysis = data.exterior_result if exterior else data.interior_result
    reinforcement = data.exterior_reinforcement if exterior else data.interior_reinforcement
    shear = data.exterior_shear if exterior else data.interior_shear
    selected = data.exterior_selected if exterior else data.interior_selected
    girder_detail = data.exterior_detail if exterior else data.interior_detail
    crack = data.exterior_crack.main if exterior else data.interior_crack.main
    fatigue = data.exterior_fatigue if exterior else data.interior_fatigue
    service = data.exterior_service if exterior else data.interior_service
    include_pl = exterior
    moment_rows = (
        combine_exterior_girder_moments(analysis)
        if exterior else combine_interior_girder_moments(analysis)
    )
    strength_moment = max(
        (row for row in moment_rows if row.limit_state == "Resistencia"),
        key=lambda row: row.combined_moment_tn_m,
    )
    shear_rows = (
        combine_exterior_girder_shears(geometry, analysis, reinforcement.parameters)
        if exterior else combine_interior_girder_shears(geometry, analysis, reinforcement.parameters)
    )
    critical_shear = max(shear_rows, key=lambda row: row.combined_shear_tn)
    document.add_heading(f"{number}. {label}", level=1)
    _body(
        document,
        "La viga se modela como un elemento simplemente apoyado. Las cargas permanentes "
        "se integran sobre sus anchos tributarios; el camión y el tándem HL-93 se desplazan "
        "longitudinalmente y sus efectos se multiplican por los factores de distribución "
        "transversal correspondientes."
    )
    document.add_heading(f"{number}.1 Geometría y distribución de carga viva", level=2)
    _table(
        document,
        ("Parámetro", "Valor"),
        (
            ("Luz L", f"{geometry.span_length_m:.3f} m"),
            ("Separación S", f"{geometry.girder_spacing_m:.3f} m"),
            ("Alma bw", f"{geometry.web_width_m:.3f} m"),
            ("Altura de viga bajo losa", f"{geometry.girder_total_height_m:.3f} m"),
            ("Espesor de losa", f"{geometry.slab_thickness_m:.3f} m"),
            ("Peralte resistente total h", f"{geometry.total_t_section_depth_m:.3f} m"),
            ("Factor g de momento", f"{analysis.distribution_factor_g:.5f}"),
            ("Factor g de cortante", f"{analysis.shear_distribution_factor_g:.5f}"),
        ),
        widths=(105, 62),
    )
    initial_effective_depth_cm = (
        geometry.total_t_section_depth_m * 100.0
        - reinforcement.parameters.concrete_cover_cm
        - reinforcement.parameters.main_bar_diameter_cm / 2.0
    )
    _calc(
        document,
        f"Peralte resistente de la {label.lower()}",
        "h = hbajo + tlosa; d = h − rec − db/2",
        "hbajo: altura de concreto de la viga bajo la losa; tlosa: espesor de losa monolítica; "
        "rec: recubrimiento; db: diámetro supuesto de la barra principal; "
        "d: peralte efectivo inicial.",
        f"h = {geometry.girder_total_height_m:.3f} + {geometry.slab_thickness_m:.3f} = "
        f"{geometry.total_t_section_depth_m:.3f} m\n"
        f"d = {geometry.total_t_section_depth_m * 100.0:.2f} − "
        f"{reinforcement.parameters.concrete_cover_cm:.2f} − "
        f"{reinforcement.parameters.main_bar_diameter_cm:.3f}/2 = "
        f"{initial_effective_depth_cm:.2f} cm.",
        f"El peralte efectivo inicial es d={initial_effective_depth_cm:.2f} cm. "
        f"Con el centroide de las capas adoptadas se utiliza d={reinforcement.main.effective_depth_cm:.2f} cm.",
        "El peso propio del alma usa hbajo; el peso de la losa se contabiliza separadamente.",
        REF_FLEX,
    )
    vehicle = inputs.live_loads.vehicular
    if exterior:
        g_formula, g_legend, g_sub, g_result, g_comment = detail.exterior_g_factor_trace(geometry, vehicle)
    else:
        g_formula, g_legend, g_sub, g_result, g_comment = detail.interior_g_factor_trace(geometry)
    _calc(
        document,
        "Cálculo de factores de distribución gM y gV",
        g_formula,
        g_legend,
        g_sub,
        g_result,
        g_comment,
        REF_DISTRIBUTION_FACTORS,
    )
    materials = inputs.materials
    q_slab = (
        materials.concrete.specific_weight_tn_m3
        * geometry.slab_thickness_m
        * geometry.tributary_width_m
    )
    q_web = (
        materials.concrete.specific_weight_tn_m3
        * geometry.web_width_m
        * geometry.girder_total_height_m
    )
    extra_components = 0.0
    extra_substitution = ""
    if exterior:
        q_sidewalk = (
            materials.sidewalk.specific_weight_tn_m3
            * materials.sidewalk.thickness_m
            * geometry.sidewalk_tributary_width_m
        )
        q_railing = materials.railing.weight_kg_m / 1000.0
        q_barrier = materials.barrier.weight_kg_m / 1000.0
        extra_components = q_sidewalk + q_railing + q_barrier
        extra_substitution = (
            f"\nqver = {materials.sidewalk.specific_weight_tn_m3:.3f}·{materials.sidewalk.thickness_m:.3f}·{geometry.sidewalk_tributary_width_m:.3f} = {q_sidewalk:.3f} Tn/m"
            f"\nqbaranda = {materials.railing.weight_kg_m:.3f}/1000 = {q_railing:.3f} Tn/m"
            f"\nqbarrera = {materials.barrier.weight_kg_m:.3f}/1000 = {q_barrier:.3f} Tn/m"
        )
    q_dc = q_slab + q_web + extra_components
    _calc(
        document,
        "Cargas permanentes distribuidas sobre la viga",
        "qDC = qlosa + qalma + qver + qbaranda + qbarrera",
        "qDC: carga permanente lineal total; qlosa: peso de la losa tributaria; qalma: peso del alma; qver: peso de vereda; qbaranda y qbarrera: pesos lineales de accesorios",
        f"qlosa = {materials.concrete.specific_weight_tn_m3:.3f}·{geometry.slab_thickness_m:.3f}·{geometry.tributary_width_m:.3f} = {q_slab:.3f} Tn/m\n"
        f"qalma = {materials.concrete.specific_weight_tn_m3:.3f}·{geometry.web_width_m:.3f}·{geometry.girder_total_height_m:.3f} = {q_web:.3f} Tn/m"
        f"{extra_substitution}\n"
        f"qDC = {q_dc:.3f} Tn/m",
        f"la carga permanente uniformemente distribuida utilizada para la {label.lower()} es qDC = {q_dc:.3f} Tn/m, además de los pesos concentrados de los diafragmas.",
        "Los elementos se asignan de acuerdo con su ancho tributario; los accesorios de borde sólo se incorporan en la viga exterior.",
        REF_DEAD_LOAD,
    )
    diaphragm_loads = tuple(
        materials.concrete.specific_weight_tn_m3
        * diaphragm.thickness_m
        * diaphragm.height_m
        * diaphragm.tributary_width_m
        for diaphragm in geometry.diaphragms
    )
    if diaphragm_loads:
        diaphragm = geometry.diaphragms[0]
        _calc(
            document,
            "Peso concentrado de los diafragmas",
            "Pd = γc·td·hd·bd",
            "Pd: peso de un diafragma asignado a la viga; γc: peso específico del concreto; td: espesor; hd: altura; bd: ancho tributario longitudinal",
            f"Pd = {materials.concrete.specific_weight_tn_m3:.3f}·{diaphragm.thickness_m:.3f}·{diaphragm.height_m:.3f}·{diaphragm.tributary_width_m:.3f} = {diaphragm_loads[0]:.3f} Tn",
            f"se aplican {len(diaphragm_loads)} carga(s) concentrada(s) en las posiciones geométricas de los diafragmas; la suma asignada es {sum(diaphragm_loads):.3f} Tn.",
            "Estas cargas intervienen en las reacciones, momentos y cortantes del caso DC sin redistribuirse como carga uniforme.",
            REF_DEAD_LOAD,
        )
    asphalt_width = (
        geometry.asphalt_tributary_width_m if exterior else geometry.tributary_width_m
    )
    q_dw = (
        materials.asphalt.specific_weight_tn_m3
        * materials.asphalt.thickness_m
        * asphalt_width
    )
    _calc(
        document,
        "Carga longitudinal de la superficie de rodadura",
        "qDW = γasf·tasf·basf",
        "qDW: carga lineal de superficie de rodadura; γasf: peso específico; tasf: espesor; basf: ancho tributario de asfalto",
        f"qDW = {materials.asphalt.specific_weight_tn_m3:.3f}·{materials.asphalt.thickness_m:.3f}·{asphalt_width:.3f} = {q_dw:.3f} Tn/m",
        f"la carga uniforme del caso DW para la {label.lower()} es qDW = {q_dw:.3f} Tn/m.",
        (
            "El ancho tributario de la viga exterior se limita a la porción efectiva de carpeta que intersecta su franja tributaria."
            if exterior
            else "La viga interior recibe la superficie de rodadura correspondiente a todo su ancho tributario."
        ),
        REF_DEAD_LOAD,
    )
    if exterior:
        q_pl = inputs.live_loads.pedestrian.load_tn_m2 * geometry.sidewalk_tributary_width_m
        _calc(
            document,
            "Carga peatonal tributaria de la viga exterior",
            "qPL = pPL·bver",
            "qPL: carga lineal peatonal; pPL: carga superficial sobre la vereda; bver: ancho tributario de vereda",
            f"qPL = {inputs.live_loads.pedestrian.load_tn_m2:.4f}·{geometry.sidewalk_tributary_width_m:.3f} = {q_pl:.3f} Tn/m",
            f"la viga exterior recibe qPL = {q_pl:.3f} Tn/m en toda la luz.",
            "La carga peatonal se mantiene como acción PL independiente para aplicar el factor que corresponde en cada estado límite.",
            REF_PEDESTRIAN,
        )
    vehicle = inputs.live_loads.vehicular
    _table(
        document,
        ("Modelo móvil", "Cargas por eje", "Separación / dato longitudinal"),
        (
            ("Camión de diseño", " – ".join(f"{value:.3f} Tn" for value in vehicle.design_truck_axles_tn), f"{vehicle.design_truck_spacings_m[0]:.3f} m y separación posterior variable"),
            ("Tándem de diseño", " – ".join(f"{value:.3f} Tn" for value in vehicle.design_tandem_axles_tn), f"{vehicle.design_tandem_spacing_m:.3f} m"),
            ("Carga de carril", f"{vehicle.lane_load_tn_m:.3f} Tn/m", "Aplicada sobre la luz cargada"),
        ),
        widths=(42, 55, 70),
        font_size=8.1,
    )
    _calc(
        document,
        "Distribución transversal de carga viva",
        "Mᵥᵢgₐ(x) = gM·MHL93(x)  ;  Vᵥᵢgₐ(x) = gV·VHL93(x)",
        "gM y gV: factores de distribución para momento y cortante; MHL93 y VHL93: efectos del modelo vehicular HL-93 sobre una línea de ruedas antes de distribuir.",
        f"gM = {analysis.distribution_factor_g:.5f}\n"
        f"gV = {analysis.shear_distribution_factor_g:.5f}\n"
        f"Los factores se aplican a cada posición del camión y tándem antes de formar la envolvente.",
        "La envolvente selecciona, en cada estación, la configuración vehicular que produce el efecto más desfavorable.",
        "El procedimiento mantiene separadas las reglas de distribución de momento y cortante.",
        REF_DISTRIBUTION_FACTORS,
    )
    g_m = float(analysis.distribution_factor_g)
    g_v = float(analysis.shear_distribution_factor_g)
    ll_formula, ll_legend, ll_sub, ll_result, ll_comment = detail.ll_im_with_g_trace(
        analysis=analysis,
        strength_moment=strength_moment,
        critical_shear=critical_shear,
        g_m=g_m,
        g_v=g_v,
    )
    _calc(
        document,
        "Aplicación numérica de gM y gV en LL+IM",
        ll_formula,
        ll_legend,
        ll_sub,
        ll_result,
        ll_comment,
        REF_DISTRIBUTION_FACTORS,
    )
    document.add_heading(f"{number}.2 Efectos de Resistencia I", level=2)
    _moment_combination_calc(document, "Momento positivo crítico", strength_moment, include_pl=include_pl)
    _shear_combination_calc(document, "Cortante crítico", critical_shear, include_pl=include_pl)
    document.add_heading(f"{number}.3 Flexión y acero adoptado", level=2)
    main = selected_option(selected, "A.")
    _flexural_calc(
        document,
        "Acero longitudinal principal",
        reinforcement.main,
        main,
        inputs.materials.concrete,
        inputs.materials.steel,
        strip=False,
    )
    temperature = selected_option(selected, "B.")
    skin = selected_option(selected, "C.")
    _calc(
        document,
        "Acero de temperatura en caras laterales",
        "As,temp = ρtemp·bw·100 / 2",
        "ρtemp: cuantía; bw: ancho del alma; la división asigna la mitad a cada cara.",
        f"As,temp = {reinforcement.temperature.ratio:.6f}·{reinforcement.temperature.web_width_cm:.2f}·100/2 = {reinforcement.temperature.required_area_cm2_m_per_face:.3f} cm²/m por cara",
        f"Se adopta {_option_text(temperature)}.",
        _compliance_comment(temperature, "El refuerzo se coloca en ambas caras del alma."),
        REF_TEMP,
    )
    ask_formula, ask_legend, ask_sub, ask_result, ask_comment = detail.skin_steel_trace(
        reinforcement.skin,
        _option_text(skin),
        _compliance_comment(skin, "Este acero controla la fisuración lateral asociada a la altura del alma."),
    )
    _calc(
        document,
        "Acero longitudinal de piel",
        ask_formula,
        ask_legend,
        ask_sub,
        ask_result,
        ask_comment,
        REF_CRACK,
    )
    document.add_heading(f"{number}.4 Diseño por cortante", level=2)
    stirrup = selected_option(selected, "D.")
    _calc(
        document,
        "Resistencia seccional a cortante",
        *_sectional_shear_calc_fields(shear, stirrup),
        REF_SHEAR,
    )
    document.add_heading(f"{number}.5 Servicio, fisuración y fatiga", level=2)
    _crack_calc(document, crack)
    svc_formula, svc_legend, svc_sub, svc_result, svc_comment = detail.service_stress_trace(
        service,
        inputs.materials.concrete,
        inputs.materials.steel,
    )
    _calc(
        document,
        "Tensiones de Servicio I",
        svc_formula,
        svc_legend,
        svc_sub,
        svc_result,
        svc_comment,
        "Manual de Puentes MTC 2018, criterios de Servicio I y análisis elástico de sección fisurada; combinación de Servicio I.",
    )
    fat_formula, fat_legend, fat_sub, fat_result, fat_comment = detail.fatigue_trace(
        fatigue,
        inputs.materials.steel,
    )
    _calc(
        document,
        "Fatiga del acero longitudinal",
        fat_formula,
        fat_legend,
        fat_sub,
        fat_result,
        fat_comment,
        REF_FATIGUE,
    )
    document.add_heading(f"{number}.6 Desarrollo y detalle constructivo", level=2)
    _calc(
        document,
        "Longitud de desarrollo del acero principal",
        "ld = max[ldb·λubic·λrec·λhorm·λconf·(As,req/As,prov); 30.48 cm]",
        "ldb: longitud básica; λ: modificadores normativos; As,req/As,prov: reducción por exceso de acero.",
        f"Para {girder_detail.selected_main_bar.bar_count} barras de {girder_detail.selected_main_bar.bar_label}\n"
        f"As,prov = {girder_detail.selected_main_bar.provided_area_cm2:.3f} cm²\n"
        f"ld = {girder_detail.development_length_m:.3f} m",
        f"Se mantienen {girder_detail.continuous_bar_count} barras continuas y se definen {girder_detail.physical_cut_count} cortes físicos con prolongación de desarrollo.",
        "Los puntos teóricos de corte se desplazan mediante ld; si la longitud disponible no fuera suficiente, las barras deberían continuarse o anclarse.",
        REF_DEVELOPMENT,
    )
    document.add_heading(f"{number}.7 Diagramas de diseño", level=2)
    slug = "exterior" if exterior else "interior"
    moment_path = save_strength_chart(
        analysis,
        chart_dir / f"{slug}_momentos.png",
        title=f"{label} · envolventes de momento Resistencia I",
        kind="moment",
        include_pl=include_pl,
    )
    shear_path = save_strength_chart(
        analysis,
        chart_dir / f"{slug}_cortantes.png",
        title=f"{label} · envolventes de cortante Resistencia I",
        kind="shear",
        include_pl=include_pl,
    )
    _picture(document, moment_path, f"Figura {number}.1. Envolventes factorizadas de momento para {label.lower()}.")
    _picture(document, shear_path, f"Figura {number}.2. Envolventes factorizadas de cortante para {label.lower()}.")


def _barrier(document: Document, data: DeckReportData) -> None:
    inputs = data.project_inputs.barrier
    result = data.barrier_result
    flex = result.flexure
    yl = result.yield_line
    geom = inputs.geometry
    impact = inputs.impact_load
    document.add_page_break()
    document.add_heading("5. Barrera de concreto", level=1)
    _body(
        document,
        "La barrera se verifica mediante el mecanismo de líneas de fluencia del paño "
        "impactado y mediante la transferencia por corte–fricción en la interfaz con la losa."
    )
    document.add_heading("5.1 Resistencia flexional del mecanismo", level=2)
    rows = tuple(
        (
            "Mw " + component.label,
            f"{component.steel_area_cm2:.3f}",
            f"{component.effective_depth_cm:.2f}",
            f"{component.compression_block_depth_cm:.3f}",
            f"{component.nominal_moment_tn_m:.3f}",
        )
        for component in flex.mw_components
    )
    _table(document, ("Componente", "As (cm²)", "d (cm)", "a (cm)", "Mn (Tn·m)"), rows, widths=(45, 30, 30, 30, 32))
    mw_formula, mw_legend, mw_sub, mw_result, mw_comment = detail.barrier_mw_mc_trace(
        flex,
        data.project_inputs.materials.concrete,
        data.project_inputs.materials.steel,
    )
    _calc(
        document,
        "Momento nominal por componente",
        mw_formula,
        mw_legend,
        mw_sub,
        mw_result,
        mw_comment,
        REF_BARRIER,
    )
    document.add_heading("5.2 Línea de fluencia y resistencia transversal", level=2)
    multiplier = 8.0 if yl.impact_pattern == "segment" else 1.0
    _calc(
        document,
        "Longitud crítica del mecanismo",
        "Lc = Lt/2 + √[(Lt/2)² + k·H·(Mb + Mw)/Mc]",
        "Lt: longitud de distribución del impacto; H: altura; Mb, Mw y Mc: resistencias flexionales; k = 8 para impacto interior.",
        f"Lc = {impact.distribution_length_m:.3f}/2 + √[({impact.distribution_length_m:.3f}/2)² + {multiplier:.0f}·{geom.height_m:.3f}·({geom.top_additional_moment_tn_m:.3f}+{flex.mw_tn_m:.3f})/{flex.mc_tn_m:.3f}] = {yl.critical_length_m:.3f} m.",
        f"La longitud crítica del mecanismo de líneas de fluencia es Lc = {yl.critical_length_m:.3f} m.",
        "Lc delimita el tramo de barrera que participa en el mecanismo resistente al impacto.",
        REF_BARRIER,
    )
    _calc(
        document,
        "Resistencia nominal por líneas de fluencia",
        "Rw = 2·[k·Mb + k·Mw + Mc·Lc²/H]/(2·Lc − Lt)",
        "Rw: resistencia transversal nominal de la barrera; "
        "Mb: resistencia flexional adicional del elemento superior; "
        "Mw: resistencia flexional de la pared alrededor de su eje vertical; "
        "Mc: resistencia flexional de la pared como voladizo respecto de su base; "
        "Lc: longitud crítica del mecanismo; "
        "Lt: longitud de distribución del impacto; "
        "H: altura de la barrera; "
        "k: coeficiente del patrón de impacto, igual a 8 para un tramo interior y 1 para un extremo o junta.",
        f"k = {multiplier:.0f}\n"
        f"Mb = {geom.top_additional_moment_tn_m:.3f} Tn·m\n"
        f"Mw = {flex.mw_tn_m:.3f} Tn·m\n"
        f"Mc = {flex.mc_tn_m:.3f} Tn·m/m\n"
        f"Lc = {yl.critical_length_m:.3f} m\n"
        f"Lt = {impact.distribution_length_m:.3f} m\n"
        f"H = {geom.height_m:.3f} m\n"
        f"Rw = 2·[{multiplier:.0f}·{geom.top_additional_moment_tn_m:.3f} + "
        f"{multiplier:.0f}·{flex.mw_tn_m:.3f} + "
        f"{flex.mc_tn_m:.3f}·({yl.critical_length_m:.3f})²/{geom.height_m:.3f}]"
        f"/[2·{yl.critical_length_m:.3f} − {impact.distribution_length_m:.3f}] "
        f"= {yl.nominal_transverse_resistance_tn:.3f} Tn.",
        f"Rw = {yl.nominal_transverse_resistance_tn:.3f} Tn frente a "
        f"Ft = {yl.demand_transverse_force_tn:.3f} Tn: {yl.resistance_status}.",
        _status_comment(yl.resistance_status, "La resistencia transversal del mecanismo supera la demanda de impacto."),
        REF_BARRIER,
    )
    document.add_heading("5.3 Interfaz, dowels y anclaje", level=2)
    shear = result.shear_transfer
    iface_formula, iface_legend, iface_sub, iface_result, iface_comment = detail.barrier_interface_trace(
        inputs,
        data.project_inputs.materials,
        shear,
    )
    _calc(
        document,
        "Transferencia por corte–fricción",
        iface_formula,
        iface_legend,
        iface_sub,
        iface_result,
        iface_comment,
        REF_BARRIER,
    )
    dowel = result.dowel
    development = result.development
    _calc(
        document,
        "Cuantía mínima y desarrollo de dowels",
        "Avf,min = 3.52·Acv/fy  ;  ldh = max[0.076·db·fy/√f'c·λ; 8db; 15.24 cm]",
        "Avf,min: acero mínimo de interfaz; db: diámetro de dowel; λ: producto de modificadores del gancho.",
        f"Avf,min = {dowel.required_avf_cm2_m:.3f} cm²/m\n"
        f"Avf,prov = {dowel.provided_avf_cm2_m:.3f} cm²/m: {dowel.status}\n"
        f"ldh,req = {development.required_ldh_cm:.2f} cm\n"
        f"ldisp = {development.available_length_cm:.2f} cm: {development.status}",
        f"La extensión del gancho adoptada es {development.hook_extension_cm:.2f} cm.",
        _status_comment(development.status, "El anclaje disponible desarrolla la barra de conexión adoptada."),
        REF_BARRIER,
    )


def _cantilever(document: Document, data: DeckReportData) -> None:
    result = data.cantilever_result
    geometry = data.project_inputs.transverse_slab.geometry
    flex = result.flexural_steel
    option = selected_option(data.cantilever_selected, "A.")
    temperature = selected_option(data.cantilever_selected, "B.")
    document.add_heading("6. Losa en voladizo", level=1)
    _body(
        document,
        "El voladizo se verifica en la raíz de la losa exterior. Se incluyen peso propio, "
        "carpeta, vereda, barrera, carga vehicular y el evento extremo de colisión cuando corresponde."
    )
    document.add_heading("6.1 Efectos y combinación gobernante", level=2)
    _table(
        document,
        ("Acción", "Carga", "Brazo", "Momento en raíz"),
        tuple(
            (effect.label, f"{effect.load_tn:.3f} Tn", f"{effect.arm_to_root_m:.3f} m", f"{effect.root_moment_tn_m:.3f} Tn·m")
            for effect in result.load_effects
        ),
        widths=(68, 31, 30, 38),
        font_size=8.1,
    )
    arm_formula, arm_legend, arm_sub, arm_result, arm_comment = detail.cantilever_arm_trace(
        result.load_effects,
        geometry.overhang_m,
    )
    _calc(
        document,
        "Brazos de carga respecto a la raíz",
        arm_formula,
        arm_legend,
        arm_sub,
        arm_result,
        arm_comment,
        REF_DEAD_LOAD,
    )
    control = result.controlling_strength
    _calc(
        document,
        "Momento factorizado en la raíz",
        "Mu = γDC·MDC + γDW·MDW + γPL·MPL + γLL·M(LL+IM)",
        "γ: factores LRFD; M: contribución con signo en la raíz del voladizo.",
        f"Mu = {control.dc_factor:.2f}·({control.dc_moment_tn_m:.3f}) + {control.dw_factor:.2f}·({control.dw_moment_tn_m:.3f}) + "
        f"{control.pl_factor:.2f}·({control.pl_moment_tn_m:.3f}) + {control.ll_im_factor:.2f}·({control.ll_im_moment_tn_m:.3f}) = {control.combined_moment_tn_m:.3f} Tn·m.",
        f"Demanda de diseño |Mu| = {control.design_moment_tn_m:.3f} Tn·m, combinación {control.combination_name}.",
        "El signo identifica tracción superior en la raíz; por ello el acero principal se coloca en la cara superior.",
        REF_COMB,
    )
    collision = result.barrier_collision
    for note in result.applicability_notes:
        _comment(document, note)
    document.add_heading("6.2 Colisión y transferencia de la barrera", level=2)
    if collision is not None:
        _comment(document, f"N simultanea={collision.axial_tension_tn_m:.3f} Tn/m; As incluye N/(phi*fy). {collision.status}: {collision.scope_note}")
        _table(document, ("Caso", "M Tn.m/m", "N Tn/m", "V Tn/m", "Estado"),
               tuple((case.name, f"{case.moment_tn_m:.3f}", f"{case.axial_tension_tn_m:.3f}", f"{case.vertical_force_tn_m:.3f}", case.status) for case in collision.cases),
               widths=(75, 30, 30, 30, 25))
        col_formula, col_legend, col_sub, col_result, col_comment = detail.cantilever_collision_trace(
            collision,
            control,
        )
        _calc(
            document,
            "Evento Extremo II — colisión de la barrera",
            col_formula,
            col_legend,
            col_sub,
            col_result,
            col_comment,
            "Manual de Puentes MTC 2018, Art. 2.4.3.5.1.2; mecanismo y transferencia de barrera: Rodríguez Serquén.",
        )
    elif result.interior_collision is not None:
        _interior_collision_docx(document, result.interior_collision)
    else:
        _comment(
            document,
            "No existe una verificación local de colisión aplicable a la posición declarada de la barrera.",
        )
    document.add_heading("6.3 Flexión y acero adoptado", level=2)
    _flexural_calc(document, "Acero superior en la raíz", flex, option, data.project_inputs.materials.concrete, data.project_inputs.materials.steel, strip=True)
    _calc(
        document,
        "Acero por temperatura del voladizo",
        "As,temp = ρtemp·Ag",
        "Ag: área bruta por metro; ρtemp: cuantía adoptada.",
        f"As,temp = {result.temperature_steel.ratio:.6f}·(100·{geometry.slab_thickness_m*100:.2f}) = {result.temperature_steel.required_area_cm2_m:.3f} cm²/m.",
        f"Se adopta {_option_text(temperature)}.",
        _compliance_comment(temperature, "La disposición se mantiene en la dirección secundaria del voladizo."),
        REF_TEMP,
    )
    document.add_heading("6.4 Cortante, fisuración y desarrollo", level=2)
    shear = result.shear
    sh_formula, sh_legend, sh_sub, sh_result, sh_comment = detail.cantilever_shear_trace(shear)
    _calc(
        document,
        "Cortante unidireccional en la raíz",
        sh_formula,
        sh_legend,
        sh_sub,
        sh_result,
        sh_comment,
        REF_SHEAR,
    )
    _crack_calc(document, data.cantilever_crack, slab_thickness_m=geometry.slab_thickness_m)
    dev = data.cantilever_development
    cover_cm = getattr(result.parameters, "concrete_cover_cm", 5.0)
    dev_formula, dev_legend, dev_sub, dev_result, dev_comment = detail.cantilever_development_trace(
        dev,
        geometry.overhang_m,
        cover_cm=cover_cm,
    )
    _calc(
        document,
        "Desarrollo de la barra superior adoptada",
        dev_formula,
        dev_legend,
        dev_sub,
        dev_result,
        dev_comment,
        REF_DEVELOPMENT,
    )


def _interior_collision_docx(document: Document, collision) -> None:
    document.add_heading("6.2.1 Transferencia local a la losa interior", level=3)
    _calc(
        document,
        "Demanda transmitida por la barrera",
        "F = max(Ft,Rw); L = Lc + 2H; N = F/L; Minterfaz = Mc·max(1,Ft/Rw)",
        "Ft: fuerza de ensayo; Rw: resistencia por líneas de fluencia; L: longitud de transferencia; N: tracción por metro.",
        f"F={collision.horizontal_force_tn:.3f} Tn; L={collision.transfer_length_m:.3f} m; "
        f"N={collision.horizontal_force_tn/collision.transfer_length_m:.3f} Tn/m; "
        f"Minterfaz={collision.interface_moment_tn_m_m:.3f} Tn·m/m.",
        "Las acciones horizontal y vertical se aplican independientemente en ambos bordes de cada barrera.",
        "La demanda se distribuye mediante una franja elástica continua de losa de 1.0 m apoyada en los ejes de las vigas.",
        "MTC 2018, Arts. 2.4.3.5.1.2 y 2.4.5.3; Tabla 2.4.3.6.3-1.",
    )
    _table(
        document,
        ("Caso independiente", "x", "M aplicado", "Fv", "N"),
        tuple(
            (
                case.name,
                f"{case.position_m:.3f}",
                f"{case.applied_moment_tn_m_m:.3f}",
                f"{case.applied_vertical_tn_m:.3f}",
                f"{case.axial_tension_tn_m:.3f}",
            )
            for case in collision.cases
        ),
        widths=(82, 21, 27, 24, 24),
        font_size=7.6,
    )
    reaction_headers = (
        "Caso",
        *(label for label, _ in collision.cases[0].reactions_tn_m),
        "Err F",
        "Err M",
    )
    _table(
        document,
        reaction_headers,
        tuple(
            (
                case.name,
                *(f"{reaction:.3f}" for _, reaction in case.reactions_tn_m),
                f"{case.vertical_equilibrium_error_tn_m:.1e}",
                f"{case.moment_equilibrium_error_tn_m_m:.1e}",
            )
            for case in collision.cases
        ),
        widths=(58,) + (17,) * (len(reaction_headers) - 1),
        font_size=7.0,
    )
    document.add_heading("6.2.2 Acero transversal por colisión", level=3)
    _table(
        document,
        ("Cara", "Caso gobernante", "x", "M", "N", "γDC", "γDW"),
        tuple(
            (
                face.face,
                face.case_name,
                f"{face.position_m:.3f}",
                f"{face.moment_tn_m_m:.3f}",
                f"{face.tension_tn_m:.3f}",
                f"{face.dc_factor:.2f}",
                f"{face.dw_factor:.2f}",
            )
            for face in collision.faces
        ),
        widths=(22, 70, 20, 22, 22, 18, 18),
        font_size=7.4,
    )
    _table(
        document,
        ("Cara", "As requerido", "Acero adoptado", "As provisto", "ld requerido", "ld disponible", "Estado"),
        tuple(
            (
                face.face,
                f"{face.option.required_area_cm2_m:.3f}" if face.option else "—",
                f"{face.option.bar.label} @ {face.option.spacing_m:.3f} m" if face.option else "Sin opción",
                f"{face.option.provided_area_cm2_m:.3f}" if face.option else "—",
                f"{face.development_length_m:.3f} m",
                f"{face.development_available_m:.3f} m",
                face.status,
            )
            for face in collision.faces
        ),
        widths=(22, 26, 40, 27, 28, 29, 24),
        font_size=7.4,
    )
    document.add_heading("6.2.3 Cortante y conexión barrera–losa", level=3)
    _calc(
        document,
        "Cortante de la franja interior",
        "φVc = φ·0.265·β·√f'c·b·dv",
        "β: factor del procedimiento general con tracción; b=100 cm; dv: peralte efectivo de corte.",
        f"Vu={collision.shear_demand_tn_m:.3f} Tn/m; β={collision.shear_beta:.3f}; "
        f"φVc={collision.shear_capacity_tn_m:.3f} Tn/m.",
        f"Vu={collision.shear_demand_tn_m:.3f} Tn/m ≤ φVc={collision.shear_capacity_tn_m:.3f} Tn/m: "
        f"{collision.shear_status}.",
        _status_comment(collision.shear_status, "La franja interior satisface el control de cortante sin estribos."),
        "Manual de Puentes MTC 2018, Art. 2.9.1.5.6.",
    )
    _table(
        document,
        ("Control", "Demanda", "Capacidad disponible", "Estado"),
        tuple(
            (
                check.control,
                f"{check.demand:.3f} {check.unit}",
                f"{check.capacity:.3f} {check.unit}",
                check.status,
            )
            for check in collision.connection_checks
        ),
        widths=(58, 39, 48, 30),
        font_size=8.0,
    )
    _comment(
        document,
        f"Resultado local: {collision.status}. Conexión barrera–losa: {collision.connection_status}.",
    )
    for note in collision.notes:
        _comment(document, note)


def _diaphragm(document: Document, data: DeckReportData, chart_dir: Path) -> None:
    geometry = data.project_inputs.diaphragm
    result = data.diaphragm_result
    reinforcement = data.diaphragm_reinforcement
    rows = [row for row in combine_diaphragm_moments(result) if row.limit_state == "Resistencia"]
    positive = max(rows, key=lambda row: row.combined_moment_tn_m)
    negative = min(rows, key=lambda row: row.combined_moment_tn_m)
    shear = max(combine_diaphragm_shears(result), key=lambda row: row.combined_shear_tn)
    document.add_heading("7. Viga diafragma", level=1)
    _body(
        document,
        "El diafragma se analiza transversalmente como viga continua entre vigas principales. "
        "La altura ingresada corresponde al concreto bajo la losa; el peralte resistente total "
        "incluye el espesor de la losa monolítica. "
        "Las cargas permanentes y peatonales permanecen fijas; las líneas de ruedas se ubican "
        "desde ambos accesos físicos al tablero y se combinan en una única envolvente superior e inferior."
    )
    document.add_heading("7.1 Modelo y estaciones críticas", level=2)
    _table(
        document,
        ("Parámetro", "Valor"),
        (
            ("Longitud transversal", f"{geometry.total_width_m:.3f} m"),
            ("Separación entre apoyos", f"{geometry.girder_spacing_m:.3f} m"),
            ("Espesor longitudinal b", f"{geometry.thickness_m:.3f} m"),
            ("Altura bajo losa", f"{geometry.height_m:.3f} m"),
            ("Espesor de losa", f"{geometry.slab_thickness_m:.3f} m"),
            ("Peralte resistente total h", f"{geometry.total_depth_m:.3f} m"),
            ("Paso vehicular", f"{geometry.vehicle_step_m:.3f} m"),
        ),
        widths=(105, 62),
    )
    _calc(
        document,
        "Peralte resistente del diafragma",
        "h = hbajo + tlosa; d = h − rec − db/2",
        "hbajo: altura de concreto del diafragma bajo la losa; tlosa: espesor de losa monolítica; d: peralte efectivo inicial.",
        f"h = {geometry.height_m:.3f} + {geometry.slab_thickness_m:.3f} = "
        f"{geometry.total_depth_m:.3f} m; d = {reinforcement.positive.effective_depth_cm:.2f} cm.",
        f"Se adopta una sección resistente rectangular de b={geometry.thickness_m:.3f} m y "
        f"h={geometry.total_depth_m:.3f} m.",
        "El peso propio adicional del diafragma utiliza solo hbajo; la losa se contabiliza separadamente.",
        REF_DEAD_LOAD,
    )
    for row, title in ((positive, "Momento positivo crítico"), (negative, "Momento negativo crítico")):
        _moment_combination_calc(document, title, row, include_pl=True)
    _shear_combination_calc(document, "Cortante crítico", shear, include_pl=True)
    document.add_heading("7.2 Flexión y cortante", level=2)
    for design, prefix in ((reinforcement.positive, "B."), (reinforcement.negative, "A.")):
        _flexural_calc(document, design.label, design, selected_option(data.diaphragm_selected, prefix), data.project_inputs.materials.concrete, data.project_inputs.materials.steel, strip=False)
    temp = selected_option(data.diaphragm_selected, "C.")
    stirrup = selected_option(data.diaphragm_selected, "D.")
    _calc(
        document,
        "Acero lateral por temperatura",
        "As,temp,prov ≥ As,temp,req",
        "As,temp: refuerzo por metro y por cara lateral.",
        f"As,temp,req = {reinforcement.temperature.required_area_cm2_m_per_face:.3f} cm²/m por cara; se adopta {_option_text(temp)}.",
        "La disposición se repite en ambas caras laterales.",
        _compliance_comment(temp, "La cuantía adoptada controla la fisuración por cambios volumétricos."),
        REF_TEMP,
    )
    skin_option = reinforcement.skin.spacing_options.recommended
    if reinforcement.skin.required_area_cm2_m_per_face > 0.0:
        _calc(
            document,
            "Acero longitudinal superficial Ask del diafragma",
            "Ask = min(0.012·(d_l − 30 in), As/flexión/4); s ≤ min(d_l/6, 300 mm)",
            "Ask se distribuye por cada cara lateral dentro de d_l/2 desde la cara traccionada. "
            "La exigencia aplica cuando d_l > 900 mm.",
            f"d_l = {reinforcement.skin.effective_depth_cm:.2f} cm; "
            f"d_l/2 = {reinforcement.skin.distribution_height_m:.3f} m; "
            f"Ask req = {reinforcement.skin.required_area_cm2_m_per_face:.3f} cm²/m por cara; "
            f"s máx = {reinforcement.skin.maximum_spacing_m:.3f} m. "
            f"Se adopta {_option_text(skin_option)}.",
            "El acero principal ubicado en las caras laterales solo puede contarse como Ask si se verifica explícitamente su área, ubicación y compatibilidad.",
            _compliance_comment(skin_option, "La disposición de Ask cumple el área por cara y la separación máxima."),
            REF_CRACK,
        )
    else:
        _comment(
            document,
            f"Acero longitudinal superficial Ask: NO APLICA porque d_l="
            f"{reinforcement.skin.effective_depth_cm:.2f} cm ≤ 90.00 cm. "
            "El acero lateral por temperatura se verifica independientemente.",
        )
    _calc(
        document,
        "Estribos del diafragma",
        *_sectional_shear_calc_fields(reinforcement.shear, stirrup),
        REF_SHEAR,
    )
    document.add_heading("7.3 Diagramas de diseño", level=2)
    moment_path = save_strength_chart(result, chart_dir / "diafragma_momentos.png", title="Diafragma · envolventes de momento Resistencia I", kind="moment", include_pl=True, transverse=True)
    shear_path = save_strength_chart(result, chart_dir / "diafragma_cortantes.png", title="Diafragma · envolventes de cortante Resistencia I", kind="shear", include_pl=True, transverse=True)
    _picture(document, moment_path, "Figura 7.1. Envolventes factorizadas de momento del diafragma.")
    _picture(document, shear_path, "Figura 7.2. Envolventes factorizadas de cortante del diafragma.")


def _reactions(document: Document, data: DeckReportData) -> None:
    from bridge_design.domain.global_reactions import project_live_reaction_cases
    live_cases = project_live_reaction_cases(data.project_inputs)
    _table(document, ("Caso global HL-93", "R fijo Tn", "R movil Tn", "Carga total Tn"),
           tuple((row.name, f"{row.left_tn:.3f}", f"{row.right_tn:.3f}", f"{row.total_load_tn:.3f}") for row in live_cases),
           widths=(77, 30, 30, 30))
    _comment(
        document,
        "LL+IM para estribos procede de estos casos globales por equilibrio. "
        "Las tablas longitudinales anteriores conservan las reacciones asociadas a Mmax; "
        "la tabla de apoyos siguiente usa envolventes locales independientes de reacción máxima. "
        "Se incluye tráfico nulo para efectos estabilizadores.",
    )
    document.add_heading("8. Reacciones por viga para apoyos y cargas globales para estribos", level=1)
    _body(
        document,
        "La primera tabla presenta reacciones individuales por viga para el diseño de apoyos. "
        "En LL+IM se maximiza independientemente la reacción de cada extremo mediante su línea de influencia, "
        "incluyendo IM y el factor de distribución de cortante gV. "
        "La tabla 8.1 presenta las acciones globales que deben transferirse al estribo. "
        "Todas las magnitudes se muestran sin factor para que cada módulo receptor forme la combinación requerida."
    )
    rows = []
    for label, result, include_pl in (
        ("Interior", data.interior_result, False),
        ("Exterior", data.exterior_result, True),
    ):
        cases = [("DC", result.dc), ("DW", result.dw)]
        if include_pl:
            cases.append(("PL", result.pl))
        cases.append(("LL+IM", result.ll_im_envelope))
        for action, case in cases:
            if action == "LL+IM":
                left, right = _maximum_support_reaction_pair(case)
            else:
                left, right = _two_reactions(case.support_reactions_tn)
            rows.append((label, action, f"{left:.3f}", f"{right:.3f}"))
    _table(
        document,
        ("Viga", "Acción", "Máx. apoyo izquierdo (Tn)", "Máx. apoyo derecho (Tn)"),
        tuple(rows),
        widths=(38, 35, 47, 47),
    )
    _comment(
        document,
        "Los máximos LL+IM izquierdo y derecho ocurren en posiciones vehiculares reflejadas y no son simultáneos. "
        "En un puente simétrico sus magnitudes deben coincidir.",
    )
    document.add_heading("8.1 Cargas globales para el diseño de estribos", level=2)
    _body(
        document,
        "Usar esta tabla para el estribo. DC, DW y PL se obtienen sumando las vigas del tablero. "
        "LL+IM se obtiene por equilibrio global HL-93 y se reporta como envolvente independiente por apoyo. "
        "Los máximos izquierdo y derecho no representan necesariamente un único posicionamiento simultáneo del vehículo."
    )
    global_rows, strength_row = _global_abutment_load_rows(data, live_cases)
    _table(
        document,
        ("Acción", "R fijo (Tn)", "R móvil (Tn)", "q fijo (Tn/m)", "q móvil (Tn/m)"),
        tuple(global_rows),
        widths=(30, 31, 31, 33, 33),
    )
    _table(
        document,
        ("Combinación", "R fijo (Tn)", "R móvil (Tn)", "Criterio"),
        (strength_row,),
        widths=(35, 32, 32, 59),
    )
    int_dc_l, _int_dc_r = _two_reactions(data.interior_result.dc.support_reactions_tn)
    ext_dc_l, _ext_dc_r = _two_reactions(data.exterior_result.dc.support_reactions_tn)
    _calc(
        document,
        "Uso de reacciones en el receptor de estribos",
        "Rtotal,acción = nint·Rviga,int + next·Rviga,ext  ;  q = Rtotal / Ltablero",
        "nint, next: número de vigas interior y exterior; Rviga: reacción no factorizada por apoyo; "
        "q: carga lineal para el modelo de estribo por franja.",
        f"Ejemplo DC en apoyo izquierdo: Rint = {int_dc_l:.3f} Tn/viga; Rext = {ext_dc_l:.3f} Tn/viga. "
        f"Las máximas reacciones locales de LL+IM ya incluyen gV y el IM; no se vuelve a aplicar ninguno al transferirlas.",
        "El estribo combina estas acciones con sus factores LRFD propios; aquí sólo se entregan valores sin factor.",
        "No deben sumarse posiciones vehiculares incompatibles al formar otra envolvente fuera de este módulo.",
        REF_DEAD_LOAD,
    )
    _comment(document, "Las reacciones de LL+IM son envolventes de carga móvil. No deben sumarse posiciones vehiculares incompatibles al usarlas en otra combinación.")


def _global_abutment_load_rows(data: DeckReportData, live_cases) -> tuple[list[tuple[str, ...]], tuple[str, ...]]:
    """Return global unfactored support loads and the Strength I envelope for abutments."""
    width = data.project_inputs.transverse_slab.geometry.total_width_m
    interior_count = max(data.project_inputs.interior_girder.girder_count - 2, 0)

    def total(interior_case, exterior_case):
        int_left, int_right = _two_reactions(interior_case.support_reactions_tn)
        ext_left, ext_right = _two_reactions(exterior_case.support_reactions_tn)
        return interior_count * int_left + 2 * ext_left, interior_count * int_right + 2 * ext_right

    dc = total(data.interior_result.dc, data.exterior_result.dc)
    dw = total(data.interior_result.dw, data.exterior_result.dw)
    ext_pl = _two_reactions(data.exterior_result.pl.support_reactions_tn)
    pl = (2 * ext_pl[0], 2 * ext_pl[1])
    ll = (max(row.left_tn for row in live_cases), max(row.right_tn for row in live_cases))
    service = tuple(sum(values) for values in zip(dc, dw, pl, ll))
    strength = tuple(1.25 * dc[i] + 1.50 * dw[i] + 1.75 * pl[i] + 1.75 * ll[i] for i in (0, 1))
    components = (("DC", dc), ("DW", dw), ("PL", pl), ("LL+IM", ll), ("Servicio I", service))
    rows = [(label, f"{values[0]:.3f}", f"{values[1]:.3f}", f"{values[0] / width:.3f}", f"{values[1] / width:.3f}") for label, values in components]
    strength_row = ("Resistencia I", f"{strength[0]:.3f}", f"{strength[1]:.3f}", "1.25DC + 1.50DW + 1.75PL + 1.75(LL+IM)")
    return rows, strength_row


def _conclusions(document: Document, data: DeckReportData) -> None:
    document.add_heading("9. Resumen de diseño adoptado", level=1)
    rows = []
    for label, option in data.slab_selected:
        status = _option_status(option)
        if label.startswith("A."):
            status = data.slab_crack.negative_main.status
        elif label.startswith("B."):
            status = data.slab_crack.positive_main.status
        rows.append(("Losa transversal", label.split(".", 1)[-1].strip(), _option_text(option), status))
    for label, option in data.interior_selected:
        rows.append(("Viga interior", label.split(".", 1)[-1].strip(), _option_text(option), _option_status(option)))
    for label, option in data.exterior_selected:
        rows.append(("Viga exterior", label.split(".", 1)[-1].strip(), _option_text(option), _option_status(option)))
    barrier_inputs = data.project_inputs.barrier.section_model
    rows.append(
        (
            "Barrera",
            "Dowel de interfaz y anclaje",
            f"1/2 in @ {barrier_inputs.dowel_spacing_m:.3f} m, Avf = {barrier_inputs.dowel_area_cm2_m:.3f} cm²/m",
            data.barrier_result.development.status,
        )
    )
    for label, option in data.cantilever_selected:
        rows.append(("Voladizo", label.split(".", 1)[-1].strip(), _option_text(option), _option_status(option)))
    collision = data.cantilever_result.barrier_collision
    interior_collision = data.cantilever_result.interior_collision
    collision_status = (
        collision.status
        if collision is not None
        else "NO APLICA - BARRERA EN LOSA INTERIOR"
        if interior_collision is not None
        else "PENDIENTE"
    )
    rows.append(("Voladizo", "Conexion y todos los casos de colision",
                 "Casos 1/2/3 + yield-line/dowel/desarrollo",
                 collision_status))
    if interior_collision is not None:
        rows.append(("Losa interior", "Transferencia local de colision",
                     "Horizontal + vertical; conexion por capacidad", interior_collision.status))
        for face in interior_collision.faces:
            rows.append(("Losa interior", f"Minimo por colision, cara {face.face}",
                         _option_text(face.option), face.status))
        _comment(document,
            "Las alternativas de colision son minimos calculados, no una sustitucion automatica del acero "
            "elegido para la losa. El armado definitivo debe cubrir ambas demandas por cara. "
            "El estado de conexion por capacidad indicado aqui prevalece sobre la comprobacion "
            "de anclaje reducido de la barrera aislada.")
    for label, option in data.diaphragm_selected:
        rows.append(("Diafragma", label.split(".", 1)[-1].strip(), _option_text(option), _option_status(option)))
    _table(document, ("Elemento", "Función", "Refuerzo adoptado", "Estado"), tuple(rows), widths=(32, 62, 48, 25), font_size=7.6)
    _comment(
        document,
        "Las opciones listadas son las disposiciones adoptadas y verificadas por resistencia, cuantía mínima, separación, servicio y detalle según corresponda. Si una verificación cambiara a NO CUMPLE, deberá incrementarse el área provista, reducirse el espaciamiento o modificarse la geometría antes de emitir planos."
    )


def _references(document: Document) -> None:
    document.add_heading("10. Referencias normativas", level=1)
    for text in (
        "Ministerio de Transportes y Comunicaciones. Manual de Puentes, Lima, 2018. Archivo: docs/Manual de Puentes MTC 2018 (PGA).pdf.",
        "Rodríguez Serquén, Arturo. Puentes con AASHTO LRFD 2020, 9th Edition. Archivo de consulta en la carpeta docs del proyecto.",
        "AASHTO LRFD Bridge Design Specifications, criterios incorporados por las referencias y numeración equivalentes indicadas en el Manual MTC 2018.",
    ):
        p = document.add_paragraph(style="List Bullet")
        p.add_run(text)


def _flexural_calc(document, title, design, option, concrete, steel, *, strip: bool) -> None:
    provided_area = None
    if option is not None:
        provided_area = getattr(
            option,
            "provided_area_cm2_m",
            getattr(option, "provided_area_cm2", None),
        )
    formula, legend, substitution, result, comment = detail.flexural_as_min_trace(
        strip=strip,
        design=design,
        concrete=concrete,
        steel=steel,
        option_text=_option_text(option),
        provided_area=provided_area,
        compliance_comment=_compliance_comment(
            option,
            "La resistencia provista es mayor o igual que la demanda y se respeta el mínimo reglamentario.",
        ),
    )
    _calc(document, title, formula, legend, substitution, result, comment, REF_FLEX)


def _crack_calc(document, crack, *, slab_thickness_m: float | None = None) -> None:
    thickness_cm = None if slab_thickness_m is None else slab_thickness_m * 100.0
    formula, legend, substitution, result, comment = detail.crack_control_trace(
        crack,
        slab_thickness_cm=thickness_cm,
    )
    _calc(
        document,
        f"Control de fisuración — {getattr(crack, 'direction', 'acero adoptado')}",
        formula,
        legend,
        substitution,
        result,
        _status_comment(getattr(crack, "status", "CUMPLE"), comment),
        REF_CRACK,
    )


def _moment_combination_calc(document, title, row, *, include_pl: bool) -> None:
    terms = [f"{row.dc_factor:.2f}·({row.dc_moment_tn_m:.3f})", f"{row.dw_factor:.2f}·({row.dw_moment_tn_m:.3f})"]
    legend = "DC, DW: efectos permanentes"
    if include_pl and hasattr(row, "pl_moment_tn_m"):
        terms.append(f"{row.pl_factor:.2f}·({row.pl_moment_tn_m:.3f})")
        legend += "; PL: efecto peatonal"
    terms.append(f"{row.ll_im_factor:.2f}·({row.ll_im_moment_tn_m:.3f})")
    _calc(
        document,
        title,
        "Mu(x) = γDC·MDC(x) + γDW·MDW(x) + γPL·MPL(x) + γLL·M(LL+IM)(x)",
        legend + "; LL+IM: efecto vehicular; γ: factores de la combinación.",
        "Mu = " + " + ".join(terms) + f" = {row.combined_moment_tn_m:.3f} Tn·m.",
        f"Estación crítica x = {row.position_m:.3f} m; combinación {row.combination_name}; dirección {getattr(row, 'direction', 'M+')}.",
        "La estación se obtiene de la envolvente combinada y no de la suma aislada de máximos de cada caso.",
        REF_COMB,
    )


def _shear_combination_calc(document, title, row, *, include_pl: bool) -> None:
    terms = [f"{row.dc_factor:.2f}·{row.dc_shear_tn:.3f}", f"{row.dw_factor:.2f}·{row.dw_shear_tn:.3f}"]
    if include_pl and hasattr(row, "pl_shear_tn"):
        terms.append(f"{row.pl_factor:.2f}·{row.pl_shear_tn:.3f}")
    terms.append(f"{row.ll_im_factor:.2f}·{row.ll_im_shear_tn:.3f}")
    _calc(
        document,
        title,
        "Vu(x) = γDC·VDC(x) + γDW·VDW(x) + γPL·VPL(x) + γLL·V(LL+IM)(x)",
        "V: cortante en la estación; γ: factores LRFD; para diseño se compara la magnitud absoluta.",
        "Vu = " + " + ".join(terms) + f" = {row.combined_shear_tn:.3f} Tn.",
        f"Estación crítica x = {row.position_m:.3f} m; combinación {row.combination_name}.",
        "La sección crítica se evalúa cerca del apoyo y el cortante móvil se toma de su envolvente superior o inferior, según gobierne.",
        REF_COMB,
    )


def _calc(document, title, formula, legend, substitution, result, comment, reference) -> None:
    heading = document.add_heading(title, level=3)
    heading.paragraph_format.keep_with_next = True

    intro = document.add_paragraph(style="Normal")
    intro.paragraph_format.keep_with_next = True
    intro.add_run(
        "Para desarrollar esta verificación se emplea la siguiente expresión:"
    )
    _add_native_equation(document, formula)

    where = document.add_paragraph(style="Normal")
    where.paragraph_format.space_before = Pt(1)
    where.paragraph_format.space_after = Pt(1)
    where.paragraph_format.keep_with_next = True
    run = where.add_run("Donde:")
    _format_run(run, 9.2, bold=True, color=NAVY)
    for definition in _legend_items(legend):
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.keep_with_next = False
        variable, separator, description = definition.partition(":")
        if separator:
            run = paragraph.add_run(variable.strip() + " = ")
            _format_run(run, 9.2, bold=True, color=NAVY)
            run = paragraph.add_run(description.strip())
            _format_run(run, 9.2, color=NAVY)
        else:
            run = paragraph.add_run(definition.strip())
            _format_run(run, 9.2, color=NAVY)

    replacing = document.add_paragraph(style="Normal")
    replacing.paragraph_format.space_before = Pt(3)
    replacing.paragraph_format.space_after = Pt(1)
    replacing.paragraph_format.keep_with_next = True
    run = replacing.add_run("Reemplazando los valores correspondientes:")
    _format_run(run, 9.2, color=NAVY)
    _add_substitution(document, substitution)

    conclusion = document.add_paragraph(style="Normal")
    conclusion.paragraph_format.space_before = Pt(2)
    conclusion.paragraph_format.space_after = Pt(3)
    run = conclusion.add_run("Por lo tanto, ")
    _format_run(run, 9.2, color=NAVY)
    run = conclusion.add_run(result)
    _format_run(run, 9.2, bold=True, color=NAVY)

    technical = document.add_paragraph(style="Normal")
    technical.paragraph_format.space_before = Pt(1)
    technical.paragraph_format.space_after = Pt(2)
    technical_color = GREEN if _looks_compliant(comment) else RED
    run = technical.add_run(comment)
    _format_run(run, 9.0, color=technical_color)

    citation = document.add_paragraph(style="Normal")
    citation.paragraph_format.space_before = Pt(0)
    citation.paragraph_format.space_after = Pt(7)
    run = citation.add_run("Referencia normativa: ")
    _format_run(run, 8.6, bold=True, color=GRAY)
    run = citation.add_run(reference)
    _format_run(run, 8.6, italic=True, color=GRAY)


def _legend_items(legend: str) -> tuple[str, ...]:
    """Split definitions without breaking semicolons inside equations."""
    items: list[str] = []
    remaining = legend.strip()
    while remaining:
        split = _split_top_level(remaining, (";",))
        if split is None:
            item = remaining.strip().rstrip(".")
            if item:
                items.append(item)
            break
        left, _, remaining = split
        item = left.strip().rstrip(".")
        if item:
            items.append(item)
        remaining = remaining.strip()
    return tuple(items)


def _add_substitution(document: Document, substitution: str) -> None:
    """Render substitution steps vertically: one assignment or note per line."""
    for line in _substitution_lines(substitution):
        if _is_equation_text(line):
            _add_native_equation(document, line)
            continue
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.left_indent = Mm(7)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(line)
        _format_run(run, 9.2, color=NAVY)


def _substitution_lines(substitution: str) -> tuple[str, ...]:
    """Split a substitution block into vertical development lines."""
    text = substitution.replace("\r\n", "\n").strip()
    if not text:
        return ()
    if "\n" in text:
        return tuple(line.strip().rstrip(".") for line in text.split("\n") if line.strip())

    lines: list[str] = []
    for sentence in _split_substitution_sentences(text):
        for piece in _equation_lines(sentence):
            cleaned = piece.strip().rstrip(".")
            if cleaned:
                lines.append(cleaned)
    return tuple(lines)


def _split_substitution_sentences(text: str) -> tuple[str, ...]:
    """Split on '. ' sentence boundaries (decimals use '.' without a following space)."""
    parts: list[str] = []
    remaining = text.strip()
    while remaining:
        index = remaining.find(". ")
        if index < 0:
            parts.append(remaining)
            break
        left = remaining[:index].strip()
        if left:
            parts.append(left)
        remaining = remaining[index + 2 :].strip()
    return tuple(parts)


def _is_equation_text(value: str) -> bool:
    return any(symbol in value for symbol in ("=", "≤", "≥", "<", ">"))


def _add_native_equation(document: Document, expression: str):
    paragraphs = []
    for line in _equation_lines(expression):
        paragraph = document.add_paragraph(style="Equation")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.keep_with_next = True
        math = OxmlElement("m:oMath")
        _append_math_expression(math, line)
        paragraph._p.append(math)
        paragraphs.append(paragraph)
    return paragraphs[-1] if paragraphs else None


def _equation_lines(expression: str) -> tuple[str, ...]:
    """Split independent equations while preserving delimiters inside groups."""
    remaining = expression.strip()
    lines: list[str] = []
    while remaining:
        split = _split_top_level(remaining, (";",))
        if split is None:
            lines.append(remaining)
            break
        left, _, remaining = split
        if left.strip():
            lines.append(left.strip())
        remaining = remaining.strip()
    return tuple(lines)


def _append_math_expression(parent, expression: str) -> None:
    expression = expression.strip()
    if not expression:
        return
    if _has_outer_group(expression, "(", ")"):
        parent.append(_math_run("("))
        _append_math_expression(parent, expression[1:-1])
        parent.append(_math_run(")"))
        return
    for operators in (("≤", "≥", "="), ("+", "−")):
        split = _split_top_level(expression, operators)
        if split is not None:
            left, operator, right = split
            _append_math_expression(parent, left)
            parent.append(_math_run(f" {operator} "))
            _append_math_expression(parent, right)
            return
    split = _split_top_level(expression, ("/",))
    if split is not None:
        numerator, _, denominator = split
        fraction = OxmlElement("m:f")
        fraction_properties = OxmlElement("m:fPr")
        fraction_type = OxmlElement("m:type")
        fraction_type.set(qn("m:val"), "bar")
        fraction_properties.append(fraction_type)
        fraction.append(fraction_properties)
        num = OxmlElement("m:num")
        den = OxmlElement("m:den")
        _append_math_expression(num, numerator)
        _append_math_expression(den, denominator)
        fraction.extend((num, den))
        parent.append(fraction)
        return
    if expression.startswith("√"):
        radicand = expression[1:].strip()
        if _has_outer_group(radicand, "[", "]") or _has_outer_group(radicand, "(", ")"):
            radicand = radicand[1:-1]
        radical = OxmlElement("m:rad")
        properties = OxmlElement("m:radPr")
        degree_hide = OxmlElement("m:degHide")
        degree_hide.set(qn("m:val"), "1")
        properties.append(degree_hide)
        radical.append(properties)
        radical.append(OxmlElement("m:deg"))
        body = OxmlElement("m:e")
        _append_math_expression(body, radicand)
        radical.append(body)
        parent.append(radical)
        return
    caret = _find_top_level_operator(expression, "^")
    if caret is not None:
        base_start = _superscript_base_start(expression, caret)
        if base_start > 0:
            _append_math_expression(parent, expression[:base_start])
        superscript = OxmlElement("m:sSup")
        base = OxmlElement("m:e")
        exponent = OxmlElement("m:sup")
        _append_math_expression(base, expression[base_start:caret])
        _append_math_expression(exponent, expression[caret + 1 :])
        superscript.extend((base, exponent))
        parent.append(superscript)
        return
    parent.append(_math_run(expression))


def _split_top_level(expression: str, operators: tuple[str, ...]):
    depths = {"(": 0, "[": 0, "{": 0}
    closers = {")": "(", "]": "[", "}": "{"}
    for index, character in enumerate(expression):
        if character in depths:
            depths[character] += 1
            continue
        if character in closers:
            opener = closers[character]
            depths[opener] = max(depths[opener] - 1, 0)
            continue
        if any(depths.values()):
            continue
        for operator in operators:
            if expression.startswith(operator, index):
                if operator in {"+", "−"} and index == 0:
                    continue
                if operator == "/" and not _slash_is_fraction(expression, index):
                    continue
                left = expression[:index].strip()
                right = expression[index + len(operator) :].strip()
                if left and right:
                    return left, operator, right
    return None


def _slash_is_fraction(expression: str, index: int) -> bool:
    left = expression[:index].rstrip()
    right = expression[index + 1 :].lstrip()
    unit_numerators = ("Tn", "Tn·m", "kg", "cm", "cm²", "m", "m²")
    unit_denominators = ("m", "m²", "cm", "cm²", "s")
    if any(left.endswith(unit) for unit in unit_numerators) and right.startswith(unit_denominators):
        return False
    if right and right[0].isdigit():
        position = 0
        while position < len(right) and (right[position].isdigit() or right[position] in ".,"):
            position += 1
        if position < len(right) and right[position] in {'"', "'"}:
            return False
    return True


def _find_top_level_operator(expression: str, operator: str) -> int | None:
    split = _split_top_level(expression, (operator,))
    if split is None:
        return None
    left, _, _ = split
    return expression.find(operator, len(left))


def _has_outer_group(expression: str, opener: str, closer: str) -> bool:
    if not expression.startswith(opener) or not expression.endswith(closer):
        return False
    depth = 0
    for index, character in enumerate(expression):
        if character == opener:
            depth += 1
        elif character == closer:
            depth -= 1
            if depth == 0 and index != len(expression) - 1:
                return False
    return depth == 0


def _superscript_base_start(expression: str, caret: int) -> int:
    end = caret - 1
    if end >= 0 and expression[end] in ")]":
        opener = "(" if expression[end] == ")" else "["
        closer = expression[end]
        depth = 0
        for index in range(end, -1, -1):
            if expression[index] == closer:
                depth += 1
            elif expression[index] == opener:
                depth -= 1
                if depth == 0:
                    return index
    index = end
    while index >= 0 and expression[index] not in " =+−·;/([":
        index -= 1
    return index + 1


def _math_run(text_value: str):
    run = OxmlElement("m:r")
    properties = OxmlElement("m:rPr")
    style = OxmlElement("m:sty")
    style.set(qn("m:val"), "p")
    properties.append(style)
    run.append(properties)
    word_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(UNIFORM_FONT_SIZE_PT * 2))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEXT_COLOR)
    word_properties.extend((fonts, size, color))
    run.append(word_properties)
    text = OxmlElement("m:t")
    text.text = text_value
    run.append(text)
    return run


def _table(document, headers, rows, *, widths, font_size=8.3, accent=False) -> None:
    total_width = sum(float(width) for width in widths)
    if total_width > MAX_CONTENT_WIDTH_MM:
        scale = MAX_CONTENT_WIDTH_MM / total_width
        widths = tuple(float(width) * scale for width in widths)
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER if accent else WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"
    for index, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = Mm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cell, "D9D9D9")
        _set_cell_margins(cell, 70, 80, 70, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.keep_with_next = True
        run = p.add_run(str(header))
        _format_run(run, 11, bold=True, color=TEXT_COLOR)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value, width) in enumerate(zip(cells, row, widths)):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if accent and row_index % 2:
                _set_cell_shading(cell, LIGHT)
            _set_cell_margins(cell, 55, 75, 55, 75)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run(str(value))
            _format_run(run, 11, color=TEXT_COLOR)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for table_row in table.rows:
        row_properties = table_row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
    _set_table_geometry(table, widths, indent_twips=0 if not accent else 120)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _set_table_geometry(table, widths, *, indent_twips: int) -> None:
    width_twips = tuple(round(float(width) * 56.692913) for width in widths)
    total = sum(width_twips)
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(total))
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(indent_twips))
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in width_twips:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, width_twips):
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.set(qn("w:type"), "dxa")
            tc_width.set(qn("w:w"), str(width))


def _picture(document, path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Mm(MAX_CONTENT_WIDTH_MM))
    p = document.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(caption)


def _body(document, text_value: str) -> None:
    p = document.add_paragraph(text_value)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _comment(document, text_value: str) -> None:
    p = document.add_paragraph(style="Formula")
    _shade_paragraph(p, LIGHT)
    _left_border(p, TEAL, 12)
    run = p.add_run(text_value)
    _format_run(run, 8.8, color=NAVY)


def _moment_row_values(row, *, include_pl: bool):
    values = [getattr(row, "direction", "—"), row.combination_name, f"{row.position_m:.3f}", f"{row.dc_moment_tn_m:.3f}", f"{row.dw_moment_tn_m:.3f}"]
    values.append(f"{getattr(row, 'pl_moment_tn_m', 0.0):.3f}" if include_pl else "—")
    values.extend((f"{row.ll_im_moment_tn_m:.3f}", f"{row.combined_moment_tn_m:.3f}"))
    return tuple(values)


def _sectional_shear_calc_fields(shear, stirrup) -> tuple[str, str, str, str, str]:
    vu = shear.controlling_shear.combined_shear_tn
    vn_max = shear.nominal_shear_limit_tn
    phi_vn_max = shear.phi * vn_max
    phi_vc = shear.phi * shear.vc_tn
    limit_ok = phi_vn_max + 1e-9 >= vu
    limit_status = "CUMPLE" if limit_ok else "NO CUMPLE"
    comment = (
        "La sección crítica se toma a una distancia dv desde el apoyo. "
        "Vu no supera φVn,max, por lo que el alma puede resistir el cortante con la armadura transversal adoptada."
        if limit_ok
        else (
            "Vu supera φVn,max. Aumentar estribos no incrementa la resistencia de la sección; "
            "debe revisarse el ancho de alma, el peralte o f'c."
        )
    )
    if not limit_ok:
        option_comment = comment
    else:
        option_comment = _compliance_comment(stirrup, comment)
    return (
        "Vu ≤ φVn  ;  Vn = min(Vc + Vs; Vn,max)  ;  Vn,max = 0.25 f'c bv dv  ;  Vs,req = max[0; Vu/φ − Vc]",
        (
            "Vu: cortante factorizado; Vn: resistencia nominal; Vc: aporte del concreto; "
            "Vs: aporte de estribos; Vn,max: límite superior del modelo seccional; φ: factor de resistencia."
        ),
        (
            f"Vu = {vu:.3f} Tn\n"
            f"Vc = {shear.vc_tn:.3f} Tn\n"
            f"φ = {shear.phi:.3f}\n"
            f"φVc = {phi_vc:.3f} Tn\n"
            f"Vs,req = {shear.required_vs_tn:.3f} Tn\n"
            f"Vn,max = 0.25 f'c bv dv = {vn_max:.3f} Tn\n"
            f"φVn,max = {shear.phi:.3f}·{vn_max:.3f} = {phi_vn_max:.3f} Tn"
        ),
        (
            f"Vu = {vu:.3f} Tn ≤ φVn,max = {phi_vn_max:.3f} Tn: {limit_status}. "
            f"Av,req = {shear.required_av_cm2_m:.3f} cm²/m; Av,min = {shear.minimum_av_cm2_m:.3f} cm²/m; "
            f"smax = {shear.maximum_spacing_m:.3f} m; se adopta {_option_text(stirrup)}."
        ),
        option_comment,
    )


def _option_text(option) -> str:
    if option is None:
        return "opción de cálculo no disponible"
    suffix = " (selección personalizada del usuario)" if getattr(option, "is_custom", False) else ""
    if hasattr(option, "bar_count"):
        layer_word = "capa" if option.layers == 1 else "capas"
        return f"{option.bar_count} barras de {option.bar_label} distribuidas en {option.layers} {layer_word}, As = {option.provided_area_cm2:.3f} cm²{suffix}"
    if hasattr(option, "legs"):
        return f"estribo {option.bar_label}, {option.legs} ramas @ {option.spacing_m:.3f} m, Av = {option.provided_av_cm2_m:.3f} cm²/m{suffix}"
    bar = getattr(option, "bar", None)
    label = getattr(bar, "label", getattr(option, "bar_label", "barra"))
    return f"{label} @ {option.spacing_m:.3f} m, As = {option.provided_area_cm2_m:.3f} cm²/m{suffix}"


def _option_status(option) -> str:
    return "CUMPLE" if option is not None and bool(getattr(option, "is_compliant", False)) else "REVISAR"


def _compliance_comment(option, ok_text: str) -> str:
    if option is not None and bool(getattr(option, "is_compliant", False)):
        return ok_text
    return "La opción no satisface la verificación; se debe aumentar el área, reducir el espaciamiento o revisar la geometría antes de emitir el diseño."


def _status_comment(status: str, ok_text: str) -> str:
    normalized = str(status).upper()
    if normalized in {"OK", "CUMPLE", "ADECUADO"}:
        return ok_text
    return "La verificación no cumple; debe incrementarse la capacidad o reducirse la demanda y repetirse el cálculo antes de emitir el diseño."


def _looks_compliant(text_value: str) -> bool:
    return "no cumple" not in text_value.lower() and "debe incrementarse" not in text_value.lower()


def _two_reactions(reactions) -> tuple[float, float]:
    values = [float(item[1]) for item in reactions]
    if not values:
        return 0.0, 0.0
    if len(values) == 1:
        return values[0], values[0]
    return values[0], values[-1]


def _maximum_support_reaction_pair(case) -> tuple[float, float]:
    reactions = case.maximum_support_reactions_tn or case.support_reactions_tn
    return _two_reactions(reactions)


def _field(paragraph, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Actualizar campo"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend((begin, instr, separate, text, end))
    _format_run(run, 8, color=GRAY)


def _format_run(run, size: float, *, bold: bool = False, italic: bool = False, color: str) -> None:
    run.font.name = FONT
    run.font.size = Pt(UNIFORM_FONT_SIZE_PT)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)


def _enforce_uniform_typography(document: Document) -> None:
    """Apply the user-requested typography to every textual paragraph."""
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    seen: set[int] = set()
    for paragraph in paragraphs:
        identity = id(paragraph._p)
        if identity in seen:
            continue
        seen.add(identity)
        is_cover_title = paragraph.style is not None and paragraph.style.name == "Title"
        paragraph.paragraph_format.line_spacing = 1.0 if is_cover_title else 1.5
        if paragraph.text.strip():
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if is_cover_title
                else WD_ALIGN_PARAGRAPH.JUSTIFY
            )
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(
                COVER_TITLE_FONT_SIZE_PT if is_cover_title else UNIFORM_FONT_SIZE_PT
            )
            run.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
            fonts = run._element.get_or_add_rPr().rFonts
            fonts.set(qn("w:ascii"), FONT)
            fonts.set(qn("w:hAnsi"), FONT)
            fonts.set(qn("w:eastAsia"), FONT)


def _shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _left_border(paragraph, color: str, size: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "5")
    left.set(qn("w:color"), color)
    borders.append(left)


def _bottom_border(paragraph, color: str, size: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int, start: int, bottom: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
