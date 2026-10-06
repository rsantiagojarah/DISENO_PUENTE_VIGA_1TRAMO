"""Memoria conjunta con el mismo desarrollo técnico de cada apoyo."""
from dataclasses import replace
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from bridge_design.reporting.deck_docx import _body,_configure_document,_table
from bridge_design.reporting.abutment_docx import _configure_abutment_header_footer,_labels
from bridge_design.reporting.bearing_a_labels import number,status,result_status
from bridge_design.reporting.bearing_a_docx import generate_bearing_a_docx,_trace


def generate_bearing_pair_docx(pair,output_path):
    path = Path(output_path).expanduser().resolve().with_suffix('.docx')
    doc = Document()
    _configure_document(doc)
    for name in ('Title','Subtitle'):
        for border in doc.styles[name]._element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    _configure_abutment_header_footer(doc,replace(_labels(is_pure_wall=False),
        header='MEMORIA DE CÁLCULO · SISTEMA DE APOYOS FIJO Y MÓVIL'))
    p = doc.add_paragraph('Memoria de cálculo\ndel sistema de apoyos fijo y móvil',style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Puente de un tramo Método A',style='Subtitle')
    _table(doc,('Dato','Descripción'),[
        ('Proyecto',pair.mobile.inputs.project),('Esquema','Un extremo fijo y otro móvil longitudinalmente'),
        ('Apoyo móvil',pair.mobile.inputs.bearing_id),('Apoyo fijo',pair.fixed.inputs.bearing_id),
        ('Estado del sistema',result_status(pair)),('Norma','Manual de Puentes MTC 2018, artículos 2.10.3 y 2.10.4'),
    ],widths=(52,107))
    _body(doc,'La memoria verifica la geometría, compresión, deformaciones, zunchos y estabilidad de los neoprenos fijo y móvil. El móvil acomoda los movimientos longitudinales por deformación del elastómero.')
    _body(doc,'El punto fijo ideal adopta desplazamiento relativo nulo, salvo el valor definido para su evaluación. Los movimientos del móvil se calculan desde ese punto. Las comprobaciones corresponden a los cuerpos de neopreno y sus zunchos.')
    _body(doc,'Las reacciones verticales se adoptan iguales para ambos extremos por hipótesis expresa.' if pair.inputs.equal_vertical_actions else 'Las reacciones verticales se definen por separado para cada extremo.')
    doc.add_heading('Comparación de los apoyos',1)
    _table(doc,('Magnitud','Móvil','Fijo'),[
        ('Reacción vertical de servicio',*(number(r.value('P'),'Tn') for r in (pair.mobile,pair.fixed))),
        ('Desplazamiento longitudinal relativo',*(number(r.value('DELTA'),'cm') for r in (pair.mobile,pair.fixed))),
        ('Planta y altura en mm',*(f'{r.adopted.length_cm*10:g} × {r.adopted.width_cm*10:g} × {r.value("HEIGHT")*10:g}' for r in (pair.mobile,pair.fixed))),
        ('Perforaciones pasantes',*(f'{r.adopted.hole_count} de Ø {r.adopted.hole_diameter_cm*10:g} mm' if r.adopted.hole_count else 'Sin perforaciones' for r in (pair.mobile,pair.fixed))),
        ('Estado del neopreno',result_status(pair.mobile),result_status(pair.fixed)),
    ],widths=(65,47,47))
    for letter,role,result in (('A','móvil',pair.mobile),('B','fijo',pair.fixed)):
        heading = doc.add_heading(f'{letter} Diseño del apoyo {role}',1)
        heading.paragraph_format.page_break_before = True
        _body(doc,f'Identificación: {result.inputs.bearing_id}. '+('El apoyo permite traslación longitudinal; su fricción se verifica en servicio.' if role=='móvil' else 'El desplazamiento del punto fijo ideal es nulo; las verificaciones corresponden a la geometría y las acciones adoptadas para el neopreno.'))
        generate_bearing_a_docx(result,path,document=doc,show_curves=(role=='móvil' or result.inputs.compression_curve != pair.mobile.inputs.compression_curve or (result.value('SI'),result.value('SE'),result.value('SIGMA')) != (pair.mobile.value('SI'),pair.mobile.value('SE'),pair.mobile.value('SIGMA'))))
    doc.save(path)
    return path
