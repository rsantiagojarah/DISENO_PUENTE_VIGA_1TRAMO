"""Componentes del estilo de memoria aprobado, independientes de cada cálculo.

La plantilla empaquetada contiene solo patrones tipográficos y de tablas.
No depende de documentos del escritorio ni de los scripts temporales de edición.
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont

from bridge_design.reporting.deck_docx import _append_math_expression, _math_run


def _text(element, text):
    """Conservar las propiedades del patrón y reemplazar su contenido completo."""
    properties = element.find('w:r/w:rPr', element.nsmap)
    properties = deepcopy(properties) if properties is not None else None
    for child in list(element):
        if child.tag != qn('w:pPr'):
            element.remove(child)
    paragraph = Paragraph(element, None)
    run = paragraph.add_run(str(text))
    if properties is not None:
        run._r.insert(0, properties)
    return paragraph


class BearingReportLayout:
    def __init__(self):
        self.doc = Document(Path(__file__).with_name('templates') / 'bearing_a_pair.docx')
        self.paragraphs = {p.text.split(':', 1)[1]: deepcopy(p._p)
                           for p in self.doc.paragraphs if p.text.startswith('@paragraph:')}
        self.tables = {len(t.columns): deepcopy(t._tbl) for t in self.doc.tables}
        for child in list(self.doc._element.body):
            if child.tag != qn('w:sectPr'):
                self.doc._element.body.remove(child)

    def append(self, element):
        self.doc._element.body.insert_element_before(element, 'w:sectPr')
        return element

    def p(self, text='', role='body'):
        node = deepcopy(self.paragraphs[role])
        _text(node, text)
        self.append(node)
        p = Paragraph(node, self.doc._body)
        if role == 'bullet':
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return p

    def link(self, paragraph, text, url):
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        run = paragraph.add_run(text)
        link.append(run._r)
        paragraph._p.append(link)

    def heading(self, text):
        return self.p(text, 'heading')

    def page_break(self):
        self.p('', 'blank').add_run().add_break(WD_BREAK.PAGE)

    def table(self, headers, rows, widths):
        table = deepcopy(self.tables[len(headers)])
        patterns = list(table.findall(qn('w:tr')))
        for row in patterns:
            table.remove(row)
        total = 8958  # ancho de tabla medido en la memoria aprobada, twips
        grid_widths = [round(total * w / sum(widths)) for w in widths]
        grid_widths[-1] += total - sum(grid_widths)
        grid = table.find(qn('w:tblGrid'))
        for child in list(grid):
            grid.remove(child)
        for width in grid_widths:
            col = OxmlElement('w:gridCol'); col.set(qn('w:w'), str(width)); grid.append(col)
        for index, values in enumerate([headers, *rows]):
            row = deepcopy(patterns[0] if index == 0 else patterns[1 + (index - 1) % (len(patterns) - 1)])
            for cell, value, width in zip(row.findall(qn('w:tc')), values, grid_widths):
                cell.find('w:tcPr/w:tcW', cell.nsmap).set(qn('w:w'), str(width))
                ps = cell.findall(qn('w:p'))
                _text(ps[0], value).alignment = WD_ALIGN_PARAGRAPH.LEFT
                for extra in ps[1:]:
                    cell.remove(extra)
            table.append(row)
        self.append(table)

    def equation(self, expression):
        p = self.p('', 'equation')
        wrapper = OxmlElement('m:oMathPara')
        math = OxmlElement('m:oMath')
        _append_math_expression(math, expression)
        wrapper.append(math); p._p.append(wrapper)
        # El analizador común resuelve fracciones y potencias; agregar subíndices
        # reales para los símbolos técnicos del apoyo, sin texto LaTeX visible.
        pattern = re.compile(r'([A-Za-zα-ωΑ-Ω]+)_([A-Za-z0-9]+)')
        for run in list(math.iter(qn('m:r'))):
            t = run.find(qn('m:t'))
            if t is None or not pattern.search(t.text or ''):
                continue
            text, previous, parts = t.text, 0, []
            for match in pattern.finditer(text):
                if match.start() > previous:
                    parts.append(_math_run(text[previous:match.start()]))
                sub = OxmlElement('m:sSub')
                base = OxmlElement('m:e'); base.append(_math_run(match[1]))
                suffix = OxmlElement('m:sub'); suffix.append(_math_run(match[2]))
                sub.extend((base, suffix)); parts.append(sub); previous = match.end()
            if previous < len(text):
                parts.append(_math_run(text[previous:]))
            parent, pos = run.getparent(), run.getparent().index(run)
            parent.remove(run)
            for offset, part in enumerate(parts):
                parent.insert(pos + offset, part)
        for run in math.iter(qn('m:r')):
            for pr in list(run.findall(qn('w:rPr'))):
                run.remove(pr)
            pr = OxmlElement('w:rPr')
            fonts = OxmlElement('w:rFonts')
            fonts.set(qn('w:ascii'), 'Cambria Math'); fonts.set(qn('w:hAnsi'), 'Cambria Math')
            size = OxmlElement('w:sz'); size.set(qn('w:val'), '20')
            pr.extend((fonts, size)); run.insert(0, pr)
        return p

    def diagram(self, result):
        self.p('', 'image').add_run().add_picture(bearing_diagram(result), width=Mm(158))


def bearing_diagram(result):
    """Sección y planta con la composición real; no supone coordenadas de agujeros."""
    g = result.adopted
    im = Image.new('RGB', (1800, 575), 'white')
    draw = ImageDraw.Draw(im)
    def font(size, bold=False):
        for name in (('arialbd.ttf' if bold else 'arial.ttf'), 'DejaVuSans.ttf'):
            try:
                return ImageFont.truetype(name, size)
            except OSError:
                pass
        return ImageFont.load_default(size=size)
    f, small, bold = font(28), font(23), font(29, True)
    draw.text((45, 15), 'SECCIÓN DEL NEOPRENO', font=bold, fill='black')
    left, right, top = 50, 1120, 145
    height = result.value('HEIGHT') * 10
    scale = 225 / height
    sequence = [(False, g.exterior_cm * 10)]
    for _ in range(g.interior_layers):
        sequence.extend(((True, g.steel_cm * 10), (False, g.interior_cm * 10)))
    sequence.extend(((True, g.steel_cm * 10), (False, g.exterior_cm * 10)))
    y = top
    for steel, thickness in sequence:
        end = y + thickness * scale
        draw.rectangle((left, y, right, end), fill='#414141' if steel else '#E6E9EC', outline='black', width=2)
        y = end
    if g.hole_count == 1:
        center = (left+right)/2
        half = (right-left)*g.hole_diameter_cm/g.length_cm/2
        draw.rectangle((center-half,top,center+half,y),fill='white')
        for edge in (center-half,center+half):
            for yy in range(int(top),int(y),14):
                draw.line((edge,yy,edge,min(yy+7,y)),fill='#555555',width=2)
    draw.line((left, 100, right, 100), fill='black', width=2)
    for x in (left, right):
        draw.line((x, 88, x, 112), fill='black', width=2)
    draw.text((490, 58), f'{g.length_cm*10:g} mm', font=f, fill='black')
    draw.line((1160, top, 1160, y), fill='black', width=2)
    for yy in (top, y):
        draw.line((1147, yy, 1173, yy), fill='black', width=2)
    draw.text((1180, 230), f'{height:g} mm', font=f, fill='black')
    draw.text((60, 410), f'{g.interior_layers} interiores de {g.interior_cm*10:g} mm   ·   2 exteriores de {g.exterior_cm*10:g} mm', font=f, fill='black')
    draw.text((60, 455), f'{g.interior_layers+1} zunchos de {g.steel_cm*10:g} mm   ·   Shore A {result.inputs.hardness}', font=f, fill='black')
    draw.text((60, 505), 'Esquema de capas   Sin escala', font=small, fill='#444444')
    px, py = 1370, 110
    scale_plan = 335 / max(g.length_cm, g.width_cm)
    width, length = g.width_cm * scale_plan, g.length_cm * scale_plan
    draw.text((1400, 15), 'PLANTA', font=bold, fill='black')
    draw.rectangle((px, py, px+width, py+length), fill='#E6E9EC', outline='black', width=2)
    # Una perforación se ilustra centrada, como en la memoria aprobada.
    # Es solo un esquema: el modelo no recibe coordenadas de fabricación.
    if g.hole_count:
        draw.text((px+10, py+35), f'{g.hole_count} de Ø{g.hole_diameter_cm*10:g} mm', font=small, fill='black')
        if g.hole_count == 1:
            cx,cy,radius = px+width/2,py+length/2,g.hole_diameter_cm*scale_plan/2
            draw.ellipse((cx-radius,cy-radius,cx+radius,cy+radius),fill='white',outline='black',width=2)
        draw.text((60, 540), 'Posición esquemática; ubicación según detalle de fabricación.', font=small, fill='#444444')
    draw.text((px, 470), f'{g.length_cm*10:g} × {g.width_cm*10:g} mm', font=f, fill='black')
    draw.text((px, 510), 'L longitudinal   W transversal', font=small, fill='black')
    stream = BytesIO(); im.save(stream, format='PNG'); stream.seek(0)
    return stream
