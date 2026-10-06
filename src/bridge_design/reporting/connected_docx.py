"""Calculation memorandum with narratives and only comparative tables."""

from pathlib import Path
from docx import Document
from bridge_design.domain.connected_2_inputs import is_connected_2
from bridge_design.reporting.connected_2_geometry import connected_2_geometry_rows

from bridge_design.reporting.connected_docx_layout import SECTIONS, connected_front_matter, connected_summary_and_references, report_reinforcement
from bridge_design.reporting.deck_docx import _body, _enforce_uniform_typography, _picture, _table
from bridge_design.reporting.connected_audit import build_connected_audit
from bridge_design.reporting.connected_case_groups import combination_groups, case_label
from bridge_design.reporting.connected_bearing_report import write_group_bearing
from bridge_design.reporting.connected_report_details import write_model_description
from bridge_design.reporting.connected_load_report import write_load_criteria


def write_connected_docx(result, path, charts):
    document = Document(Path(__file__).parent / "templates" / "connected_reference.docx")
    connected_front_matter(document, result)
    data, mesh = result.inputs, result.mesh
    uniform = is_connected_2(data)
    audit = build_connected_audit(result)

    def section(number):
        document.add_heading(f"{number}. {SECTIONS[number-1]}", level=1)

    section(1)
    reinforcement = report_reinforcement(result)
    _body(document, "Se documentan el análisis estático y seudoestático, el contacto con el terreno y el armado "
          + ("adoptado de dos estribos iguales con pantalla de espesor variable y cimentación continua uniforme. "
             if uniform else "adoptado de dos estribos con cajuela unidos mediante una cimentación continua zapata–losa–zapata. ") +
          "La memoria permite seguir las cargas desde su origen hasta las comprobaciones gobernantes.")
    _body(document, f"Se resolvieron {len(result.results)} combinaciones de carga. "
          "Se desarrollan las comprobaciones de flexión, cortante, fisuración, acero mínimo y longitudes de desarrollo.")
    _body(document, "Bases de referencia: Manual de Puentes MTC 2018. Se usan tn, m y radianes en el FRAME; "
          "kgf/cm² y cm en el concreto armado. Se analiza una franja transversal de 1,00 m "
          "con las reacciones del tablero por metro de estribo.")
    _body(document, "El análisis emplea secciones brutas elásticas de primer orden.")

    section(2)
    _picture(document, charts["geometria"], "Geometría del conjunto y ubicación de los resortes reales")
    geometry_note = (f"Separación libre {data.clear_span_m:.3f} m; longitud total {data.total_length_m:.3f} m; "
          f"losa de {data.slab_thickness_m:.3f} m; transiciones {data.left_transition_m:.3f} y "
          f"{data.right_transition_m:.3f} m. Origen en el extremo izquierdo y la cara superior de cimentación.")
    if uniform:
        geometry_note = (f"Luz libre superior {data.upper_clear_span_m:.3f} m; luz libre inferior {data.clear_span_m:.3f} m; "
                         f"ancho total {data.total_length_m:.3f} m; espesor uniforme {data.slab_thickness_m:.3f} m. "
                         "Origen en el extremo izquierdo y la cara superior de cimentación.")
    _body(document, geometry_note)
    if uniform:
        _body(document, "El conjunto está formado por dos estribos conectados. El trasdós es vertical; la cara interior "
              "se inclina entre la base y el asiento. El relleno exterior llega a la coronación del parapeto "
              "y no existe relleno interior sobre la zapata combinada. Los talones exteriores, los tramos bajo las pantallas "
              "y el tramo entre estribos forman una única zapata combinada de espesor uniforme. "
              "El armado superior e inferior se diseña con todas las secciones de esa zapata continua. "
              "No se define una puntera interior ni un bloque o transición bajo el asiento.")
    labels = (("retained_height_m", "Altura de relleno desde fondo de zapata"),
              ("footing_width_m", "Ancho de zapata"), ("footing_thickness_m", "Espesor de zapata"),
              ("toe_length_m", "Longitud de puntera"), ("heel_length_m", "Longitud de talón"),
              ("lower_stem_thickness_m", "Espesor inferior de pantalla"),
              ("upper_stem_thickness_m", "Espesor superior de pantalla"),
              ("bearing_seat_length_m", "Longitud horizontal de cajuela"),
              ("seat_wall_width_m", "Espesor del parapeto"), ("seat_block_height_m", "Altura del parapeto"),
              ("backwall_drop_m", "Altura del bloque de cajuela"),
              ("backwall_taper_height_m", "Altura de transición de cajuela"),
              ("small_batter_width_m", "Transición frontal t1"), ("backfill_step_width_m", "Retiro posterior t2"),
              ("front_soil_depth_m", "Altura de suelo frontal desde fondo de zapata"),
              ("bridge_seat_to_bearing_height_m", "Brazo adicional para frenado"))
    geometry_rows = [(label, f"{getattr(data.left.geometry, key):.3f}")
            for key, label in labels] + [
                ("Espesor de losa de fondo central que une las zapatas", f"{data.slab_thickness_m:.3f}"),
           ]
    _table(document, ("Dimensión en metros", "Ambos estribos"),
           connected_2_geometry_rows(data) if uniform else geometry_rows, widths=(110, 50), font_size=10)
    m = data.left.materials
    _table(document, ("Material", "Valor común"), (
        ("f'c kgf/cm²", f"{m.concrete_strength_kg_cm2:.1f}"),
        ("fy kgf/cm²", f"{m.steel_yield_kg_cm2:.1f}"),
        ("γ concreto tn/m³", f"{m.concrete_unit_weight_kg_m3/1000:.3f}"),
        ("γ relleno tn/m³", f"{m.soil_unit_weight_kg_m3/1000:.3f}"),
    ), widths=(110, 50), font_size=10)
    _body(document, f"Módulo de balasto vertical {data.soil.subgrade_tn_m3:.3f} tn/m³; "
          f"qadm={data.soil.allowable_tn_m2:.3f} tn/m²; FS nominal={data.soil.nominal_bearing_fs:.3f}. "
          f"Recubrimiento de pantalla izquierdo/derecho {data.left.reinforcement.stem_cover_cm:.2f}/"
          f"{data.right.reinforcement.stem_cover_cm:.2f} cm; zapata {data.left.reinforcement.footing_cover_cm:.2f}/"
          f"{data.right.reinforcement.footing_cover_cm:.2f} cm" +
          (" en toda la zapata combinada." if uniform else f"; losa {data.slab_cover_cm:.2f} cm."))

    section(3)
    write_model_description(document, result)

    section(4)
    write_initial_load_figures(document, result, charts)
    write_load_criteria(document, result)

    section(5)
    factors = result.combination_factors
    _table(document, ("Acción", *(f.name for f in factors)),
           ((label, *(f"{getattr(f, key):.2f}" for f in factors))
            for label, key in (("DC", "dc"), ("DW", "dw"), ("EV", "ev"), ("LL", "ll"),
                               ("LS vertical", "ls_vertical"), ("LS horizontal", "ls_horizontal"),
                               ("EH", "eh"), ("EQ", "eq"), ("BR", "br"))),
           widths=(32, *(128/len(factors) for f in factors)), font_size=10)
    pairs = [(c.name, c.left, c.right) for c in data.cases] or [("Par simultáneo de cargas", data.left.loads, data.right.loads)]
    for label, left, right in pairs:
        _body(document, f"Reacciones por metro transversal del caso {case_label(label)}.")
        _table(document, ("Lado", "DC", "DW", "PL", "LL+IM", "BR"),
               ((name, *(f"{getattr(loads, key):.3f}" for key in
                          ("pdc_tn_m", "pdw_tn_m", "ppl_tn_m", "pll_im_tn_m", "braking_tn_m")))
                for name, loads in (("Izquierda", left), ("Derecha", right))), widths=(35, 25, 25, 25, 25, 25))
    _body(document, ("Cada variante Ia o Ib se aplica globalmente a ambos estribos y la zapata combinada, sin "
          if uniform else "Cada variante Ia o Ib se aplica globalmente a ambos estribos, losa y transiciones, sin ") +
          "cruces entre regiones. Las reacciones izquierda y derecha son del mismo posicionamiento del tablero; "
          "no se cruzan máximos independientes. Se revisan ambos sentidos globales de frenado y sismo.")
    _body(document, "Evento Extremo I A utiliza PAE+0,5PIR; B utiliza max(0,5PAE,EH)+PIR. "
          "La condición sin tablero elimina sus reacciones e inercias y mantiene las demás acciones.")
    for i, case in enumerate(result.results, 1):
        _body(document, case_label(case.name))

    section(6)
    groups = combination_groups(result)
    for index, group in enumerate(groups,1):
        document.add_heading(f"6.{index}. Envolvente {group.name}",level=2)
        _body(document,"Casos incluidos: "+"; ".join(case_label(c) for c in group.cases)+".")
        for field,label in (("axial","axial N"),("shear","de cortante V"),("moment","de momento M")):
            _picture(document,charts[f"envolvente_{index}_{field}"],
                     f"Envolvente {label} sobre la estructura completa de {group.name}")
    _body(document, "N positivo corresponde a tracción. Las envolventes conservan signos y no son una única "
          "combinación simultánea. Las secciones críticas de flexión, corte y servicio se identifican por separado "
          "con sus coordenadas y esfuerzos simultáneos.")
    critical = max(result.results, key=lambda r:max(abs(u) for u in r.displacements[0::3]))
    _body(document, f"Desplazamiento horizontal absoluto máximo de nudos "
          f"{1000*max(abs(u) for u in critical.displacements[0::3]):.3f} mm en {case_label(critical.name)}. "
          "Corresponde al desplazamiento del modelo FRAME.")
    document.add_heading(f"6.{len(groups)+1}. Envolvente general de todas las combinaciones", level=2)
    _body(document, "Se incluyen todos los casos de Resistencia Ia, Resistencia Ib, Servicio I y Evento Extremo I, "
          "con y sin tablero según las condiciones analizadas. Se conservan los mínimos y máximos con signo "
          "de N, V y M en cada sección de la estructura.")
    for key, label in (("modelo_axial", "axial N"), ("modelo_cortante", "de cortante V"),
                       ("modelo_momento", "de momento M")):
        _picture(document, charts[key], f"Envolvente general {label} de todas las combinaciones")
    _body(document, "Esta envolvente general es la referencia de esfuerzos para el diseño de concreto armado "
          "del punto 8. Cada verificación conserva el caso gobernante y sus esfuerzos simultáneos en la sección "
          "correspondiente; las comprobaciones de servicio utilizan los casos de Servicio I.")

    section(7)
    from bridge_design.reporting.deck_docx import _add_native_equation
    _add_native_equation(document, "q_i = R_i/A_i")
    _body(document, "La presión máxima procede del análisis estructural de cada caso: es el mayor valor "
          "de presión entre sus resortes. Únicamente se desarrolla el cálculo de q límite: qadm en servicio, "
          "0,55·FS·qadm en resistencia y 0,80·FS·qadm en evento extremo.")
    _body(document, "Se presentan cuatro familias de envolventes: Resistencia Ia, Resistencia Ib, Servicio I "
          "y Evento Extremo I. Cada familia reúne todos sus casos simultáneos, con y sin tablero, "
          "y conserva los valores mínimos y máximos en cada resorte.")
    _body(document, "La resistencia nominal adoptada se estima como FS·qadm. Los factores de resistencia "
          "son 0.55 para resistencia y 0.80 para evento extremo.")
    _body(document, "Los asentamientos presentados son los desplazamientos elásticos del modelo Winkler "
          "con el módulo de balasto ks adoptado.")
    write_group_bearing(document, result, charts)

    from bridge_design.reporting.connected_distribution_report import write_distributed_design
    write_distributed_design(document, result, reinforcement, audit, charts)

    section(10)
    force_error = max(max(abs(r.equilibrium_error[0]), abs(r.equilibrium_error[1])) for r in result.results)
    moment_error = max(abs(r.equilibrium_error[2]) for r in result.results)
    _body(document, f"Máximo residuo global de fuerzas {force_error:.3e} tn y de momentos {moment_error:.3e} tn·m. "
          f"Máximo de iteraciones de contacto {max(r.iterations for r in result.results)}. "
          "Se comprueban equilibrio nodal y global y se rechaza inestabilidad o falta de convergencia.")
    if result.mesh_comparison:
        c = result.mesh_comparison
        _body(document, f"Refinamiento de {c.coarse_nodes} a {c.fine_nodes} resortes; paso de referencia "
              f"{c.coarse_step:.3f} a {c.fine_step:.3f} m. Cambios relativos: momento {c.moment_change:.2%}, "
              f"presión {c.pressure_change:.2%} y asentamiento {c.settlement_change:.2%}.")
    connected_summary_and_references(document, result)
    _enforce_uniform_typography(document)
    destination = Path(path).expanduser().resolve().with_suffix(".docx")
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


def write_initial_load_figures(document, result, charts):
    """Show the base actions before the combination factors."""
    data = result.inputs
    _body(document, "Los siguientes esquemas muestran las acciones iniciales sin factores LRFD. "
          "Los perfiles continuos muestran "
          "la variación por tramo; las flechas indican su sentido. Cada componente tiene una escala "
          "gráfica propia, con intensidades numéricas en tn/m para la franja de 1 m. "
          "Los rangos conservan el signo global: +x hacia la derecha y +z hacia arriba.")
    for key, caption in (("cargas_pesos", "Peso propio y relleno sin factorizar sobre la geometría real"),
                         ("cargas", "Perfil del empuje estático en ambas caras reales sin factorizar")):
        _picture(document, charts[key], caption)
    from bridge_design.domain.connected_cases import paired_cases
    for index, pair in enumerate(paired_cases(data), 1):
        _picture(document, charts[f"cargas_sobrecarga_{index}"],
                 f"Sobrecarga del relleno sin factores LRFD para {case_label(pair.name)}")
        _body(document, f"Ocupación de sobrecarga izquierda/derecha: {pair.left_surcharge:g}/{pair.right_surcharge:g}. "
              "Estos multiplicadores describen la condición simultánea de carga y se conservan "
              "en los perfiles; no son factores LRFD. Las flechas y círculos en el talón representan "
              "la transferencia nodal de equilibrio de la misma acción aplicada sobre la cara real. "
              "Los rangos de momentos indican los traslados equivalentes, no cargas adicionales.")
        _picture(document, charts[f"cargas_tablero_{index}"],
                 f"Reacciones iniciales y alturas físicas de aplicación para {case_label(pair.name)}")
    _body(document, "Los esquemas sísmicos siguientes son acciones base de evento extremo. "
          "Se presentan separados de las cargas de servicio. A y B son alternativas, "
          "y cada perfil ya contiene el empuje completo adoptado: no se suma nuevamente EH. "
          "Los coeficientes físicos kh y As se conservan; no se aplican factores LRFD. "
          "PIR se muestra antes del multiplicador 0,5 de la alternativa A; su sentido se invierte "
          "para EQ−. PEQ se muestra a su altura física en los esquemas del tablero.")
    for direction in ("menos", "mas"):
        for kind in ("A", "B"):
            _picture(document, charts[f"cargas_sismo_{kind}_{direction}"],
                     f"Forma real del empuje sísmico base alternativa {kind} sentido {direction}")
    _picture(document, charts["cargas_inercia"], "Inercia base del concreto y del relleno PIR")
