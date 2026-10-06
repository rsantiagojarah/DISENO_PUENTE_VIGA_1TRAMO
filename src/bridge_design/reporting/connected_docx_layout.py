"""Connected report front matter and summary using the abutment Word house style."""

from dataclasses import replace
from datetime import datetime, timedelta, timezone

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from bridge_design.reporting.connected_case_groups import case_label
from bridge_design.domain.connected_reinforcement import region_has_reinforcement_design
from bridge_design.domain.anchorage_status import anchorage_passes
from bridge_design.domain.connected_2_inputs import is_connected_2

from bridge_design.reporting.abutment_docx import _configure_abutment_header_footer, _labels
from bridge_design.reporting.deck_docx import (
    GRAY, TEAL, _body, _bottom_border, _configure_document, _format_run, _table,
)


SECTIONS = (
    "Bases y alcance del cálculo", "Geometría materiales y suelo",
    "Idealización estructural y resortes", "Determinación y aplicación de cargas",
    "Casos simultáneos y combinaciones", "Resultados del análisis estructural",
    "Contacto y verificaciones geotécnicas", "Diseño del concreto armado por región",
    "Servicio y detallado", "Verificación numérica",
    "Resumen y conclusiones", "Referencias",
)


def report_reinforcement(result):
    return tuple(steel for steel in result.reinforcement if region_has_reinforcement_design(steel.region))


def failing_controls(steel):
    controls = (("flexión", steel.flexural_utilization), ("cortante", steel.shear_utilization),
                ("servicio", steel.crack_utilization), ("acero mínimo", steel.minimum_utilization),
                ("acero transversal", steel.transverse_utilization))
    return tuple(f"{name} ({ratio:.3f})" for name, ratio in controls if ratio > 1+1e-8)


def connected_front_matter(document, result):
    _configure_document(document)
    title_properties = document.styles["Title"].element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    uniform = is_connected_2(result.inputs)
    labels = replace(_labels(is_pure_wall=False), header="MEMORIA DE CÁLCULO · ESTRIBOS CONECTADOS" + (" 2" if uniform else ""))
    _configure_abutment_header_footer(document, labels)
    for section in document.sections:
        for text in section.footer._element.xpath('.//w:t'):
            if text.text == "Actualizar campo":
                text.text = ""
    data = result.inputs
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(34)
    _format_run(paragraph.add_run("INGENIERÍA ESTRUCTURAL"), 11, bold=True, color=TEAL)
    _bottom_border(paragraph, TEAL, 16)
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(42)
    paragraph.add_run("Memoria de cálculo\nde estribos conectados" + (" 2" if uniform else ""))
    document.add_paragraph("Geometría · empujes · contacto · diseño estructural · detalle", style="Subtitle")
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(32)
    paragraph.paragraph_format.space_after = Pt(4)
    _format_run(paragraph.add_run("DOCUMENTO TÉCNICO DE DISEÑO"), 11, bold=True, color=TEAL)
    _table(document, ("Dato", "Descripción"), (
        ("Sistema", "Estribos iguales con base uniforme FRAME 2D" if uniform else "Dos estribos con cajuela y cimentación continua FRAME 2D"),
        ("Franja de cálculo", "1,00 m"),
        ("Separación libre", f"{data.clear_span_m:.2f} m"),
        ("Longitud de cimentación", f"{data.total_length_m:.2f} m"),
        ("q admisible", f"{data.soil.allowable_tn_m2:.3f} tn/m²"),
        ("Norma principal", "Manual de Puentes MTC 2018"),
        ("Fecha de emisión", datetime.now(timezone(timedelta(hours=-5))).strftime("%d/%m/%Y")),
    ), widths=(52, 108))
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(30)
    _format_run(paragraph.add_run(
        "La presente memoria presenta las hipótesis, cálculos, verificaciones y armaduras adoptadas."
    ), 11, italic=True, color=GRAY)
    document.add_page_break()
    document.add_heading("Contenido", level=1)
    _body(document, "La memoria sigue la secuencia del cálculo y presenta las armaduras adoptadas para cada región estructural.")
    _table(document, ("Sección", "Contenido"),
           ((str(index), label) for index, label in enumerate(SECTIONS, 1)), widths=(24, 136))
    document.add_page_break()


def connected_summary_and_references(document, result):
    document.add_heading("11. " + SECTIONS[10], level=1)
    _table(document, ("Distribución", "Acero elegido", "As cm²/m", "Estado"),
           ((steel.region, f"{steel.bar_label} @ {steel.spacing_m:.3f} m",
             f"{steel.area_per_face_cm2_m:.3f}", steel.status)
            for steel in report_reinforcement(result)), widths=(66, 40, 28, 24), accent=True)
    _body(document, "Las áreas y verificaciones corresponden a las barras y separaciones adoptadas. "
          "Los estados de la tabla resumen corresponden a flexión, cortante, servicio y acero mínimo.")
    for steel in report_reinforcement(result):
        controls = failing_controls(steel)
        if steel.status != "OK":
            _body(document, f"NO CUMPLE — {steel.region}: " + "; ".join(controls or ("mínimos o separación del armado",)) + ".")
        if steel.role == "primary" and steel.available_anchor_cm is not None and not anchorage_passes(steel.anchor_status):
            label = steel.anchor_status
            available = f"L disponible={steel.available_anchor_cm:.2f} cm"
            _body(document, f"{label} — {steel.region}: ld={steel.required_straight_anchor_cm:.2f} cm; "
                  f"{available}.")
    worst_pressure = max(result.foundation_checks, key=lambda row: row.bearing_utilization)
    _body(document, f"Gobierna la presión local {case_label(worst_pressure.case)}, "
          f"con índice {worst_pressure.bearing_utilization:.3f}.")
    for check in result.foundation_checks:
        if check.bearing_status != "OK":
            _body(document, f"NO CUMPLE — presión en {case_label(check.case)}: "
                  f"q={check.maximum_pressure:.3f} tn/m²; q límite={check.pressure_limit:.3f} tn/m².")
    document.add_heading("12. " + SECTIONS[11], level=1)
    for text in (
        "Manual de Puentes MTC 2018. Referencias específicas de materiales, empujes, resistencia, fisuración y desarrollo "
        "se indican junto a los cálculos que las utilizan.",
        "Euler Bernoulli y trabajo virtual para el FRAME elástico; modelo Winkler para la respuesta vertical del terreno.",
    ):
        _body(document, text)
