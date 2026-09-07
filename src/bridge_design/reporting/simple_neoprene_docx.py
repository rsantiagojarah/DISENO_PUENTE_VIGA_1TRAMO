"""Memoria de cálculo Word trazable para ``diseno-apoyos-neopreno``."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from bridge_design.domain.simple_neoprene_support import SimpleSupportResult
from bridge_design.reporting.deck_docx import (
    GRAY,
    RULE,
    TEAL,
    _body,
    _bottom_border,
    _calc,
    _comment,
    _configure_document,
    _enforce_uniform_typography,
    _field,
    _format_run,
    _table,
)
from bridge_design.reporting.simple_neoprene_docx_detail import (
    check_reference,
    effective_area_trace,
    free_area_trace,
    gross_area_trace,
    horizontal_demand_trace,
    h_pad_trace,
    modulus_trace,
    pin_interaction_trace,
    removed_area_trace,
    service_reaction_trace,
    shape_factor_trace,
    support_check_trace,
    thermal_movement_trace,
)

REF_GENERAL = "Manual de Puentes MTC 2018, Art. 2.10.4; AASHTO LRFD 14.7.6 (Método A, neopreno simple sin zunchos)."
REF_GEOMETRY = "AASHTO LRFD 14.7.5.1-1; Manual de Puentes MTC 2018, Art. 2.10.4."
REF_DEMANDS = "Manual de Puentes MTC 2018, Art. 2.4.5; AASHTO LRFD combinaciones de carga."
REF_THERMAL = "Manual de Puentes MTC 2018, criterios de movimiento térmico; AASHTO LRFD 14.7.6.3.4."
REF_HORIZONTAL = "AASHTO LRFD 3.10.9 / Resistencia I; Manual de Puentes MTC 2018, Art. 2.10.4.3.8."

HEADER = "MEMORIA DE CÁLCULO · APOYO NEOPRENO SIMPLE"
FILENAME_PREFIX = "MEMORIA_CALCULO_APOYO_NEOPRENO"


def generate_simple_neoprene_docx(result: SimpleSupportResult, output_path: str | Path) -> Path:
    """Crea la memoria Word del apoyo de neopreno simple y devuelve su ruta absoluta."""
    path = Path(output_path).expanduser().resolve()
    if path.suffix.lower() != ".docx":
        path = path.with_suffix(".docx")
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_document(document)
    _configure_header_footer(document)
    _cover(document, result)
    _contents(document, result)
    _design_basis(document, result)
    _geometry(document, result)
    _demands(document, result)
    _neoprene_verifications(document, result)
    _detail_verifications(document, result)
    _summary(document, result)
    _references(document)
    _enforce_uniform_typography(document)
    document.save(path)
    return path


def select_simple_neoprene_docx_save_path(result: SimpleSupportResult | None = None) -> Path | None:
    """Abre el diálogo nativo para elegir destino de la memoria Word."""
    try:
        import tkinter as tk
        from tkinter import filedialog
    except ImportError as exc:
        raise RuntimeError("Tkinter no está disponible para elegir el destino Word.") from exc
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    suffix = "FIJO" if result is None or result.inputs.support_type == "FIJO_BARRAS" else "MOVIL"
    default_name = f"{FILENAME_PREFIX}_{suffix}_{datetime.now():%Y%m%d_%H%M}.docx"
    try:
        selected = filedialog.asksaveasfilename(
            parent=root,
            title="Guardar memoria de cálculo del apoyo de neopreno simple",
            defaultextension=".docx",
            initialfile=default_name,
            filetypes=(("Documento de Word", "*.docx"),),
        )
    finally:
        root.destroy()
    return Path(selected) if selected else None


def generate_simple_neoprene_docx_with_dialog(result: SimpleSupportResult) -> Path | None:
    """Solicita destino y genera la memoria Word cuando el usuario confirma."""
    path = select_simple_neoprene_docx_save_path(result)
    if path is None:
        print("Generación de la memoria Word del apoyo de neopreno cancelada por el usuario.")
        return None
    generated = generate_simple_neoprene_docx(result, path)
    print(f"Memoria Word del apoyo de neopreno guardada en: {generated}")
    return generated


def _configure_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(HEADER)
    _format_run(run, 11, bold=True, color=TEAL)
    _bottom_border(header, RULE, 5)
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("DISEÑO ESTRUCTURAL  ·  ")
    _format_run(run, 11, color=GRAY)
    _field(footer, "PAGE")


def _support_title(result: SimpleSupportResult) -> str:
    if result.inputs.support_type == "FIJO_BARRAS":
        return "Apoyo fijo con pasadores lisos interiores"
    return "Apoyo móvil con planchas externas superior e inferior"


def _cover(document: Document, result: SimpleSupportResult) -> None:
    inp = result.inputs
    g = inp.geometry
    kind = "FIJO PASADORES" if inp.support_type == "FIJO_BARRAS" else "MOVIL PLACAS"
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(34)
    run = p.add_run("INGENIERÍA ESTRUCTURAL")
    _format_run(run, 11, bold=True, color=TEAL)
    _bottom_border(p, TEAL, 16)
    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(42)
    p.add_run("Memoria de cálculo\ndel apoyo de neopreno simple")
    p = document.add_paragraph(style="Subtitle")
    p.add_run(f"{kind} · geometría · demandas · verificaciones del neopreno · detalle constructivo")
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(32)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("DOCUMENTO TÉCNICO DE DISEÑO")
    _format_run(run, 11, bold=True, color=TEAL)
    _table(
        document,
        ("Dato", "Descripción"),
        (
            ("Sistema", _support_title(result)),
            ("Neopreno", f"{g.length_cm * 10:.0f} × {g.width_cm * 10:.0f} × {g.thickness_cm * 10:.0f} mm"),
            ("Dureza Shore A", f"{inp.hardness}"),
            ("Reacción de servicio", f"{inp.demands.r_service_tn:.3f} Tn"),
            ("f'c pedestal/estribo", f"{inp.fc_kg_cm2:.1f} kg/cm²"),
            ("Estado global", "CONFORME" if result.overall_ok else "NO CONFORME"),
            ("Norma principal", "Manual de Puentes MTC 2018"),
            ("Fecha de emisión", datetime.now().strftime("%d/%m/%Y")),
        ),
        widths=(42, 125),
        accent=True,
    )
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    run = p.add_run(
        "La presente memoria desarrolla las expresiones, leyendas, sustituciones numéricas, "
        "resultados y criterios técnicos correspondientes al apoyo de neopreno simple adoptado."
    )
    _format_run(run, 11, italic=True, color=GRAY)
    document.add_page_break()


def _contents(document: Document, result: SimpleSupportResult) -> None:
    detail_row = (
        ("5", "Verificaciones del detalle constructivo (pasadores o planchas)")
        if result.inputs.support_type == "FIJO_BARRAS"
        else ("5", "Verificaciones del detalle constructivo (planchas A36 y fricción)")
    )
    document.add_heading("Contenido", level=1)
    _body(
        document,
        "La memoria se organiza siguiendo la secuencia del modelo de cálculo y conserva "
        "estilos jerárquicos de Word con ecuaciones nativas trazables.",
    )
    _table(
        document,
        ("Sección", "Contenido desarrollado"),
        (
            ("1", "Bases de diseño, materiales y acciones de entrada"),
            ("2", "Geometría del neopreno, áreas y factor de forma"),
            ("3", "Demandas de servicio, movimiento térmico y acciones horizontales"),
            ("4", "Verificaciones del núcleo de neopreno simple (Método A PEP)"),
            detail_row,
            ("6", "Resumen de verificaciones y diseño adoptado"),
            ("7", "Referencias normativas"),
        ),
        widths=(25, 142),
        accent=True,
    )
    document.add_page_break()


def _design_basis(document: Document, result: SimpleSupportResult) -> None:
    inp = result.inputs
    dem = inp.demands
    temp = dem.temperature
    document.add_heading("1. Bases de diseño", level=1)
    _body(
        document,
        "El apoyo consiste en una sola capa de neopreno simple sin zunchos internos de acero. "
        "Las placas o pasadores son elementos externos al espesor h del elastómero. "
        "El módulo diseno-apoyos-neopreno es independiente de "
        "diseno-apoyos (apoyos elastoméricos con zunchos).",
    )
    document.add_heading("1.1 Materiales y parámetros", level=2)
    _table(
        document,
        ("Parámetro", "Símbolo", "Valor adoptado"),
        (
            ("Dureza Shore A", "Shore A", f"{inp.hardness}"),
            ("Módulo mínimo (capacidad)", "G_min", f"{result.g_min_kg_cm2:.2f} kg/cm²"),
            ("Módulo máximo (fuerza H)", "G_max", f"{result.g_max_kg_cm2:.2f} kg/cm²"),
            ("Coeficiente de fricción", "μ", f"{inp.friction_coefficient:.2f}"),
            ("Resistencia del concreto", "f'c", f"{inp.fc_kg_cm2:.1f} kg/cm²"),
            ("Factor térmico", "γ_TU", f"{dem.gamma_tu:.2f}"),
            ("Dilatación térmica", "α", f"{dem.alpha_per_c:.8f} /°C"),
        ),
        widths=(78, 30, 59),
    )
    document.add_heading("1.2 Acciones verticales por apoyo", level=2)
    _table(
        document,
        ("Acción", "Símbolo", "Valor (Tn)"),
        (
            ("Carga permanente estructural", "R_DC", f"{dem.r_dc_tn:.3f}"),
            ("Superficie de rodadura", "R_DW", f"{dem.r_dw_tn:.3f}"),
            ("Carga peatonal", "R_PL", f"{dem.r_pl_tn:.3f}"),
            ("Carga vehicular con impacto", "R_LL+IM", f"{dem.r_ll_im_tn:.3f}"),
            ("Reacción de servicio", "R", f"{dem.r_service_tn:.3f}"),
        ),
        widths=(83, 29, 55),
    )
    document.add_heading("1.3 Acciones horizontales y térmicas", level=2)
    horizontal_rows = [
        ("Luz del tramo", "L", f"{dem.span_length_m:.2f} m"),
        ("PGA", "PGA", f"{dem.pga:.3f}"),
        ("Factor sísmico del sitio", "Fpga", f"{dem.fpga:.3f}"),
        ("Coeficiente sísmico", "As = PGA·Fpga", f"{dem.as_coeff:.3f}"),
    ]
    if inp.support_type == "FIJO_BARRAS":
        horizontal_rows.extend(
            (
                ("Frenado nominal por apoyo", "BR", f"{dem.h_long_tn:.3f} Tn"),
                ("Frenado Resistencia I", "H_BR,RI", f"{dem.h_br_strength_tn:.3f} Tn"),
            )
        )
    horizontal_rows.append(("Sismo Evento Extremo I", "H_EQ,EE-I", f"{dem.h_eq_long_tn:.3f} Tn"))
    if temp is not None:
        horizontal_rows.extend(
            (
                ("Temperatura superior MTC", "T_sup", f"{temp.t_sup_c:.1f} °C"),
                ("Temperatura inferior MTC", "T_inf", f"{temp.t_inf_c:.1f} °C"),
                ("Temperatura de instalación", "T_inst", f"{temp.t_install_c:.1f} °C"),
                ("Contracción térmica", "ΔT", f"{temp.contraction_delta_t_c:.1f} °C"),
            )
        )
    _table(document, ("Parámetro", "Símbolo", "Valor"), tuple(horizontal_rows), widths=(78, 30, 59))
    _comment(document, REF_GENERAL)


def _geometry(document: Document, result: SimpleSupportResult) -> None:
    g = result.inputs.geometry
    document.add_heading("2. Geometría del neopreno y factor de forma", level=1)
    _body(document, "Se determinan las áreas resistentes y el factor de forma S empleado en el Método A.")
    _table(
        document,
        ("Parámetro", "Símbolo", "Valor"),
        (
            ("Largo en planta", "L", f"{g.length_cm:.2f} cm"),
            ("Ancho en planta", "W", f"{g.width_cm:.2f} cm"),
            ("Espesor total", "h", f"{g.thickness_cm:.2f} cm"),
            ("Área bruta", "A_g", f"{result.gross_area_cm2:.2f} cm²"),
            ("Área descontada", "A_h", f"{result.removed_area_cm2:.2f} cm²"),
            ("Área efectiva", "A_eff", f"{result.effective_area_cm2:.2f} cm²"),
            ("Área libre", "A_libre", f"{result.free_area_cm2:.2f} cm²"),
            ("Factor de forma", "S", f"{result.shape_factor:.3f}"),
        ),
        widths=(78, 30, 59),
    )
    _calc(document, "Área bruta del neopreno", *gross_area_trace(result), REF_GEOMETRY)
    _calc(document, "Área descontada por pasadores", *removed_area_trace(result), REF_GEOMETRY)
    _calc(document, "Área efectiva en compresión", *effective_area_trace(result), REF_GEOMETRY)
    _calc(document, "Área libre lateral", *free_area_trace(result), REF_GEOMETRY)
    _calc(document, "Factor de forma", *shape_factor_trace(result), REF_GEOMETRY)
    _calc(document, "Módulos del elastómero y límite de compresión", *modulus_trace(result), REF_GEOMETRY)


def _demands(document: Document, result: SimpleSupportResult) -> None:
    document.add_heading("3. Demandas de servicio y movimiento", level=1)
    _calc(document, "Reacción de servicio por apoyo", *service_reaction_trace(result), REF_DEMANDS)
    _calc(document, "Movimiento térmico de contracción", *thermal_movement_trace(result), REF_THERMAL)
    if result.inputs.support_type == "FIJO_BARRAS":
        _calc(document, "Demanda horizontal de diseño del apoyo fijo", *horizontal_demand_trace(result), REF_HORIZONTAL)
    else:
        _calc(document, "Fuerza horizontal por corte del neopreno móvil", *h_pad_trace(result), REF_HORIZONTAL)


def _neoprene_verifications(document: Document, result: SimpleSupportResult) -> None:
    document.add_heading("4. Verificaciones del núcleo de neopreno", level=1)
    _body(
        document,
        "Las verificaciones siguientes corresponden al elastómero simple sin zunchos internos, "
        "con los límites PEP del Método A.",
    )
    core_names = {
        "Compresion neopreno",
        "Deflexion por compresion",
        "Estabilidad h <= L/3",
        "Estabilidad h <= W/3",
        "Aplastamiento concreto bajo apoyo",
    }
    for check in result.checks:
        if check.name in core_names:
            _calc(document, check.name, *support_check_trace(check), check_reference(check.name))


def _detail_verifications(document: Document, result: SimpleSupportResult) -> None:
    document.add_heading("5. Verificaciones del detalle constructivo", level=1)
    core_names = {
        "Compresion neopreno",
        "Deflexion por compresion",
        "Estabilidad h <= L/3",
        "Estabilidad h <= W/3",
        "Aplastamiento concreto bajo apoyo",
    }
    if result.inputs.support_type == "FIJO_BARRAS":
        _body(document, "El apoyo fijo transfiere la restricción horizontal mediante pasadores lisos ASTM A108 Gr. 1020.")
        bars = result.inputs.fixed_bars
        if bars is not None:
            _table(
                document,
                ("Parámetro", "Valor"),
                (
                    ("Número de pasadores", f"{bars.n_bars}"),
                    ("Diámetro", f"{bars.diameter_cm:.2f} cm ({bars.diameter_cm / 2.54:.2f} pulg)"),
                    ("Separación longitudinal", f"{bars.spacing_long_cm:.2f} cm"),
                    ("Separación transversal", f"{bars.spacing_trans_cm:.2f} cm"),
                    ("Brazo e entre contactos", f"{bars.moment_arm_cm:.2f} cm"),
                    ("Fluencia pasador", f"{bars.fy_kg_cm2:.0f} kg/cm²"),
                ),
                widths=(90, 77),
            )
        _calc(document, "Interacción corte-flexión del pasador", *pin_interaction_trace(result), REF_HORIZONTAL)
    else:
        _body(document, "El apoyo móvil permite desplazamiento térmico mediante planchas externas ASTM A36.")
        plates = result.inputs.plates
        g = result.inputs.geometry
        if plates is not None:
            _table(
                document,
                ("Parámetro", "Valor"),
                (
                    ("Plancha superior/inferior", f"{g.length_cm * 10:.0f} × {g.width_cm * 10:.0f} × {plates.thickness_cm * 10:.0f} mm"),
                    ("Fluencia A36", f"{plates.fy_kg_cm2:.0f} kg/cm²"),
                    ("Resistencia última A36", f"{plates.fu_kg_cm2:.0f} kg/cm²"),
                    ("H_pad por corte", f"{result.h_pad_tn:.3f} Tn"),
                    ("Capacidad por fricción", f"{result.friction_capacity_tn:.3f} Tn"),
                ),
                widths=(90, 77),
            )
    for check in result.checks:
        if check.name not in core_names:
            _calc(document, check.name, *support_check_trace(check), check_reference(check.name))


def _summary(document: Document, result: SimpleSupportResult) -> None:
    inp = result.inputs
    g = inp.geometry
    document.add_heading("6. Resumen de verificaciones y diseño adoptado", level=1)
    rows = tuple(
        (
            check.name,
            f"{check.demand:.3f} {check.unit}",
            f"{check.limit:.3f} {check.unit}",
            f"{check.ratio:.3f}" if check.ratio is not None else "—",
            check.status,
        )
        for check in result.checks
    )
    _table(
        document,
        ("Verificación", "Demanda", "Límite", "Ratio", "Estado"),
        rows,
        widths=(52, 28, 28, 20, 19),
    )
    document.add_heading("6.1 Diseño adoptado", level=2)
    summary_rows = [
        ("Estado global", "CONFORME" if result.overall_ok else "NO CONFORME"),
        ("Neopreno", f"{g.length_cm * 10:.0f} × {g.width_cm * 10:.0f} × {g.thickness_cm * 10:.0f} mm, Shore {inp.hardness}"),
        ("Reacción de servicio", f"{inp.demands.r_service_tn:.3f} Tn"),
    ]
    if inp.support_type == "FIJO_BARRAS" and inp.fixed_bars is not None:
        bars = inp.fixed_bars
        summary_rows.extend(
            (
                ("Detalle fijo", f"{bars.n_bars} pasadores lisos d = {bars.diameter_cm / 2.54:.2f} pulg"),
                ("Demanda horizontal", f"{result.h_design_tn:.3f} Tn ({result.h_design_case})"),
            )
        )
    elif inp.plates is not None:
        summary_rows.extend(
            (
                ("Planchas A36", f"t = {inp.plates.thickness_cm * 10:.0f} mm"),
                ("H_pad", f"{result.h_pad_tn:.3f} Tn"),
            )
        )
    _table(document, ("Elemento", "Valor adoptado"), tuple(summary_rows), widths=(70, 97))


def _references(document: Document) -> None:
    document.add_heading("7. Referencias normativas", level=1)
    _table(
        document,
        ("Referencia", "Aplicación"),
        (
            ("Manual de Puentes MTC 2018, Art. 2.10.4", "Apoyos elastoméricos de neopreno simple"),
            ("Manual de Puentes MTC 2018, Art. 2.10.4.3.2", "Compresión σ ≤ min(G·S, 0.80 ksi)"),
            ("Manual de Puentes MTC 2018, Art. 2.10.4.3.3", "Deflexión δ ≤ 0.09 h"),
            ("Manual de Puentes MTC 2018, Art. 2.10.4.3.4", "Cortante h ≥ 2·Δ_s y γ ≤ 0.50"),
            ("Manual de Puentes MTC 2018, Art. 2.10.4.3.6", "Estabilidad h ≤ L/3 y h ≤ W/3"),
            ("AASHTO LRFD 14.7.6", "Método A para neopreno simple (PEP)"),
            ("AASHTO LRFD 14.6.3.1", "Fuerza horizontal H = G·A·Δ/h"),
            ("AASHTO LRFD 14.8.3", "Fricción y anclaje de apoyos"),
            ("AASHTO LRFD 5.7.5", "Aplastamiento del concreto"),
            ("AASHTO LRFD 6.13.1", "Interacción corte-flexión de pasadores"),
        ),
        widths=(78, 89),
    )
