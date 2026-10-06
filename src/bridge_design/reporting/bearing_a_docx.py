"""Word Método A con el formato tipográfico de diseno-apoyos-neopreno.

Las expresiones y verificaciones provienen del mismo registro que consola y
JSON. La figura se dibuja con la composición realmente adoptada.
"""

from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt
from PIL import Image, ImageDraw, ImageFont

from bridge_design.domain.bearing_method_a import BearingAResult, MTC_URL
from bridge_design.reporting.deck_docx import (
    GRAY, TEAL, _add_native_equation, _append_math_expression, _body, _configure_document,
    _enforce_uniform_typography, _field, _format_run, _math_run, _split_top_level, _table,
)


def select_bearing_a_docx_path() -> Path | None:
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost",True)
    try:
        path = filedialog.asksaveasfilename(
            parent=root,title="Guardar memoria de cálculo del apoyo Método A",
            defaultextension=".docx",initialfile="MEMORIA_CALCULO_APOYO_METODO_A.docx",
            filetypes=(("Documento de Word","*.docx"),),
        )
    finally:
        root.destroy()
    return Path(path) if path else None


def _diagram(result: BearingAResult) -> BytesIO:
    g = result.adopted
    canvas = Image.new("RGB",(1700,1000),"white")
    d = ImageDraw.Draw(canvas)
    try:
        font = ImageFont.truetype("arial.ttf",30)
        small = ImageFont.truetype("arial.ttf",26)
    except OSError:
        font = ImageFont.load_default(size=30)
        small = ImageFont.load_default(size=26)
    d.text((70,30),"APOYO ELASTOMERICO REFORZADO CON ACERO",font=font,fill="black")
    d.text((100,100),"Planta",font=font,fill="black")
    scale = min(600/g.width_cm,520/g.length_cm)
    x,y = 100,165
    w,l = g.width_cm*scale,g.length_cm*scale
    d.rectangle((x,y,x+w,y+l),fill="#ededed",outline="black",width=4)
    d.text((x,y+l+20),f"W = {g.width_cm*10:g} mm (transversal)",font=small,fill="black")
    d.text((x,y+l+65),f"L = {g.length_cm*10:g} mm (longitudinal)",font=small,fill="black")
    d.line((x+w*.5,y+l*.8,x+w*.5,y+l*.2),fill="black",width=4)
    d.polygon(((x+w*.5,y+l*.12),(x+w*.5-13,y+l*.2),(x+w*.5+13,y+l*.2)),fill="black")
    d.text((x+25,y+40),"Direccion del trafico",font=small,fill="black")
    d.text((850,100),"Seccion y composicion",font=font,fill="black")
    scale_h = 480/result.value("HEIGHT")
    x,y = 850,165
    composition = [g.exterior_cm]
    for _ in range(g.interior_layers):
        composition.extend((g.steel_cm,g.interior_cm))
    composition.extend((g.steel_cm,g.exterior_cm))
    for index,h in enumerate(composition):
        steel = index%2==1
        end = y+h*scale_h
        d.rectangle((x,y,x+620,end),fill="#383838" if steel else "#e5e5e5",outline="black",width=2)
        y = end
    d.text((x,y+20),f"H = {result.value('HEIGHT')*10:g} mm",font=small,fill="black")
    d.text((850,730),f"{g.interior_layers} capas interiores de {g.interior_cm*10:g} mm",font=small,fill="black")
    d.text((850,775),f"2 capas exteriores de {g.exterior_cm*10:g} mm",font=small,fill="black")
    d.text((850,820),f"{g.interior_layers+1} zunchos de {g.steel_cm*10:g} mm",font=small,fill="black")
    if g.hole_count:
        d.text((100,855),f'{g.hole_count} perforaciones pasantes de diametro {g.hole_diameter_cm*10:g} mm',font=small,fill='black')
        d.text((100,890),'Ubicacion de agujeros no representada en este esquema.',font=small,fill='black')
    d.text((100,920),"Esquema de capas; dimensiones y recubrimientos según la geometría adoptada.",font=small,fill="black")
    stream = BytesIO()
    canvas.save(stream,format="PNG")
    stream.seek(0)
    return stream


from dataclasses import replace
from bridge_design.reporting.bearing_a_labels import (
    CONTROL_IDS, compression_basis, conclusion, design_checks, number, status, title,
)
from bridge_design.reporting.abutment_docx import _configure_abutment_header_footer, _labels


def _source(text):
    if text.startswith('Ingreso manual del usuario') and 'LL/IM aproximados' in text:
        return 'Reacciones de cálculo con separación aproximada de carga vehicular e impacto mediante IM = 0.33 LL.'
    if text.startswith('Ingreso manual del usuario'):
        return text.replace('Ingreso manual del usuario', 'Reacciones adoptadas', 1)
    if text.startswith('Indicar modelo'):
        return 'Reacciones adoptadas para el apoyo.'
    if text.startswith('Serquén p. 231'):
        return 'Serquén, p. 231, Fig. C14.7.6.3.3-1 AASHTO; interpolación de lecturas gráficas.'
    return text


def _expression(text):
    for before, after in (
        ('epsilon_est', 'ε_est'), ('epsilon', 'ε'), ('sigma', 'σ'),
        ('Delta', 'Δ'), ('delta_creep', 'δ_fluencia'), ('delta', 'δ'),
        ('gamma', 'γ'), ('alpha', 'α'), ('phi', 'φ'), ('creep', 'fluencia'),
        ('<=', '≤'), ('>=', '≥'), (' - ', ' − '),
    ):
        text = text.replace(before, after)
    return text


def _trace(document, step, heading, result):
    """Usar el mismo bloque de desarrollo que tablero y estribos conectados."""
    from bridge_design.reporting.deck_docx import _calc

    first = len(document.paragraphs)
    legend = step.legend
    if step.id == 'FRICTION':
        legend = 'μ: coeficiente de fricción; P_min: reacción vertical mínima adoptada para la condición de servicio (Tn)'
    reference = step.reference
    if step.id == 'P':
        reference = 'MTC 2.4.5.3.1, Servicio I. ' + _source(result.inputs.actions.source)
    elif step.id.startswith('EPS_'):
        reference = _source(reference)
    elif step.id == 'DELTA' and result.inputs.horizontal_displacement_cm is not None:
        reference = 'Condición cinemática adoptada del esquema del tramo; MTC 2.10.2.1.1 y 2.10.4.3.4 definen los movimientos de diseño'

    substitution = _expression(step.substitution)
    # Identificar el valor calculado permite escribir una igualdad numérica,
    # en lugar de dejar una operación aislada como texto corriente.
    if not any(operator in substitution for operator in ('=', '≤', '≥', '<', '>')):
        left = _expression(step.formula).split('=')[0].strip()
        substitution = left + ' = ' + substitution
    if step.value is not None and not any(operator in substitution for operator in ('≤', '≥', '<', '>')):
        substitution += ' = ' + number(step.value, step.unit)

    outcome = number(step.value, step.unit, percent=step.id == 'STRAIN_CHECK')
    if step.limit is not None:
        outcome += '; límite ' + ('< ' if step.strict else '≤ ') + number(step.limit, step.unit, percent=step.id == 'STRAIN_CHECK')
    if step.status != 'CALCULADO':
        outcome += '; ' + status(step) + '.'
    comment = ''
    if step.id == 'SCOPE':
        comment = 'El límite 22 corresponde al MTC; los límites 20 y 16 proceden del comentario C14.7.6.1 aplicado por Serquén.'
    elif step.id == 'JOINT':
        comment = 'Serquén reproduce la recomendación de 1/8 in para deflexión relativa por carga viva de C14.7.5.3.6. En su ejemplo de la p. 238 incluye además la fluencia; aquí se adopta ese criterio conservador.'
    elif step.id == 'HS_MIN':
        comment = 'El MTC p. 508 imprime 0.625 in junto a 1.588 mm, valores incompatibles. Se adopta 1/16 in = 1.5875 mm, concordante con el valor métrico del MTC y con Serquén.'
    elif step.status in ('REFERENCIAL', 'CUMPLE (ESTIMADO)'):
        comment = compression_basis(result)
    _calc(document, heading + ' ' + title(step), _expression(step.formula),
          _expression(legend), substitution, outcome, comment, reference)
    for paragraph in document.paragraphs[first:]:
        if paragraph.text.startswith('Referencia normativa:') and (step.id in ('AREA', 'HS_HOLES', 'AREA_REQ', 'HRT', 'HEIGHT', 'DEF_T', 'DEF_LL', 'JOINT', 'FRICTION') or (step.id == 'DELTA' and result.inputs.horizontal_displacement_cm is not None)):
            for run in paragraph.runs:
                run.text = run.text.replace('Referencia normativa:', 'Base de cálculo y referencia técnica:')
        if paragraph.style.name == 'Equation':
            paragraph.paragraph_format.keep_with_next = False
        if not paragraph.text and not paragraph._p.xpath('.//m:oMath'):
            paragraph._p.getparent().remove(paragraph._p)


def generate_bearing_a_docx(result: BearingAResult, output_path: str | Path, *, document=None, show_curves=True) -> Path:
    path = Path(output_path).expanduser().resolve().with_suffix('.docx')
    path.parent.mkdir(parents=True, exist_ok=True)
    g, i = result.adopted, result.inputs
    doc = document
    if doc is None:
        doc = Document()
        _configure_document(doc)
        for name in ('Title', 'Subtitle'):
            ppr = doc.styles[name]._element.pPr
            if ppr is not None:
                for child in list(ppr):
                    if child.tag.endswith('}pBdr'):
                        ppr.remove(child)
        _configure_abutment_header_footer(doc, replace(_labels(is_pure_wall=False), header='MEMORIA DE CÁLCULO · APOYO ELASTOMÉRICO MÉTODO A'))
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(34)
        _format_run(p.add_run('INGENIERÍA ESTRUCTURAL'), 11, bold=True, color=TEAL)
        p = doc.add_paragraph(style='Title')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(42)
        p.add_run('Memoria de cálculo\ndel apoyo elastomérico Método A')
        doc.add_paragraph('Geometría acciones deformaciones refuerzo y estabilidad', style='Subtitle')
        _table(doc, ('Dato', 'Descripción'), (
            ('Proyecto', i.project), ('Apoyo', i.bearing_id),
            ('Planta y altura', f'{g.length_cm*10:g} × {g.width_cm*10:g} × {result.value("HEIGHT")*10:g} mm'),
            ('Elastómero', f'Shore A {i.hardness}'),
            ('Norma principal', 'Manual de Puentes MTC 2018, Artículo 2.10.4'),
            ('Procedimiento de cálculo', 'Serquén, capítulo 4, problema 4.1'),
            ('Fecha de emisión', datetime.now(timezone(timedelta(hours=-5))).strftime('%d/%m/%Y')),
        ), widths=(52, 107))
        _body(doc, 'La memoria desarrolla las acciones, geometría, deformaciones, refuerzo de acero, estabilidad y fricción del apoyo. ' + conclusion(result))
        doc.add_page_break()
    groups = tuple(dict.fromkeys(s.section for s in result.steps))
    group_names = {'Fuerzas horizontales y conexiones': 'Fuerzas horizontales y fricción' if i.neoprene_only else 'Fuerzas horizontales y conexiones'}
    summary_no = len(groups) + 2
    sections = ['Bases de diseño e hipótesis'] + [group_names.get(x, x) for x in groups] + ['Resumen de verificaciones y conclusiones', 'Geometría y composición del apoyo', 'Referencias']
    doc.add_heading('Contenido', 1)
    _table(doc, ('Sección', 'Contenido'), tuple((str(k), x) for k, x in enumerate(sections, 1)), widths=(24, 135))
    doc.add_page_break()
    doc.add_heading('1 Bases de diseño e hipótesis', 1)
    _body(doc, 'El apoyo es rectangular y está reforzado con zunchos de acero, ' + ('con perforaciones pasantes en el elastómero y los zunchos' if g.hole_count else 'sin perforaciones') + ' y sin superficie deslizante. L corresponde a la dirección longitudinal y W a la transversal. Las capas interiores tienen igual espesor y están adheridas por ambas caras; las exteriores están adheridas por una cara.')
    _body(doc, 'Las fuerzas se expresan en toneladas-fuerza, las longitudes en centímetros y los esfuerzos en kgf/cm². La reacción de servicio incluye el incremento dinámico; el límite de deformación por compresión se evalúa con las acciones de servicio sin impacto.')
    if i.horizontal_displacement_cm is not None:
        _body(doc, 'La evaluación corresponde al cuerpo de neopreno del extremo fijo. Su desplazamiento relativo longitudinal adoptado es '+number(i.horizontal_displacement_cm,'cm')+'. La condición cinemática del punto fijo se adopta del esquema estructural.')
    else:
        _body(doc, 'La evaluación corresponde al apoyo elastomérico y a su fricción en servicio.' if i.neoprene_only else 'La evaluación incluye el apoyo elastomérico, las demandas de conexión y la resistencia local del concreto de soporte.')
    _body(doc, compression_basis(result))
    if i.compression_curve and i.compression_curve.kind == 'referencia':
        _body(doc, 'La base de compresión corresponde a curvas gráficas de referencia y no a ensayos de un producto específico.')
    if g.hole_count:
        _body(doc, f'Se adoptan {g.hole_count} perforaciones circulares iguales de diámetro {g.hole_diameter_cm*10:g} mm, pasantes y sin contribución resistente de elementos alojados en ellas. Sus paredes se consideran libres de abombamiento. Para el zuncho se adopta conservadoramente que todos los agujeros atraviesan una misma sección del lado menor del núcleo, con ancho neto b n = min(Lc,Wc) − N d. Esta hipótesis no define su ubicación constructiva.')
    direction = 'alrededor del eje transversal' if g.principal_rotation_transverse else 'en una dirección distinta del eje transversal'
    _body(doc, 'La rotación principal se considera ' + direction + '. ' + ('Su evaluación está implícita en los límites geométricos y de esfuerzo del Método A, según MTC 2.10.4.1 y Serquén p. 239.' if g.principal_rotation_transverse else 'Esta condición no satisface el alcance del Método A definido en MTC 2.10.4.1.'))
    if g.total_height_cm is not None:
        _body(doc, f'Altura especificada H = {g.total_height_cm*10:g} mm; altura de la composición H = {result.value("HEIGHT")*10:g} mm. Compatibilidad geométrica: {status(result.step("HEIGHT_TARGET"))}.')
    pmin = i.actions.min_vertical_tn if i.actions.min_vertical_tn is not None else i.actions.dc_tn
    t = i.movements.temperature
    parameters = (
        ('Procedencia de las reacciones', _source(i.actions.source)),
        ('DC / DW / LL / PL / IM', f'{i.actions.dc_tn:g} / {i.actions.dw_tn:g} / {i.actions.ll_tn:g} / {i.actions.pl_tn:g} / {i.actions.im_tn:g} Tn'),
        ('Reacción mínima para fricción', f'{pmin:g} Tn' + ('; hipótesis P mínima = DC' if i.actions.min_vertical_tn is None else '')),
        ('Longitud efectiva de expansión', f'{i.movements.span_length_m:g} m'),
        ('Temperaturas mínima instalación máxima', f'{t.t_inf_c:g} / {t.t_install_c:g} / {t.t_sup_c:g} °C'),
        ('Retracción postensado otros acortamientos', f'{i.movements.shrinkage_cm:g} / {i.movements.prestress_shortening_cm:g} / {i.movements.other_permanent_cm:g} cm'),
        ('Coeficiente térmico y factor de movimiento', f'α = {i.movements.alpha_per_c:g}/°C; γ TU = {i.movements.gamma_tu:g}; ' + ('rango térmico completo' if not i.movements.use_install_to_min else 'desde la temperatura de instalación')),
        ('Módulos de corte mínimo y máximo', f'{({50:(6.68,9.14),60:(9.14,14.06)}[i.hardness])[0]:g} / {({50:(6.68,9.14),60:(9.14,14.06)}[i.hardness])[1]:g} kgf/cm²'),
        ('Fluencia del elastómero', f'C d = {0.25 if i.hardness == 50 else 0.35:g}'),
        ('Perforaciones pasantes en neopreno y zunchos', f'{g.hole_count} de diámetro {g.hole_diameter_cm*10:g} mm' if g.hole_count else 'Sin perforaciones'),
        ('Límite elástico del zuncho', f'{i.fy_kg_cm2:g} kgf/cm²'),
        ('Coeficiente de fricción', f'μ = {i.connections.friction_mu:g}'),
        ('Límite de deflexión de junta', f'{i.joint_limit_cm:g} cm; criterio adoptado del comentario C14.7.5.3.6'),
    )
    if i.horizontal_displacement_cm is not None:
        parameters = tuple(row for row in parameters if row[0] not in {
            "Reacción mínima para fricción", "Longitud efectiva de expansión",
            "Temperaturas mínima instalación máxima", "Retracción postensado otros acortamientos",
            "Coeficiente térmico y factor de movimiento", "Coeficiente de fricción"})
        parameters += (("Desplazamiento relativo adoptado", number(i.horizontal_displacement_cm,"cm")),)
    _table(doc, ("Parámetro", "Valor adoptado"), parameters, widths=(65,94))
    if not i.neoprene_only:
        _table(doc, ('Parámetro de soporte', 'Valor'), (
            ("Resistencia del concreto f c", f'{i.fc_kg_cm2:g} kgf/cm²; φ = {i.concrete_phi:g}'),
            ('Aceleración de sitio', f'{i.connections.as_site:g}'),
            ('Fuente de resistencia de conexiones', i.connections.resistance_source or 'Resistencia de conexión no determinada'),
        ), widths=(65, 94))
    figure_count = sum(p.style.name == "Caption" and p.text.startswith("Figura ") for p in doc.paragraphs)
    for chapter, group in enumerate(groups, 2):
        heading = doc.add_heading(f'{chapter} {group_names.get(group, group)}', 1)
        if show_curves and i.compression_curve and any(s.section == group and s.id == 'EPS_ID' for s in result.steps):
            heading.paragraph_format.page_break_before = True
        count = 0
        for step in result.steps:
            if step.section != group or step.id in CONTROL_IDS:
                continue
            if step.id.startswith('EPS_') and step.id != 'EPS_ID':
                continue
            count += 1
            if step.id.startswith('EPS_'):
                # Agrupar las cuatro lecturas de curvas en una sola tabla física.
                if step.id != 'EPS_ID':
                    continue
                doc.add_heading(f'{chapter} {count} Deformaciones unitarias por compresión', 2)
                if show_curves and i.compression_curve:
                    from bridge_design.reporting.bearing_a_curves import compression_chart, original_serquen_curve
                    original = original_serquen_curve(i.compression_curve)
                    if original:
                        _body(doc, f'Curvas esfuerzo deformación de referencia para elastómero de dureza Shore A {i.hardness}. Las familias de curvas corresponden al factor de forma S de la capa.')
                        doc.add_picture(str(original), width=Mm(159))
                        figure_count += 1
                        doc.add_paragraph(f'Figura {figure_count} Curvas originales de compresión para Shore A {i.hardness}', style='Caption')
                        _body(doc, 'Fuente: Arturo Rodríguez Serquén, Puentes, página impresa 231, Fig. C14.7.6.3.3-1 AASHTO; reproducción del documento APOYOS.pdf.')
                        doc.add_page_break()
                    doc.add_heading('Lecturas de las curvas adoptadas', 2)
                    doc.add_picture(compression_chart(result), width=Mm(159))
                    figure_count += 1
                    doc.add_paragraph(f'Figura {figure_count} Curvas utilizadas y puntos de cálculo del apoyo', style='Caption')
                    _body(doc, 'Las líneas finas representan las curvas de los datos adoptados; las líneas resaltadas corresponden a los factores de forma interior y exterior. Los puntos 1 a 4 se presentan en el mismo orden que las filas de la tabla siguiente. Fuente: ' + _source(i.compression_curve.source))
                _table(doc, ('Capa', 'Acción', 'Factor S', 'Esfuerzo kgf/cm²', 'Deformación %'), tuple(
                    ('Interior' if x == 'I' else 'Exterior', 'Permanente' if y == 'D' else 'Total sin impacto',
                     f'{result.value("SI" if x == "I" else "SE"):.4g}',
                     f'{1000*((i.actions.dc_tn+i.actions.dw_tn) if y == "D" else result.value("P")-i.actions.im_tn)/result.value("AREA"):.5g}',
                     number(result.step('EPS_'+x+y).value, percent=True))
                    for x in ('I', 'E') for y in ('D', 'T')
                ), widths=(28, 40, 23, 36, 32))
                _body(doc, 'Las deformaciones se obtienen para el factor de forma y el esfuerzo de cada capa. La interpolación se realiza dentro del dominio de las curvas, sin extrapolación. Fuente: ' + (_source(i.compression_curve.source) if i.compression_curve else 'modelo elástico aproximado σ / [3 G (1 + 2 k S²)].'))
                continue
            _trace(doc, step, f'{chapter} {count}', result)
    section_heading = doc.add_heading(f'{summary_no} Resumen de verificaciones y conclusiones', 1)
    section_heading.paragraph_format.page_break_before = True
    _table(doc, ('Verificación', 'Valor calculado', 'Límite', 'Resultado'), tuple(
        (title(s), number(s.value, s.unit, percent=s.id == 'STRAIN_CHECK'),
         ('< ' if s.strict else '≤ ') + number(s.limit, s.unit, percent=s.id == 'STRAIN_CHECK'), status(s))
        for s in design_checks(result)
    ), widths=(53, 33, 33, 40))
    _body(doc, conclusion(result))
    _body(doc, 'Los límites de compresión, corte, refuerzo y estabilidad corresponden al Manual de Puentes. El procedimiento de Serquén complementa la aplicación del Método A; los límites conservadores de aplicabilidad, el criterio de junta y el coeficiente de fricción proceden de comentarios AASHTO.')
    section_heading = doc.add_heading(f'{summary_no+1} Geometría y composición del apoyo', 1)
    section_heading.paragraph_format.page_break_before = True
    _table(doc, ('Elemento', 'Cantidad', 'Espesor mm'), (
        ('Capas interiores', str(g.interior_layers), f'{g.interior_cm*10:g}'),
        ('Capas exteriores', '2', f'{g.exterior_cm*10:g}'),
        ('Zunchos de acero', str(g.interior_layers+1), f'{g.steel_cm*10:g}'),
        ('Recubrimiento lateral', 'Por lado', f'{g.cover_cm*10:g}'),
    ), widths=(75, 40, 44))
    doc.add_picture(_diagram(result), width=Mm(159))
    doc.add_paragraph(f'Figura {figure_count+1} Planta y composición del apoyo elastomérico', style='Caption')
    _body(doc, 'Los zunchos y las capas de elastómero forman una unidad adherida. La geometría de fabricación comprende los espesores, recubrimientos y orientación indicados.')
    section_heading = doc.add_heading(f'{summary_no+2} Referencias', 1)
    section_heading.paragraph_format.page_break_before = True
    _table(doc, ('Fuente', 'Aplicación'), (
        ('Manual de Puentes MTC 2018, 2.10.4', 'Método A; alcance, compresión, corte, estabilidad y refuerzo.'),
        ('Manual de Puentes MTC 2018, 2.10.3.3.5 y 2.10.3.3.6', 'Zunchos y deflexiones por compresión.'),
        ('Manual de Puentes MTC 2018, 2.4.3.9.2 y 2.4.5.3.1', 'Temperatura y combinaciones de acciones.'),
        ('Manual de Puentes MTC 2018, Tabla 2.9.4.6.1.2.3-3', 'Umbral de fatiga de categoría A: 24 ksi; Serquén adopta 1687 kgf/cm².'),
        ('AASHTO LRFD correlativos impresos en MTC 2018', '14.7.6; 14.7.5.1; 14.7.5.3.5; 14.7.5.3.6 y 14.6.3.1. Numeración cotejada mediante el MTC.'),
        ('FHWA Comprehensive Design Example, Design Step 6.1.2.1 y 6.1.2.6', 'Perforaciones: factor de forma y aumento de espesor del zuncho. https://www.fhwa.dot.gov/bridge/lrfd/pscus06.cfm'),
        ('Arturo Rodríguez Serquén, Puentes, capítulo 4, APOYOS.pdf', 'Páginas impresas 229 a 239; incluye los comentarios C14.7.6.1, C14.7.5.3.6 y C14.8.3.1 y el problema 4.1.'),
    ) + (() if i.neoprene_only else (('Manual de Puentes MTC 2018, 2.4.3.11.8 y 2.8.1.4', 'Conexiones sísmicas y resistencia local del concreto.'),)), widths=(73, 86))
    _enforce_uniform_typography(doc)
    # La bibliografía conserva tipografía y márgenes; el espaciado compacto
    # evita dejar una única referencia en una página adicional.
    for row in doc.tables[-1].rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(0)
    if document is None:
        doc.save(path)
    return path
