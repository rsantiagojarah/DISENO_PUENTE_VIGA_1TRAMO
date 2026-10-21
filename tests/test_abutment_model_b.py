"""Independent geometry checks and shared-workflow regression for model B."""

from dataclasses import replace
from math import hypot

import pytest

from bridge_design.domain.abutment import (
    AbutmentGeometryInputs, AbutmentInputs, AbutmentKeyInputs, AbutmentSoilInputs,
    _bar_detail_length_m, _heel_weight_breakdown,
    _stem_thickness_at_height_m, _stem_temperature_result, _upper_concrete_weight_and_arm,
    solve_abutment_design,
)
from bridge_design.domain.abutment_model_b import (
    back_face_height_m, back_face_length_m, concrete_regions, model_b_geometry,
    soil_regions, stem_faces_at_height,
)
from bridge_design.cli.yaml_inputs import abutment_inputs_from_yaml, abutment_yaml_template
from bridge_design.domain.wall_friction import _back_x
from bridge_design.reporting.abutment_docx_charts import abutment_stem_outline


@pytest.fixture
def inputs_b():
    return AbutmentInputs(geometry=model_b_geometry(), key=AbutmentKeyInputs(enabled=False))


def test_regions_match_hand_calculated_areas_and_centroids(inputs_b):
    regions = concrete_regions(inputs_b.geometry)
    assert [r.area_m2 for r in regions] == pytest.approx([.375, .260, .105, 1.200, 1.320, 5.170])
    assert [r.centroid_x_m for r in regions] == pytest.approx([2.225, 2.025, 2+.35/3, 1.85, 1.5, 2.35])
    assert [r.centroid_y_m for r in regions] == pytest.approx([6.25, 5.3, 5.1-.6/3, 3.1, 1.1+4.4/3, .55])
    soil = soil_regions(inputs_b.geometry)
    assert [r.area_m2 for r in soil] == pytest.approx([2.35*5.9, .105, .35*3.4, .6*.4**2/(2*4.4), 1.1*.4])
    assert [r.centroid_x_m for r in soil] == pytest.approx([3.525, 2+2*.35/3, 2.175, 1.1+(.6*.4/4.4)/3, .55])
    assert [r.centroid_y_m for r in soil] == pytest.approx([4.05, 4.5+.6/3, 2.8, 1.1+2*.4/3, 1.3])


@pytest.mark.parametrize("height", [5., 7., 10.])
@pytest.mark.parametrize("base_thickness", [.3, .9, 1.4])
def test_section_integrals_match_regions_and_continuous_outer_batter(height, base_thickness):
    g = replace(model_b_geometry(), retained_height_m=height, lower_stem_thickness_m=base_thickness)
    data = AbutmentInputs(geometry=g)
    h = g.seat_body_height_m
    boundaries = [0., h-g.backwall_drop_m-g.backwall_taper_height_m,
                  h-g.backwall_drop_m, h, g.stem_height_above_footing_m]
    area = first_x = first_y = 0.0
    for lo, hi in zip(boundaries, boundaries[1:]):
        def sample(y):
            front, back = stem_faces_at_height(g, y)
            return back-front, (back**2-front**2)/2, (back-front)*(y+g.footing_thickness_m)
        # Simpson is exact for these piecewise linear faces and their moments.
        samples = [sample(lo+1e-10), sample((lo+hi)/2), sample(hi-1e-10)]
        integrals = [(hi-lo)/6 * (samples[0][j]+4*samples[1][j]+samples[2][j]) for j in range(3)]
        area += integrals[0]
        first_x += integrals[1]
        first_y += integrals[2]
    stem = concrete_regions(g)[:-1]
    assert area == pytest.approx(sum(r.area_m2 for r in stem), abs=1e-8)
    assert first_x == pytest.approx(sum(r.area_m2*r.centroid_x_m for r in stem), abs=1e-8)
    assert first_y == pytest.approx(sum(r.area_m2*r.centroid_y_m for r in stem), abs=1e-8)
    for y in [0., h/3, h/2, h-.01]:
        expected_front = g.toe_length_m+(base_thickness-g.upper_stem_thickness_m)*y/h
        front, back = stem_faces_at_height(g, y)
        assert front == pytest.approx(expected_front)
        assert _back_x(g, y) == back
        assert _stem_thickness_at_height_m(data, y) == pytest.approx(back-front)


def test_pdf_regions_fill_rear_envelope_without_gaps_or_overlap(inputs_b):
    g = inputs_b.geometry
    # Everything behind the fixed shoulder, from footing top to the crown,
    # is accounted for once as concrete or soil. Add the outer batter separately.
    concrete = concrete_regions(g)
    soil = soil_regions(g)
    rear_total = sum(r.area_m2 for r in concrete[:4]) + sum(r.area_m2 for r in soil[:3])
    shoulder = g.toe_length_m+g.lower_stem_thickness_m-g.upper_stem_thickness_m
    # The empty region above the bearing seat is on the exposed side of parapet.
    empty_seat = g.bearing_seat_length_m*g.seat_block_height_m
    assert rear_total + empty_seat == pytest.approx((g.footing_width_m-shoulder)*g.stem_height_above_footing_m)


def test_shared_solver_uses_model_b_masses_inertia_and_seat_position(inputs_b):
    r = solve_abutment_design(inputs_b)
    assert r.dc_self_weight_tn_m == pytest.approx(20.232)
    assert r.ev_weight_tn_m == pytest.approx(30.051)
    assert inputs_b.geometry.superstructure_load_x_m == pytest.approx(1.9)
    assert r.pressures.pir_tn_m == pytest.approx(.18*(20.232+30.051))
    assert r.pressures.lsy_tn_m == pytest.approx(2.35*.6*1.925)
    weight, arm = _upper_concrete_weight_and_arm(inputs_b)
    stem = concrete_regions(inputs_b.geometry)[:-1]
    assert weight == pytest.approx(7.824)
    assert arm == pytest.approx(sum(x.area_m2*(x.centroid_y_m-1.1) for x in stem)/sum(x.area_m2 for x in stem))
    seat_load = next(c for c in r.components.vertical_with_bridge if c.name == "PDC tablero")
    assert seat_load.arm_m == pytest.approx(1.9)
    assert len(r.extreme_seismic_with_bridge) == len(r.extreme_seismic_without_bridge) == 2
    assert r.stem_design.required_as_cm2_m > 0
    assert r.heel_design.required_as_cm2_m > 0
    assert r.toe_design.required_as_cm2_m > 0
    assert r.crack_checks and r.development_checks and r.secondary_reinforcement


def test_heel_free_body_uses_same_regions_as_global_stability(inputs_b):
    g = inputs_b.geometry
    heel = _heel_weight_breakdown(inputs_b)
    soil = soil_regions(g)[:3]
    assert heel.soil_weight_tn_m == pytest.approx(sum(r.area_m2 for r in soil)*1.925)
    assert heel.soil_weight_tn_m*heel.soil_arm_m == pytest.approx(
        sum(r.area_m2*(r.centroid_x_m-2.) for r in soil)*1.925)
    assert heel.seat_concrete_weight_tn_m == pytest.approx((.35*(.4+.6/2)+.25*1.5)*2.4)


def test_bar_path_and_cut_lengths_include_posterior_transition(inputs_b):
    g = inputs_b.geometry
    height = g.stem_height_above_footing_m
    expected = height-.6+hypot(.6, .35)
    assert back_face_length_m(g, height) == pytest.approx(expected)
    for y in [0., 2., 3.5, 3.9, 4.4, height]:
        assert back_face_height_m(g, back_face_length_m(g, y)) == pytest.approx(y, abs=1e-12)
    assert _bar_detail_length_m(inputs_b, "Pantalla", 50.) == pytest.approx(expected+.5)
    cut = solve_abutment_design(inputs_b).stem_reinforcement_cut
    assert cut is not None
    assert cut.continuous_bar_length_m == pytest.approx(expected+cut.footing_anchorage_m)
    assert cut.lower_cut_bar_length_m == pytest.approx(back_face_length_m(g, cut.constructive_cut_height_m)+cut.footing_anchorage_m)


def test_temperature_uses_actual_average_panel_thickness(inputs_b):
    result = _stem_temperature_result(inputs_b)
    # Panel below the posterior transition: h=3.4 m. The front batter runs
    # over 4.4 m, so its upper thickness is 0.9-0.6*3.4/4.4, not t_sup.
    assert result.b_cm == pytest.approx((.9 + (.9-.6*3.4/4.4))*50)
    assert result.h_cm == pytest.approx(340.)


def test_yaml_round_trip_default_A_legacy_and_B_derived_N():
    assert abutment_inputs_from_yaml({}).geometry.abutment_type == "A"
    template = abutment_yaml_template(geometry_defaults=model_b_geometry())
    assert template["geometria"]["tipo_estribo"] == "B"
    assert "transicion_frontal_superior_t1_m" not in template["geometria"]
    assert abutment_inputs_from_yaml(template).geometry == model_b_geometry()
    template["geometria"].pop("longitud_cajuela_m")
    template["geometria"]["espesor_superior_pantalla_m"] = .4
    assert abutment_inputs_from_yaml(template).geometry.bearing_seat_length_m == pytest.approx(.5)
    assert abutment_inputs_from_yaml({"geometria": {"tipo_estribo": "b"}}).geometry.abutment_type == "B"


@pytest.mark.parametrize("changes,match", [
    ({"bearing_seat_length_m": .9}, "N.*b_parapeto"),
    ({"small_batter_width_m": .1}, "t1"),
    ({"front_soil_depth_m": 6.}, "suelo frontal"),
    ({"backwall_taper_height_m": 0.}, "transicion positiva"),
    ({"seat_wall_width_m": 0., "bearing_seat_length_m": .65}, "parapeto"),
    ({"abutment_type": "C"}, "tipo_estribo"),
])
def test_invalid_model_b_geometry_rejected(changes, match):
    with pytest.raises(ValueError, match=match):
        replace(model_b_geometry(), **changes)


def test_model_b_unsupported_slope_wall_and_connected_are_explicit(inputs_b):
    from bridge_design.domain.connected_inputs import ConnectedInputs, FoundationSoil
    with pytest.raises(ValueError, match="relleno horizontal"):
        replace(inputs_b, soil=AbutmentSoilInputs(backfill_slope_deg=5))
    with pytest.raises(ValueError, match="diseno-muros"):
        replace(inputs_b, is_pure_wall=True)
    with pytest.raises(ValueError, match="diseno-estribos"):
        abutment_inputs_from_yaml({"geometria": {"tipo_estribo": "B"}}, pure_wall=True)
    with pytest.raises(ValueError, match="FRAME conectado requiere modelo A"):
        ConnectedInputs(soil=FoundationSoil(3000., .5, 26.7), left=inputs_b, right=inputs_b)


def test_model_b_without_toe_or_front_fill():
    data = AbutmentInputs(geometry=replace(model_b_geometry(), toe_length_m=0., front_soil_depth_m=0.))
    result = solve_abutment_design(data)
    assert result.toe_design is None
    assert [c.name for c in result.soil_components] == ["7 Relleno posterior", "8 Relleno triangular t2", "9 Relleno junto pantalla"]


def test_model_b_friction_and_adopted_reinforcement_use_shared_checks(inputs_b):
    data = replace(inputs_b, soil=replace(inputs_b.soil, wall_soil_friction_deg=10.))
    preliminary = solve_abutment_design(data)
    options = preliminary.stem_design.spacing_options.options
    selected = options[-1]
    result = solve_abutment_design(data, {"Pantalla": selected})
    assert result.stem_design.selected_bar_label == selected.bar.label
    assert result.stem_design.selected_spacing_m == selected.spacing_m
    assert result.stem_design.provided_as_cm2_m == pytest.approx(selected.provided_area_cm2_m)
    assert result.pressures.stem_force_angle_deg == 10.


def test_model_b_interactive_geometry_calculates_N(monkeypatch, capsys):
    from bridge_design.cli.abutment_input_prompts import _collect_geometry
    monkeypatch.setattr("builtins.input", lambda prompt: "B" if "Tipo de estribo" in prompt else "")
    result = _collect_geometry()
    assert result == model_b_geometry()
    assert "N - Longitud horizontal del asiento calculada: 0.400" in capsys.readouterr().out


def test_model_b_template_cli_and_help(tmp_path, capsys):
    from bridge_design.abutment_main import main
    from bridge_design.cli.yaml_io import load_yaml_file
    target = tmp_path / "modelo_b.yaml"
    main(["output", str(target), "--tipo", "B"])
    data = abutment_inputs_from_yaml(load_yaml_file(target))
    assert data.geometry.abutment_type == "B"
    main(["--help"])
    assert "--tipo A|B" in capsys.readouterr().out


def test_model_b_yaml_runs_complete_selection_flow(monkeypatch, capsys):
    from pathlib import Path
    from bridge_design import abutment_main as workflow
    from bridge_design.cli.yaml_io import load_yaml_file
    prompts, results = [], []

    def answer(prompt):
        prompts.append(prompt)
        if len(prompts) > 60:
            raise AssertionError("El flujo no converge con las propuestas de consola.")
        return ""

    monkeypatch.setattr("builtins.input", answer)
    monkeypatch.setattr(workflow, "generate_abutment_docx_with_dialog", results.append)
    data = abutment_inputs_from_yaml(load_yaml_file(Path("modelo_estribo_B.yaml")))
    workflow.run_abutment_design(inputs=data)
    assert len(results) == 1
    result = results[0]
    assert result.inputs.geometry.abutment_type == "B"
    assert result.inputs.geometry.footing_width_m >= data.geometry.footing_width_m
    assert all(state.overturning_status == "OK" and state.bearing_status == "OK"
               and (state.sliding_with_key_status or state.sliding_status) == "OK"
               for state in result.with_bridge+result.without_bridge)
    assert "Modelo B" in capsys.readouterr().out


def test_model_b_console_and_drawing_describe_same_section(inputs_b, tmp_path):
    from bridge_design.cli.abutment_ascii_output import format_abutment_design_result
    from bridge_design.reporting.abutment_docx_charts import save_abutment_geometry
    from bridge_design.reporting.abutment_docx_detail import concrete_geometry_rows
    result = solve_abutment_design(inputs_b)
    text = format_abutment_design_result(result)
    assert "Modelo B" in text and "5 Ensanche exterior continuo" in text
    assert "11 Relleno frontal" in text
    assert "armado local requiere diseno especifico" in text
    outline = abutment_stem_outline(inputs_b)
    # At the head top the exposed face meets the single continuous batter.
    seat_y = inputs_b.geometry.retained_height_m-inputs_b.geometry.seat_block_height_m
    seat_xs = [x for x, y in outline if abs(y-seat_y) < 1e-5]
    assert min(seat_xs) == pytest.approx(1.7, abs=1e-6)
    expressions = concrete_geometry_rows(result)
    assert expressions[4][1].startswith("(t_inf-t_sup)")
    assert save_abutment_geometry(result, tmp_path / "b.png").stat().st_size > 10000


def test_model_b_word_identifies_geometry_and_reports_its_scope(inputs_b, tmp_path):
    from docx import Document
    from bridge_design.reporting.abutment_docx import generate_abutment_docx
    path = generate_abutment_docx(solve_abutment_design(inputs_b), tmp_path / "b.docx")
    document = Document(path)
    text = "\n".join(p.text for p in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "Modelo B con cara exterior inclinada continua" in text
    assert "su armado local y la zona de apoyo requieren diseno especifico" in text
    assert "No se ha incluido dentellón" in text
    assert "fricción de base satisface el deslizamiento en todos los estados" not in text
    assert "5 Ensanche exterior continuo" in table_text and "6 Zapata" in table_text
    assert "11 Relleno frontal" in table_text


def test_model_A_baseline_unchanged():
    result = solve_abutment_design(AbutmentInputs(geometry=AbutmentGeometryInputs()))
    assert result.dc_self_weight_tn_m == pytest.approx(20.016, abs=1e-9)
    assert result.ev_weight_tn_m == pytest.approx(30.057176470588246, abs=1e-9)
    assert result.with_bridge[0].vu_tn_m == pytest.approx(60.04157647058824, abs=1e-9)
