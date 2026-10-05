from copy import deepcopy
from dataclasses import replace
import json

from docx import Document
import pytest

from bridge_design.cli.connected_2_prompts import collect_connected_2_inputs
from bridge_design.cli.connected_2_yaml import (
    COMMAND, GEOMETRY_FIELDS, connected_2_inputs_from_yaml, connected_2_yaml_template,
)
from bridge_design.cli.connected_yaml import connected_inputs_from_yaml, connected_yaml_template
from bridge_design.cli.connected_selection import format_connected_options
from bridge_design.cli.yaml_io import load_yaml_file
from bridge_design.connected_2_main import main
from bridge_design.domain.abutment import _stem_thickness_at_height_m
from bridge_design.domain.connected_2_inputs import Connected2Geometry, Connected2Inputs
from bridge_design.domain.connected_defaults import connected_geometry_defaults, connected_load_defaults
from bridge_design.domain.connected_design import solve_connected_abutments
from bridge_design.domain.connected_geometry import build_connected_mesh, global_x
from bridge_design.domain.connected_loads import gravity_builders, soil_density
from bridge_design.domain.connected_options import connected_reinforcement_options
from bridge_design.domain.wall_friction import _back_x
from bridge_design.reporting.connected_chart_geometry import structure_polygons


def model():
    return connected_2_inputs_from_yaml(connected_2_yaml_template(example=True))


def test_pdf_dimensions_and_actual_profile():
    data = model()
    g = data.left.geometry
    assert isinstance(data, Connected2Inputs)
    assert data.left == data.right
    assert data.left.loads == connected_load_defaults()
    assert data.total_length_m == pytest.approx(17.40)
    assert data.upper_clear_span_m == pytest.approx(14.00)
    assert data.clear_span_m == pytest.approx(12.90)
    assert g.retained_height_m == pytest.approx(10.70)
    assert g.stem_height_above_footing_m == pytest.approx(9.20)
    assert g.seat_block_height_m == 1.17
    assert g.seat_wall_width_m == .25
    assert g.bearing_seat_length_m == .55
    assert g.upper_stem_thickness_m == .8
    assert g.lower_stem_thickness_m == 1.35
    assert g.heel_length_m == pytest.approx(.9)
    assert g.toe_length_m == 0
    assert g.front_soil_depth_m == g.footing_thickness_m == data.slab_thickness_m == 1.5
    for y, expected in ((0, 1.35), (4.015, 1.075), (8.03-1e-7, .8), (8.5, .25)):
        assert _stem_thickness_at_height_m(data.left, y) == pytest.approx(expected)
        assert global_x(data, 0, _back_x(g, y)) == pytest.approx(.9)
    assert global_x(data, 0, g.superstructure_load_x_m) == pytest.approx(1.425)
    assert global_x(data, 1, g.superstructure_load_x_m) == pytest.approx(15.975)


def test_continuous_base_mesh_and_drawn_concrete():
    data = model()
    mesh = build_connected_mesh(data)
    assert len(mesh.frame.springs) == 41
    assert sum(s.tributary_area for s in mesh.frame.springs) == pytest.approx(17.4)
    assert mesh.frame.nodes[mesh.reference_node].x == pytest.approx(8.7)
    for element in mesh.frame.elements[:mesh.foundation_count-1]:
        assert element.region == "Zapata combinada"
        assert element.area == 1.5
        assert element.inertia == pytest.approx(1.5**3/12)
    assert {mesh.frame.elements[i].region for i in mesh.side_elements[0]} == {"Pantalla izquierda", "Parapeto izquierda"}
    assert not any(e.region.startswith(("Cajuela", "Transicion")) for e in mesh.frame.elements)
    floor, left, right = structure_polygons(data, mesh)
    assert min(y for _, y in floor) == -1.5
    assert max(x for x, _ in floor) == pytest.approx(17.4)
    assert min(x for x, _ in left) == pytest.approx(.9)
    assert max(x for x, _ in right) == pytest.approx(16.5)


def test_load_integration_matches_independent_areas_and_no_interior_soil():
    data = model()
    parts = gravity_builders(data, build_connected_mesh(data))
    concrete = -sum(sum(parts[k].vector()[1::3]) for k in ("DC0", "DC1", "DCs"))
    soil = -sum(sum(parts[k].vector()[1::3]) for k in ("EV0", "EV1"))
    expected_concrete = (17.4*1.5 + 2*((1.35+.8)/2*8.03 + .25*1.17))*2.4
    assert concrete == pytest.approx(expected_concrete)
    assert soil == pytest.approx(2*.9*(8.03+1.17)*1.925)
    assert soil_density(data.left, .4) == (0, 0)
    assert soil_density(data.left, 1.8) == pytest.approx((9.2, 4.6))


@pytest.mark.parametrize("key,attr,_label", GEOMETRY_FIELDS)
def test_every_independent_dimension_is_editable(key, attr, _label):
    raw = connected_2_yaml_template(example=True)
    before = model()
    old = raw["estribo"]["geometria"][key]
    raw["estribo"]["geometria"][key] = old * 1.1
    after = connected_2_inputs_from_yaml(raw)
    assert after.left.geometry != before.left.geometry or after.clear_span_m != before.clear_span_m
    assert after.left.geometry == after.right.geometry
    assert raw["estribo"]["geometria"][key] == old * 1.1


@pytest.mark.parametrize("values", [
    {"foundation_thickness_m": 0}, {"stem_body_height_m": float("nan")},
    {"outer_heel_m": -1}, {"seat_length_m": True},
    {"stem_base_thickness_m": .7}, {"upper_clear_span_m": 1.0},
])
def test_impossible_geometry_is_rejected(values):
    with pytest.raises(ValueError):
        Connected2Geometry(**values)


def test_new_schema_cannot_silently_load_old_geometry_or_duplicate_dimensions():
    with pytest.raises(ValueError, match="no corresponde"):
        connected_2_inputs_from_yaml(connected_yaml_template(example=True))
    with pytest.raises(ValueError, match="no corresponde"):
        connected_inputs_from_yaml(connected_2_yaml_template(example=True))
    raw = connected_2_yaml_template(example=True)
    raw["estribo"]["geometria"]["longitud_puntera_m"] = 1
    with pytest.raises(ValueError, match="Dimensiones no admitidas"):
        connected_2_inputs_from_yaml(raw)
    raw = connected_2_yaml_template(example=True)
    raw["cimentacion"]["espesor_losa_central_m"] = .75
    with pytest.raises(ValueError, match="una vez"):
        connected_2_inputs_from_yaml(raw)


def test_yaml_overrides_cases_and_shared_validation_survive_adaptation():
    raw = connected_2_yaml_template(example=True)
    raw["cargas_tablero_derecho"] = {"pdc_carga_muerta_tablero_tn_m": 15}
    loads = raw["estribo"]["cargas_tablero"]
    raw["casos_simultaneos"] = [{"nombre": "posicion alternativa", "cargas_izquierda": loads,
                               "cargas_derecha": {**loads, "br_frenado_tn_m": 0}}]
    original = deepcopy(raw)
    data = connected_2_inputs_from_yaml(raw)
    assert raw == original
    assert data.left.loads.pdc_tn_m == 8.959
    assert data.right.loads.pdc_tn_m == 15
    assert data.cases[0].right.braking_tn_m == 0
    assert replace(data, foundation_node_count=81).geometry_model == "ver_2l"
    with pytest.raises(ValueError, match="uniforme"):
        replace(data, slab_thickness_m=.75)
    with pytest.raises(ValueError, match="modulo_balasto"):
        connected_2_inputs_from_yaml(connected_2_yaml_template())


@pytest.mark.parametrize("same_right", [True, False])
def test_fast_console_defaults_and_editing_without_irrelevant_questions(monkeypatch, capsys, same_right):
    prompts = []
    def answer(prompt):
        prompts.append(prompt)
        if prompt.startswith("Modulo de balasto"):
            return "3000"
        if prompt.startswith("Usar las mismas reacciones"):
            return "s" if same_right else "n"
        if prompt.startswith("PDC") and sum(p.startswith("PDC") for p in prompts) == 2:
            return "11"
        return ""
    monkeypatch.setattr("builtins.input", answer)
    data = collect_connected_2_inputs()
    assert data.left.geometry == model().left.geometry
    assert data.total_length_m == pytest.approx(17.4)
    assert data.right.loads.pdc_tn_m == (8.959 if same_right else 11)
    assert len([p for p in prompts if p.startswith("Espesor uniforme")]) == 1
    assert not any(any(term in p.lower() for term in ("puntera", "transicion", "theta", "suelo frontal", "h activo", "ancho total")) for p in prompts)
    assert "luz inferior=12.900" in capsys.readouterr().out


def test_analysis_equilibrium_common_steel_and_original_defaults():
    result = solve_connected_abutments(model(), check_mesh=True)
    assert len(result.results) == 17
    assert len(result.reinforcement) == 12
    assert not any("puntera" in steel.region for steel in result.reinforcement)
    options = format_connected_options(connected_reinforcement_options(result))
    assert "OPCIONES - Zapata combinada - longitudinal superior" in options
    assert "OPCIONES - Zapata combinada - longitudinal inferior" in options
    assert "losa" not in options.lower()
    for case in result.results:
        assert max(abs(v) for v in case.equilibrium_error) < 1e-6
    assert result.mesh_comparison.fine_nodes == 81
    assert result.mesh_comparison.pressure_change < .01
    original = connected_inputs_from_yaml(connected_yaml_template(example=True))
    assert original.left.geometry == connected_geometry_defaults()
    assert original.clear_span_m == 13.1
    assert original.slab_thickness_m == .75
    assert original.left_transition_m == original.right_transition_m == 1
    assert original.total_length_m == pytest.approx(18.1)


def test_command_generates_complete_separate_outputs(tmp_path, capsys):
    template = tmp_path / "modelo2.yaml"
    main(["output", str(template), "--ejemplo"])
    assert load_yaml_file(template)["comando"] == COMMAND
    destination = tmp_path / "resultados2"
    main(["input", str(template), "--automatico", "--resultados", str(destination)])
    saved = json.loads((destination / "resultados.json").read_text(encoding="utf-8"))
    assert saved["inputs"]["geometry_model"] == "ver_2l"
    assert len(saved["results"]) == 17
    assert len(saved["reinforcement"]) == 12
    assert "losa central" not in (destination / "resultados.json").read_text(encoding="utf-8").lower()
    for name in ("resumen.txt", "auditoria.txt", "elementos.csv", "esfuerzos.csv"):
        assert "losa central" not in (destination / name).read_text(encoding="utf-8-sig").lower()
    for name in ("geometria.png", "esfuerzos.csv", "nodos.csv", "memoria_estribos_conectados_2.docx"):
        assert (destination / name).exists()
    for name in ("cortes_zapata.csv", "cortes_zapata.png"):
        assert (destination / name).exists()
    assert not (destination / "memoria_estribos_conectados.docx").exists()
    doc = Document(destination / "memoria_estribos_conectados_2.docx")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "estribos conectados 2" in text
    assert "no existe relleno interior" in text
    assert "Cortes del acero de zapata - superior" in text
    assert "Cortes del acero de zapata - inferior" in text
    assert "losa" not in text.lower()
    assert "losa" not in " ".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells).lower()
    assert not any("puntera" in p.text for p in doc.paragraphs if p.style.name.startswith("Heading"))
    output = capsys.readouterr().out
    assert "BASE UNIFORME" in output
    assert "Zapata combinada - longitudinal superior" in output
    assert "Zapata combinada - longitudinal inferior" in output
    assert "CORTES DE ACERO - Zapata combinada - longitudinal superior" in output
    assert "CORTES DE ACERO - Zapata combinada - longitudinal inferior" in output
    assert "losa" not in output.lower()
