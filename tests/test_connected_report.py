from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from bridge_design.domain.connected_design import solve_connected_abutments
from bridge_design.domain.connected_inputs import ConnectedInputs, FoundationSoil
from bridge_design.reporting.connected_charts import save_connected_charts
from bridge_design.reporting.connected_docx import write_connected_docx
from bridge_design.reporting.connected_case_groups import case_label


def test_connected_memory_preserves_inputs_results_and_correct_header(tmp_path):
    inputs = ConnectedInputs(FoundationSoil(3000, 0.50, 26.7), foundation_node_count=41,
                             include_without_bridge=False)
    result = solve_connected_abutments(inputs, check_mesh=True)
    charts = save_connected_charts(result, tmp_path)
    destination = write_connected_docx(result, tmp_path / "memoria.docx", charts)
    document = Document(destination)
    paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
    tables = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "ESTRIBOS CONECTADOS" in document.sections[0].header.paragraphs[0].text
    assert "TABLERO DE PUENTE" not in document.sections[0].header.paragraphs[0].text
    assert "Módulo de balasto vertical 3000.000 tn/m³" in paragraphs
    assert "Módulo de balasto vertical" not in tables
    assert "Longitud horizontal de cajuela" in tables
    assert "PENDIENTE DETALLE" in paragraphs
    assert case_label(result.results[-1].name) in paragraphs
    assert len(document.inline_shapes) == 29
    assert "Perfil del empuje estático en ambas caras reales sin factorizar" in paragraphs
    assert "Forma real del empuje sísmico base alternativa A sentido menos" in paragraphs
    assert "Los siguientes esquemas muestran las acciones iniciales sin factores LRFD" in paragraphs
    assert "Deformada nodal de servicio con amplificación indicada en la figura" not in paragraphs
    assert "6.5. Envolvente general de todas las combinaciones" in paragraphs
    assert "Envolvente general de momento M de todas las combinaciones" in paragraphs
    assert paragraphs.index("Envolvente general de momento M de todas las combinaciones") < paragraphs.index("7. Contacto")
    assert "Cargas aplicadas de la primera combinación" not in paragraphs
    assert "Esquema de armadura por cara" not in paragraphs
    import re
    assert re.search(r"\bC\d{3}\b",paragraphs+tables) is None
    assert "Capacidad portante y presión límite factorizada" in paragraphs
    assert paragraphs.count("Capacidad portante y presión límite factorizada")==4
    assert "q máximo de la envolvente=" in paragraphs
    assert "Presiones y asentamientos de servicio en la cimentación completa" not in paragraphs
    assert "deslizamiento" not in (paragraphs + tables).lower()
    assert "Índice desl" not in tables
    assert "Nudo de presión máxima" not in tables
    assert "Meyerhof" not in paragraphs
    assert "6.1. Envolvente Resistencia Ia" in paragraphs
    assert "6.4. Envolvente Evento Extremo I" in paragraphs
    from bridge_design.reporting.connected_docx_layout import report_reinforcement
    reported_steel = report_reinforcement(result)
    assert len(document.tables) == 8 + 3*len({s.base_region for s in reported_steel})
    design_headings = [p.text for p in document.paragraphs if p.style.name.startswith('Heading')]
    assert [h for h in design_headings if h.startswith('8.') and h != '8. Diseño del concreto armado por región'] == [
        '8.1. Pantalla', '8.2. Zapata', '8.3. Losa central', '8.4. Parapeto']
    assert len(reported_steel) == 16
    assert 'único armado de estribo' in paragraphs
    assert 'Pantalla izquierda;' in paragraphs or 'Pantalla derecha;' in paragraphs
    assert all('izquierda' not in s.region and 'derecha' not in s.region for s in reported_steel)
    assert not any('Transicion' in heading or 'Cajuela' in heading for heading in design_headings)
    assert not any(row.cells[0].text.startswith(('Transicion', 'Cajuela')) for row in document.tables[-1].rows)
    assert document.tables[0].cell(0, 0).text == "Dato"
    assert document.tables[1].cell(0, 0).text == "Sección"
    assert "Volumen m³" not in tables and "W·x tn·m" not in tables
    assert "descomposicion geometrica" not in paragraphs
    assert "Ejemplo nodal" not in paragraphs
    assert "4.4. Empuje e inercias sísmicas" in paragraphs
    assert "KAE EQ−" in tables and "KAE EQ+" in tables
    assert "4.6. Frenado y traslado de la fuerza al modelo" in paragraphs
    assert "La norma refiere los 1.80 m a la superficie de calzada" in paragraphs
    assert "kh=0.5·As" in paragraphs
    assert "FS·qadm no identifica automáticamente la capacidad última" in paragraphs
    assert "Relleno" in tables and "Exterior" in tables
    assert "Superior" in tables and "Inferior" in tables
    assert "8.1. Envolvente" not in paragraphs
    equations = " ".join(document.element.xpath(".//m:t/text()"))
    assert "gamma" not in equations and "phi" not in equations and "eta" not in equations
    assert "Envolvente axial N sobre la estructura completa" in paragraphs
    assert "Envolvente de cortante V sobre la estructura completa" in paragraphs
    assert "Envolvente de momento M sobre la estructura completa" in paragraphs
    assert "sin verificacion de interaccion axial-momento" in paragraphs
    assert "Sección gobernante de flexión" in paragraphs
    assert "Interacción N M" not in tables
    assert "Secciones gobernantes independientes" in paragraphs
    assert "As requerido flexión y mínimos=" in paragraphs
    assert "Peralte y acero de la seccion gobernante de flexion" in paragraphs
    assert "Verificación de cortante" in paragraphs
    assert "Cortante y procedimiento simplificado" not in paragraphs
    assert "Fisuracion y tension del acero en servicio" in paragraphs
    assert "Propiedades FRAME por elemento" not in paragraphs
    assert "Accion / lado" not in tables
    assert "Fx base Tn" not in tables
    assert "contacto, presion y asentamiento" not in paragraphs
    assert "K=ks*Area" not in paragraphs
    assert "Factores incluyen sentido" not in paragraphs
    assert "se aplican también cruzadas" not in paragraphs
    assert "se aplica globalmente a ambos estribos, losa y transiciones" in paragraphs
    assert "Asentamientos sin límite admisible ingresado" in paragraphs
    assert "cargas_combinadas.csv" in paragraphs
    assert "Envolventes N V M del estribo" not in paragraphs
    section = document.sections[0]
    assert round(section.page_width.mm) == 210
    assert round(section.page_height.mm) == 297
    assert abs(section.left_margin.mm - 25.4) < 0.05
    title = next(paragraph for paragraph in document.paragraphs if paragraph.style.name == "Title")
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.style.element.get_or_add_pPr().find(qn("w:pBdr")) is None
    assert all(run.font.size.pt == 22 for run in title.runs)
    assert "Contenido" in paragraphs
    assert "11. Resumen y conclusiones" in paragraphs
    assert "12. Referencias y anexos de trazabilidad" in paragraphs
    from bridge_design.reporting.connected_docx_layout import SECTIONS
    assert [p.text for p in document.paragraphs if p.style.name == "Heading 1" and p.text != "Contenido"] == [
        f"{i}. {label}" for i, label in enumerate(SECTIONS, 1)]
    assert "Control de cobertura:" in paragraphs
    assert "Se solicitaron y se colocaron 41 nudos con resorte" in paragraphs
    assert "4.3. Empuje estático y transferencia al talón" in paragraphs
    assert "z_BR=11.000 m, z_nudo=8.050 m y brazo=2.950 m" in paragraphs
    assert "NO CUMPLE — Pantalla - vertical relleno:" in paragraphs
    assert "Refinamiento de 41 a 81 resortes" in paragraphs
    assert "Reemplazando los valores correspondientes:" in paragraphs
    # Every adopted principal distribution develops its own governing section,
    # rather than printing one shared formula followed by compressed values.
    principal = [s for s in reported_steel if s.role == 'primary']
    assert paragraphs.count('Resistencia a flexion del acero elegido') == len(principal)
    assert paragraphs.count('Peralte y acero de la seccion gobernante de flexion') == len(principal)
    assert paragraphs.count('Desarrollo recto y gancho') == len(principal)
    assert paragraphs.count('Donde:') == paragraphs.count('Reemplazando los valores correspondientes:')
    assert 'se desarrolla con expresión, definición de símbolos, sustitución numérica y conclusión' in paragraphs
    assert 'Procedimiento general: Mv=' not in paragraphs
    normalized_math = equations.replace(' ', '')
    assert 'd=100.00-7.50-(2.54/2)=91.23cm' in normalized_math
    assert 'η_M=' in normalized_math and 'η_s=' in normalized_math
    assert paragraphs.count('Verificación de cortante') == len(principal)
    assert paragraphs.count('Sección gobernante de cortante:') == len(principal)
    assert 'As,prov=' in normalized_math and 'ag=0.75in' in normalized_math
    assert 'k_z' not in equations and 'k: multiplicador' not in paragraphs
    assert 'Md=max(abs(Mu),min(Mcr,1.33*abs(Mu)))' in normalized_math
    assert 'Md=max(1.33*abs(Mu),min(Mcr,1.33*abs(Mu)))' not in normalized_math
    assert 'amplificación adicional' not in paragraphs
    anchor_equations = [p for p in document.paragraphs
                        if ''.join(p._p.xpath('.//m:t/text()')).strip().startswith('ldb =')]
    assert len(anchor_equations) == 2*len(principal)
    for paragraph in anchor_equations:
        fraction = paragraph._p.xpath('.//m:f')[0]
        assert '2.54' in ''.join(fraction.find(qn('m:num')).itertext())
        assert '2.54' not in ''.join(fraction.find(qn('m:den')).itertext())
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Title":
            continue
        if paragraph.text.strip():
            assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert paragraph.paragraph_format.line_spacing == 1.5
        assert all(run.font.name == "Arial Narrow" and run.font.size.pt == 11 for run in paragraph.runs)
    assert all(table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None for table in document.tables)
    assert any(table.cell(0, 0).text == "Acción" and len(table.columns) == 5 for table in document.tables)
    assert all(steel.area_per_face_cm2_m >= steel.temperature_cm2_m for steel in result.reinforcement)
    failed = [steel for steel in result.reinforcement if steel.status != "OK"]
    assert failed
    assert "NO CUMPLE" in tables
    assert all(max(steel.flexural_utilization, steel.shear_utilization,
                   steel.crack_utilization) > 1 for steel in failed)
