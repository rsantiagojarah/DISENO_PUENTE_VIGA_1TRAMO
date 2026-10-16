"""Wall soil geometry and rejected cut details must reach every output."""

from dataclasses import replace
from io import BytesIO
from math import radians, tan

import pytest
from docx import Document

from bridge_design.cli.abutment_ascii_output import (
    _format_stem_reinforcement_cut,
    _format_stem_reinforcement_cut_option,
)
from bridge_design.domain import abutment as a
from bridge_design.domain.rebar_catalog import custom_spacing_option
from bridge_design.reporting import abutment_docx as report


def wall_inputs():
    base = a.pure_wall_inputs()
    return replace(
        base,
        geometry=replace(
            base.geometry, retained_height_m=8.3, footing_width_m=6.,
            footing_thickness_m=1., toe_length_m=2.,
            lower_stem_thickness_m=1., upper_stem_thickness_m=.5,
            front_soil_depth_m=1.,
        ),
        key=replace(base.key, enabled=False),
    )


@pytest.mark.parametrize("theta", [90., 89.999999, 88.])
@pytest.mark.parametrize("front_depth", [1., 2., 3.])
def test_front_soil_wedge_has_geometric_area_and_centroid(theta, front_depth):
    base = wall_inputs()
    inputs = replace(base, geometry=replace(base.geometry, front_soil_depth_m=front_depth),
                     soil=replace(base.soil, wall_backface_angle_deg=theta))
    g = inputs.geometry
    height = front_depth - g.footing_thickness_m
    setback = g.stem_height_above_footing_m * tan(radians(90. - theta))
    width = (g.lower_stem_thickness_m - g.upper_stem_thickness_m - setback) * height / g.stem_height_above_footing_m
    wedges = [c for c in a._soil_components(inputs) if c.name == "Relleno triangular frontal"]
    if height == 0.:
        assert wedges == []
        return
    assert len(wedges) == 1
    wedge = wedges[0]
    assert wedge.value_tn_m == pytest.approx(.5 * width * height * 1.925)
    assert wedge.arm_m == pytest.approx(g.toe_length_m + width / 3)
    assert wedge.vertical_arm_m == pytest.approx(g.footing_thickness_m + 2 * height / 3)


def test_front_soil_weight_and_stability_are_continuous_at_vertical_backface():
    base = wall_inputs()
    inputs = replace(base, geometry=replace(base.geometry, front_soil_depth_m=3.))
    vertical = a.solve_abutment_design(inputs)
    almost = a.solve_abutment_design(replace(inputs, soil=replace(inputs.soil, wall_backface_angle_deg=89.999999)))
    expected_wedge = 1.925 * .5 * (.5 / 7.3) * 2**2
    assert vertical.ev_weight_tn_m == pytest.approx(42.1575 + 7.7 + expected_wedge)
    assert vertical.ev_weight_tn_m == pytest.approx(almost.ev_weight_tn_m, abs=1e-6)
    for left, right in zip(vertical.without_bridge, almost.without_bridge):
        assert left.vu_tn_m == pytest.approx(right.vu_tn_m, abs=2e-6)
        assert left.friction_resistance_tn_m == pytest.approx(right.friction_resistance_tn_m, abs=2e-6)
        assert left.eccentricity_m == pytest.approx(right.eccentricity_m, abs=1e-6)


def test_word_weights_include_the_same_front_wedge_as_engine():
    base = wall_inputs()
    result = a.solve_abutment_design(replace(base, geometry=replace(base.geometry, front_soil_depth_m=3.)))
    document = Document()
    report._configure_document(document)
    report._weights(document, result, report._labels(is_pure_wall=True))
    matches = [row for table in document.tables for row in table.rows
               if row.cells[0].text == "Relleno triangular frontal"]
    assert len(matches) == 1
    assert matches[0].cells[3].text == "0.264"
    assert matches[0].cells[4].text == "2.046"


@pytest.fixture(scope="module")
def rejected_cut_result():
    base = wall_inputs()
    inputs = replace(base, geometry=replace(
        base.geometry, retained_height_m=1.2, footing_thickness_m=.5,
        lower_stem_thickness_m=.3, upper_stem_thickness_m=.2,
        footing_width_m=2., toe_length_m=.5, front_soil_depth_m=.5,
    ))
    preliminary = a.solve_abutment_design(inputs)
    selected = custom_spacing_option(preliminary.stem_design.spacing_options, '1"', .1)
    result = a.solve_abutment_design(inputs, {"Pantalla": selected})
    assert result.stem_reinforcement_cut.status == "NO CONVIENE"
    return result


@pytest.mark.parametrize("formatter", [_format_stem_reinforcement_cut, _format_stem_reinforcement_cut_option])
def test_terminal_rejected_cut_keeps_all_principal_bars_continuous(rejected_cut_result, formatter):
    output = "\n".join(formatter(rejected_cut_result))
    assert "NO CONVIENE" in output
    assert "Sin corte; todas las barras principales continuan" in output
    assert '1" @ 0.100 m' in output
    assert "Continua 1 de cada" not in output
    assert "Acero continuo superior" not in output


def test_word_rejected_cut_matches_continuous_bar_schedule(rejected_cut_result):
    result = rejected_cut_result
    document = Document()
    report._configure_document(document)
    report._service_and_detailing(document, result)
    report._summary(document, result)
    saved = BytesIO()
    document.save(saved)
    saved.seek(0)
    document = Document(saved)
    text = " ".join(document.element.xpath('//*[local-name()="t"]/text()'))
    assert "NO CONVIENE" in text
    assert "El corte no se adopta." in text
    assert "acero principal continuo en toda la altura" in text
    assert "Patrón constructivo adoptado" not in text
    assert "Superior continuo (1 de cada" not in text
    principal = [row for row in result.bar_details if row.element == "Pantalla" and row.anchorage_m > 0]
    assert len(principal) == 1
    assert principal[0].spacing_m == .1
    assert principal[0].length_m == pytest.approx(result.stem_reinforcement_cut.continuous_bar_length_m)


def test_valid_cut_keeps_its_adopted_pattern_in_word_and_terminal():
    result = a.solve_abutment_design(wall_inputs())
    assert result.stem_reinforcement_cut.status == "OK"
    document = Document()
    report._configure_document(document)
    report._service_and_detailing(document, result)
    text = " ".join(document.element.xpath('//*[local-name()="t"]/text()'))
    assert "Patrón constructivo adoptado: continúa 1 de cada 2" in text
    assert "Continua 1 de cada 2" in "\n".join(_format_stem_reinforcement_cut(result))
