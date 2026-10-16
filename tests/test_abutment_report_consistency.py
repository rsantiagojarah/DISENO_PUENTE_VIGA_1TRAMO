"""Keep bar schedules and Word developments consistent with adopted steel."""

from dataclasses import replace
import re

import pytest
from docx import Document

from bridge_design.domain import abutment
from bridge_design.domain.concrete_flexure import rectangular_flexural_response
from bridge_design.domain.rebar_catalog import custom_spacing_option, reinforcing_bar_by_label
from bridge_design.reporting import abutment_docx, abutment_docx_detail as detail


@pytest.fixture(scope="module")
def default_result():
    return abutment.solve_abutment_design()


@pytest.fixture(scope="module")
def wall_result():
    # Geometry of the reported 8.30 m wall; fixed here to keep the test independent
    # of subsequent edits to the user's YAML file.
    inputs = abutment.AbutmentInputs(
        geometry=abutment.AbutmentGeometryInputs(
            retained_height_m=8.30, footing_width_m=6.0, footing_thickness_m=1.0,
            toe_length_m=2.0, lower_stem_thickness_m=1.0, upper_stem_thickness_m=.5,
            front_soil_depth_m=1.0,
        ),
        is_pure_wall=True,
    )
    return abutment.solve_abutment_design(inputs)


def test_cut_families_preserve_selected_area_and_positions(wall_result):
    cut = wall_result.stem_reinforcement_cut
    assert cut is not None and cut.status == "OK"
    families = [row for row in wall_result.bar_details
                if row.element == "Pantalla" and row.anchorage_m > 0]
    assert len(families) == 2
    assert cut.lower_spacing_m == .1 and cut.upper_spacing_m == .2
    assert [row.spacing_m for row in families] == [.2, .2]
    area = sum(reinforcing_bar_by_label(row.bar_label).area_cm2 / row.spacing_m
               for row in families)
    assert area == pytest.approx(wall_result.stem_design.provided_as_cm2_m)
    assert area == pytest.approx(28.4)
    assert "intercalada a 0.100 m" in families[0].note
    assert families[0].length_m == cut.lower_cut_bar_length_m
    assert families[1].length_m == cut.continuous_bar_length_m


def test_saved_one_in_three_pattern_has_three_nonoverlapping_families(wall_result):
    cut = replace(wall_result.stem_reinforcement_cut, continuous_every_n_bars=3,
                  upper_spacing_m=.3, upper_provided_as_cm2_m=2.84/.3)
    check = next(c for c in wall_result.development_checks if c.element == "Pantalla")
    families = abutment._stem_cut_bar_details(wall_result.stem_design, check, cut)
    assert len(families) == 3
    assert [row.spacing_m for row in families] == [.3, .3, .3]
    assert sum(2.84 / row.spacing_m for row in families) == pytest.approx(28.4)
    assert "intercalada a 0.100 m" in families[0].note
    assert "intercalada a 0.200 m" in families[1].note


def test_word_cut_lengths_and_schedule_use_family_spacing(wall_result):
    cut = wall_result.stem_reinforcement_cut
    trace = detail.stem_cut_constructive_length_trace(wall_result, cut)
    assert 'Barras cortadas (1 familia(s) de 3/4" @ 0.200 m)' in trace[3]
    assert 'parrilla inferior de 0.100 m' in trace[4]
    document = Document()
    abutment_docx._configure_document(document)
    abutment_docx._summary(document, wall_result)
    rows = [row for row in document.tables[0].rows
            if "barras cortadas" in row.cells[2].text or "continua 1" in row.cells[2].text]
    assert len(rows) == 2
    assert [row.cells[4].text for row in rows] == ["0.200", "0.200"]


@pytest.mark.parametrize("width,yielded", [(.20, False), (.40, True)])
def test_word_flexure_uses_actual_steel_stress_and_reproduces_capacity(width, yielded):
    inputs = abutment.AbutmentInputs()
    inputs = replace(inputs, key=replace(inputs.key, width_m=width))
    initial = abutment.solve_abutment_design(inputs)
    option = custom_spacing_option(initial.key_design.spacing_options, '1"', .10)
    result = abutment.solve_abutment_design(inputs, {"Diente de concreto": option})
    case = result.key_design
    response = rectangular_flexural_response(
        case.provided_as_cm2_m, 100, case.effective_depth_cm,
        inputs.materials.concrete_strength_kg_cm2, inputs.materials.steel_yield_kg_cm2,
        abutment.STEEL_ELASTIC_MODULUS_KG_CM2,
    )
    assert (response.steel_stress_kg_cm2 >= inputs.materials.steel_yield_kg_cm2) == yielded
    formula, _, substitution, _, _ = detail.flexural_resistance_trace(result, case)
    assert "As,prov·fs" in formula and "fs = min(fy, Es·εs)" in formula
    assert f"= {response.compression_block_depth_cm:.6f} cm" in substitution
    assert f"= {response.steel_stress_kg_cm2:.3f} kg/cm²" in substitution
    # Recompute the final printed operation from its own numbers, not the
    # unrounded engine values. This catches formulas with a copied final result.
    match = re.search(r"Mr = ([\d.]+)·([\d.]+)·([\d.]+)·\(([\d.]+) − ([\d.]+)/2\)/100000 = ([\d.]+)", substitution)
    assert match is not None
    phi, area, stress, depth, block, printed = map(float, match.groups())
    assert phi * area * stress * (depth - block/2) / 100000 == pytest.approx(printed, abs=.00051)
    assert printed == pytest.approx(case.moment_resistance_tn_m_m, abs=.00051)
    if not yielded:
        assert printed == pytest.approx(7.126, abs=.00051)
        assert "8.423" not in substitution


def test_structural_word_block_consumes_shared_compatible_trace(default_result, monkeypatch):
    blocks = {}
    monkeypatch.setattr(abutment_docx, "_calc", lambda doc, title, *fields: blocks.update({title: fields}))
    document = Document()
    abutment_docx._configure_document(document)
    case = default_result.heel_design
    abutment_docx._structural_case(document, default_result, case)
    assert blocks["Refuerzo adoptado y momento resistente"][:5] == detail.flexural_resistance_trace(default_result, case)
    assert "0.59" not in blocks["Acero por resistencia a flexión"][0]
    assert "Mr(As)" in blocks["Acero por resistencia a flexión"][0]


@pytest.mark.parametrize("attribute", ["stem_design", "heel_design", "toe_design", "key_design"])
def test_word_capacity_agrees_with_each_structural_case(default_result, attribute):
    case = getattr(default_result, attribute)
    response = rectangular_flexural_response(
        case.provided_as_cm2_m, 100, case.effective_depth_cm,
        default_result.inputs.materials.concrete_strength_kg_cm2,
        default_result.inputs.materials.steel_yield_kg_cm2,
        abutment.STEEL_ELASTIC_MODULUS_KG_CM2,
    )
    assert case.applied_flexural_phi * response.nominal_moment_tn_m == pytest.approx(case.moment_resistance_tn_m_m)
    text = detail.flexural_resistance_trace(default_result, case)[2]
    assert f"= {case.moment_resistance_tn_m_m:.3f} Tn·m/m" in text
    if attribute == "stem_design":
        assert f"Mr,R = {case.strength_moment_resistance_tn_m_m:.3f}" in text
        assert f"Mr,E = {case.extreme_moment_resistance_tn_m_m:.3f}" in text


def test_failed_comparisons_survive_native_word_equation_serialization(default_result):
    from io import BytesIO

    case = replace(default_result.stem_design, shear_demand_tn_m=44.354,
                   shear_resistance_tn_m=31.501, shear_status="NO")
    check = replace(default_result.crack_checks[0], steel_stress_kg_cm2=3062.2,
                    steel_stress_limit_kg_cm2=2520., stress_status="NO CUMPLE",
                    status="NO CUMPLE")
    document = Document()
    abutment_docx._configure_document(document)
    abutment_docx._calc(document, "Cortante", *detail.shear_beta_trace(default_result, case), "")
    abutment_docx._calc(document, "Servicio", *detail.crack_control_trace(default_result, check), "")
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    reopened = Document(buffer)
    assert reopened._element.xpath(".//m:oMath")
    text = "".join(reopened._element.xpath(".//m:t/text() | .//w:t/text()"))
    compact = re.sub(r"\s+", "", text)
    assert "44.354Tn/m>Vr=31.501Tn/m:NO" in compact
    assert "3062.2kg/cm²>0.60fy=2520.0kg/cm²:NOCUMPLE" in compact


@pytest.mark.parametrize("factor,sign,status", [(.5, "≤", "OK"), (1., "≤", "OK"), (2., ">", "NO")])
def test_shear_comparison_follows_numbers(default_result, factor, sign, status):
    case = replace(default_result.stem_design, shear_demand_tn_m=10*factor,
                   shear_resistance_tn_m=10, shear_status=status)
    result = detail.shear_beta_trace(default_result, case)[3]
    assert f"Tn/m {sign} Vr = 10.000 Tn/m: {status}" in result


@pytest.mark.parametrize("stress_factor", [.5, 1., 2.])
@pytest.mark.parametrize("spacing_factor", [.5, 1., 2.])
def test_crack_stress_and_spacing_have_independent_comparisons(default_result, stress_factor, spacing_factor):
    stress_ok, spacing_ok = stress_factor <= 1, spacing_factor <= 1
    check = replace(default_result.crack_checks[0], steel_stress_kg_cm2=2520*stress_factor,
                    steel_stress_limit_kg_cm2=2520, provided_spacing_m=.2*spacing_factor,
                    maximum_spacing_m=.2, stress_status="CUMPLE" if stress_ok else "NO CUMPLE",
                    spacing_status="CUMPLE" if spacing_ok else "NO CUMPLE",
                    status="CUMPLE" if stress_ok and spacing_ok else "NO CUMPLE")
    trace = detail.crack_control_trace(default_result, check)
    assert f"kg/cm² {'≤' if stress_ok else '>'} 0.60fy" in trace[2]
    assert f"m {'≤' if spacing_ok else '>'} smax" in trace[3]
    assert f"resultado conjunto: {check.status}" in trace[3]


@pytest.mark.parametrize("factor,upper,lower,status", [(.5, "≤", "<", "NO"), (1., "≤", "≥", "OK"), (2., ">", "≥", "OK")])
def test_stability_word_comparisons(default_result, monkeypatch, factor, upper, lower, status):
    state = replace(default_result.with_bridge[0], eccentricity_m=factor,
                    eccentricity_limit_m=1., geotechnical_pressure_kg_cm2=factor,
                    q_allow_kg_cm2=1., key_resistance_tn_m=factor,
                    hu_tn_m=1., sliding_with_key_status=status)
    assert f"m {upper} e_lím" in detail.eccentricity_limit_trace(default_result, state)[3]
    assert f"kg/cm² {upper} qlím" in detail.allowable_bearing_trace(default_result, state)[3]
    blocks = {}
    monkeypatch.setattr(abutment_docx, "_calc", lambda doc, title, *fields: blocks.update({title: fields}))
    document = Document()
    abutment_docx._configure_document(document)
    abutment_docx._stability_state_calculations(document, default_result, state, "con puente")
    assert f"{factor:.3f} {lower} 1.000" in blocks["Verificación al deslizamiento"][2]
    assert status in blocks["Verificación al deslizamiento"][3]


@pytest.mark.parametrize("area_factor", [.5, 2.])
@pytest.mark.parametrize("spacing_factor", [.5, 2.])
def test_secondary_steel_comparisons(default_result, area_factor, spacing_factor):
    case = replace(default_result.secondary_reinforcement[0], required_as_cm2_m=5.,
                   provided_as_cm2_m=5*area_factor, selected_spacing_m=.2*spacing_factor,
                   maximum_spacing_m=.2)
    text = detail.secondary_temperature_trace(case)[3]
    assert f"cm²/m {'≥' if area_factor >= 1 else '<'} As,adic" in text
    assert f"cm {'≤' if spacing_factor <= 1 else '>'} smax" in text


@pytest.mark.parametrize("complies", [True, False])
def test_cut_word_comparisons(wall_result, complies):
    cut = replace(wall_result.stem_reinforcement_cut,
                  upper_spacing_m=.2 if complies else .4, upper_spacing_limit_m=.3,
                  upper_provided_as_cm2_m=10 if complies else 2, minimum_as_cm2_m=5,
                  required_as_at_cut_cm2_m=5, moment_resistance_at_cut_tn_m_m=20 if complies else 5,
                  required_moment_at_cut_tn_m_m=10)
    pattern = detail.stem_cut_continuous_pattern_trace(wall_result, cut)
    theoretical = detail.stem_cut_theoretical_height_trace(wall_result, cut)
    assert f"m {'≤' if complies else '>'} smax,pat" in pattern[2]
    assert f"cm²/m {'≥' if complies else '<'} As,temp" in pattern[3]
    assert f"Tn·m/m {'≥' if complies else '<'} Mreq" in theoretical[3]
    assert f"As,sup {'≥' if complies else '<'} As,req" in theoretical[3]
