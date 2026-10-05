from copy import deepcopy
import json
from pathlib import Path
import tomllib

import pytest
from docx import Document

from bridge_design.cli.connected_yaml import connected_inputs_from_yaml, connected_yaml_template
from bridge_design.cli.yaml_io import load_yaml_file
from bridge_design.connected_main import main


def test_template_requires_subgrade_but_not_interface_friction():
    data = connected_yaml_template()
    with pytest.raises(ValueError, match="balasto"):
        connected_inputs_from_yaml(data)
    data["suelo_cimentacion"]["modulo_balasto_vertical_tn_m3"] = 3000
    assert connected_inputs_from_yaml(data).soil.friction_coefficient == 0.0


def test_common_editable_geometry_and_simultaneous_load_pairs():
    data = connected_yaml_template(example=True)
    data["estribo"]["geometria"]["altura_relleno_activo_m"] = 8.0
    left_loads = deepcopy(data["estribo"]["cargas_tablero"])
    right_loads = deepcopy(left_loads)
    right_loads["pll_im_vehicular_tn_m"] = 2.5
    data["casos_simultaneos"] = [{"nombre": "Vehiculo cerca de izquierda",
                                 "cargas_izquierda": left_loads, "cargas_derecha": right_loads,
                                 "factor_sobrecarga_derecha": 0.0}]
    result = connected_inputs_from_yaml(data)
    assert result.right.geometry.retained_height_m == 8
    assert result.left.geometry == result.right.geometry
    assert result.cases[0].right.pll_im_tn_m == 2.5
    assert result.cases[0].right_surcharge == 0


def test_legacy_yaml_preserves_reactions_and_rejects_different_geometry():
    data = connected_yaml_template(example=True)
    common = data.pop("estribo")
    data["estribo_izquierdo"] = deepcopy(common)
    data["estribo_derecho"] = deepcopy(common)
    data["estribo_derecho"]["cargas_tablero"]["pll_im_vehicular_tn_m"] = 2.5
    inputs = connected_inputs_from_yaml(data)
    assert inputs.left.loads.pll_im_tn_m == 6.92
    assert inputs.right.loads.pll_im_tn_m == 2.5
    data["estribo_derecho"]["geometria"]["altura_relleno_activo_m"] = 8.0
    with pytest.raises(ValueError, match="diseno comun"):
        connected_inputs_from_yaml(data)


def test_console_commands_registered_and_dispatched(monkeypatch):
    import bridge_design.main as bridge_main
    import bridge_design.connected_main as connected_main
    calls = []
    monkeypatch.setattr(connected_main, "main", lambda args: calls.append(args))
    bridge_main.main(["estribos-conectados", "--help"])
    assert calls == [["--help"]]
    with Path("pyproject.toml").open("rb") as stream:
        assert tomllib.load(stream)["project"]["scripts"]["diseno-estribos-conectados"] == "bridge_design.connected_main:main"


@pytest.mark.parametrize("without_word", [False, True])
def test_output_generates_template_and_cli_runs_complete_example(tmp_path, capsys, without_word):
    template = tmp_path / "modelo.yaml"
    main(["output", str(template), "--ejemplo"])
    assert load_yaml_file(template)["suelo_cimentacion"]["modulo_balasto_vertical_tn_m3"] == 3000
    destination = tmp_path / "resultados"
    main(["input", str(template), "--resultados", str(destination), "--automatico", "--verificar-malla"]
         + (["--sin-word"] if without_word else []))
    document = json.loads((destination / "resultados.json").read_text(encoding="utf-8"))
    assert len(document["results"]) == 17
    assert all("Mixta" not in case["name"] for case in document["results"])
    assert document["reinforcement"]
    assert document["mesh_comparison"]["coarse_nodes"] == 41
    assert document["mesh_comparison"]["fine_nodes"] == 81
    assert len(document["mesh"]["frame"]["springs"]) == 41
    assert document["mesh_comparison"]["fine_step"] == document["mesh_comparison"]["coarse_step"] / 2
    assert (destination / "nodos.csv").exists()
    assert (destination / "esfuerzos.csv").exists()
    assert (destination / "geometria.png").exists()
    word_path = destination / "memoria_estribos_conectados.docx"
    assert word_path.exists() == (not without_word)
    if not without_word:
        memory = Document(word_path)
        paragraphs = "\n".join(p.text for p in memory.paragraphs)
        assert "6.9. Envolvente general de todas las combinaciones" in paragraphs
        assert "Envolvente general de momento M de todas las combinaciones" in paragraphs
        assert "Deformada nodal de servicio con amplificación indicada en la figura" not in paragraphs
        assert not any("Transicion" in p.text for p in memory.paragraphs if p.style.name.startswith("Heading"))
        assert paragraphs.count("Capacidad portante y presión límite factorizada") == 4
    assert "ANALISIS Y DISENO" in capsys.readouterr().out


def test_cli_invalid_input_exits_nonzero_without_reports(tmp_path):
    template = tmp_path / "obligatorio.yaml"
    main(["output", str(template)])
    destination = tmp_path / "sin_resultados"
    with pytest.raises(SystemExit) as error:
        main(["input", str(template), "--resultados", str(destination)])
    assert error.value.code == 2
    assert not destination.exists()


def test_nonfinite_data_are_rejected():
    data = connected_yaml_template(example=True)
    data["estribo"]["cargas_tablero"]["pdc_carga_muerta_tablero_tn_m"] = float("nan")
    with pytest.raises(ValueError, match="finito"):
        connected_inputs_from_yaml(data)
