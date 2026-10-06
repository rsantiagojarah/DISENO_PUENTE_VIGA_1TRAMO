from dataclasses import replace
import pytest

from bridge_design.cli.bearing_a_pair_yaml import bearing_pair_from_yaml,bearing_pair_template
from bridge_design.domain.bearing_a_pair import design_bearing_pair


def test_pair_calculates_both_without_connection_checks():
    data = bearing_pair_template()
    pair = design_bearing_pair(bearing_pair_from_yaml(data))
    assert pair.mobile.value('DELTA') > 0
    assert pair.fixed.value('DELTA') == 0
    assert pair.fixed.value('P') == pair.mobile.value('P')
    assert 'SLIP' in {s.id for s in pair.mobile.steps}
    assert 'SLIP' not in {s.id for s in pair.fixed.steps}
    assert all(not s.id.startswith(('EQ_','CONNECTION_','CONCRETE')) for r in (pair.mobile,pair.fixed) for s in r.steps)


def test_vertical_reactions_and_geometry_are_independent():
    data = bearing_pair_template()
    data['fijo']['acciones']['dc_tn'] = 40
    data['fijo']['geometria']['largo_cm'] = 40
    data['desplazamiento_fijo_cm'] = .12
    pair = design_bearing_pair(bearing_pair_from_yaml(data))
    assert pair.fixed.value('P') != pair.mobile.value('P')
    assert pair.fixed.adopted.length_cm == 40
    assert pair.fixed.value('DELTA') == .12


def test_legacy_file_uses_explicit_equal_reaction_hypothesis():
    from bridge_design.cli.bearing_a_yaml import bearing_a_template
    inputs = bearing_pair_from_yaml(bearing_a_template())
    assert inputs.equal_vertical_actions
    assert inputs.mobile.actions == inputs.fixed.actions


def test_failed_fixed_check_controls_combined_status():
    data = bearing_pair_template()
    data['fijo']['geometria']['zuncho_cm'] = .01
    pair = design_bearing_pair(bearing_pair_from_yaml(data))
    assert pair.fixed.status == 'NO CONFORME'
    assert pair.status == 'NO CONFORME'


def test_pair_output_is_compact_and_has_both_supports():
    from bridge_design.cli.bearing_a_pair_output import format_bearing_pair
    pair = design_bearing_pair(bearing_pair_from_yaml(bearing_pair_template()))
    text = format_bearing_pair(pair,width=112)
    assert len(text.splitlines()) < 90
    assert 'Fijo' in text and 'Móvil' in text
    assert 'anclaje' not in text.lower() and 'Conexión y combinación' not in text


def test_manual_pair_never_asks_for_anchors(monkeypatch):
    from bridge_design.cli.bearing_a_cli import collect_bearing_pair_inputs
    prompts = []
    def answer(prompt):
        prompts.append(prompt)
        return '.45' if 'Acortamiento por retracción calculado' in prompt else ''
    monkeypatch.setattr('builtins.input',answer)
    inputs = collect_bearing_pair_inputs()
    assert inputs.equal_vertical_actions
    assert not any('anclaj' in p.lower() or 'Evento Extremo' in p for p in prompts)


def test_pair_report_has_no_external_verification_requirements(tmp_path):
    from docx import Document
    from bridge_design.reporting.bearing_a_pair_docx import generate_bearing_pair_docx
    pair = design_bearing_pair(bearing_pair_from_yaml(bearing_pair_template()))
    doc = Document(generate_bearing_pair_docx(pair,tmp_path/'par.docx'))
    text = '\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert 'A Diseño del apoyo móvil' in text and 'B Diseño del apoyo fijo' in text
    assert 'anclaj' not in text.lower()
    assert 'C Restricciones' not in text and 'PENDIENTE' not in text
