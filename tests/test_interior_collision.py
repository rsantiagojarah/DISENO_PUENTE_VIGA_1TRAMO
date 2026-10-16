from dataclasses import replace

import pytest

from bridge_design.codes.mtc_2018 import mtc_tension_development_length_cm
from bridge_design.domain.barrier import BarrierDesignInputs, design_concrete_barrier
from bridge_design.domain.cantilever_slab import CantileverSlabParameters, design_cantilever_slab
from bridge_design.domain import interior_collision
from bridge_design.domain.rebar_catalog import ReinforcementSpacingOption, reinforcing_bar_by_label
from bridge_design.domain.loads import LiveLoads, PedestrianLoad, VehicleLoadModel
from bridge_design.domain.transverse_slab import PointMoment, solve_load_case
from bridge_design.domain.reinforcement import flexural_steel_area_cm2
from test_cantilever_slab import _geometry, _layout, _materials


def _option(label='5/8"', spacing=.15):
    bar = reinforcing_bar_by_label(label)
    return ReinforcementSpacingOption(99, bar, spacing, 1.0, bar.area_cm2/spacing,
                                      True, is_custom=True)


def _slab_selection(upper=None, lower=None):
    return (("A. Acero negativo", upper or _option()),
            ("B. Acero positivo", lower or _option()))


def _result(level="TL-4", overhang=1.0, *, barrier_outer_face=None, cover_cm=5.0, slab_selected=None,
            slab_cover_cm=None):
    g = replace(_geometry(), overhang_m=overhang, girder_spacing_m=2.56,
                girder_width_m=.5, girder_count=5)
    outer_face = overhang + .7 if barrier_outer_face is None else barrier_outer_face
    layout = replace(_layout(), barrier_left_m=outer_face, barrier_width_m=.375)
    inputs = BarrierDesignInputs()
    inputs = replace(inputs, impact_load=replace(inputs.impact_load, test_level=level))
    barrier = design_concrete_barrier(inputs, _materials())
    live = LiveLoads(PedestrianLoad.mtc_sidewalk_default(), VehicleLoadModel.mtc_hl93_default())
    params = CantileverSlabParameters(concrete_cover_cm=cover_cm)
    selection = _slab_selection() if slab_selected is None else slab_selected
    return g, design_cantilever_slab(g, _materials(), live, layout, barrier, params,
                                    slab_selected=selection, slab_cover_cm=slab_cover_cm), barrier


def test_point_couple_matches_simply_supported_closed_form():
    g = replace(_geometry(), girder_count=2, overhang_m=0, girder_spacing_m=4)
    result = solve_load_case(g, _materials(), "couple", (), (),
                             point_moments=(PointMoment(1.5, 8, "couple"),))
    assert [r for _, r in result.support_reactions_tn] == pytest.approx([2, -2])
    assert [m for x, m in result.moment_samples_tn_m if x == 1.5] == pytest.approx([3, -5])
    assert result.moment_samples_tn_m[-1][1] == pytest.approx(0, abs=1e-10)


def test_interior_cases_conserve_equilibrium_and_mirror_without_combining():
    _, result, barrier = _result()
    local = result.interior_collision
    assert result.barrier_collision is None
    assert len(local.cases) == 8
    assert local.transfer_length_m == pytest.approx(barrier.yield_line.critical_length_m + 2*barrier.geometry.height_m)
    assert local.interface_moment_tn_m_m == pytest.approx(barrier.flexure.mc_tn_m)
    for case in local.cases:
        assert abs(case.vertical_equilibrium_error_tn_m) < 1e-9
        assert abs(case.moment_equilibrium_error_tn_m_m) < 1e-9
        if "vertical" in case.name:
            assert case.axial_tension_tn_m == case.applied_moment_tn_m_m == 0
        else:
            assert case.applied_vertical_tn_m == 0
            assert case.axial_tension_tn_m > 0
    for left, right in zip(local.cases[:4], local.cases[4:]):
        assert [r for _, r in left.reactions_tn_m] == pytest.approx([r for _, r in reversed(right.reactions_tn_m)])
    assert "PENDIENTE: verificar transferencia" not in " ".join(result.applicability_notes)
    assert not result.overall_ok  # A local strip cannot certify the global bridge.


def test_vertical_collision_uses_test_level_not_default():
    _, tl4, _ = _result()
    _, tl5, _ = _result("TL-5")
    v4 = tl4.interior_collision.cases[1].applied_vertical_tn_m
    v5 = tl5.interior_collision.cases[1].applied_vertical_tn_m
    assert v5/v4 == pytest.approx((80/40)/(18/18))


def test_selected_steel_is_verified_with_its_actual_diameter():
    g, result, _ = _result()
    for face in result.interior_collision.faces:
        option = face.option
        assert option is not None
        d = 100*g.slab_thickness_m - result.parameters.concrete_cover_cm - option.bar.diameter_cm/2
        required = flexural_steel_area_cm2(face.moment_tn_m_m, 100, d, 280, 4200, phi=.9)
        required += face.tension_tn_m*1000/(.75*4200)
        assert option.provided_area_cm2_m >= required
        if face.development_length_m > face.development_available_m:
            assert face.status == "NO CUMPLE"


def test_actual_bars_spacing_and_area_are_preserved_without_adding_collision_steel():
    upper = replace(_option('3/4"', .137), provided_area_cm2_m=999, is_compliant=False)
    lower = _option('1/2"', .121)
    selection = _slab_selection(upper, lower)
    g, result, _ = _result(slab_selected=selection)
    for face, (_, original) in zip(result.interior_collision.faces, selection):
        assert face.option.bar == original.bar
        assert face.option.spacing_m == original.spacing_m
        assert face.option.is_custom
        assert face.option.provided_area_cm2_m == pytest.approx(original.bar.area_cm2/original.spacing_m)
        d = 100*g.slab_thickness_m - 5 - original.bar.diameter_cm/2
        required = flexural_steel_area_cm2(face.moment_tn_m_m, 100, d, 280, 4200, phi=.9)
        required += face.tension_tn_m*1000/(.75*4200)
        assert face.option.required_area_cm2_m == pytest.approx(max(required, .0018*100*100*g.slab_thickness_m))
        ld, _ = mtc_tension_development_length_cm(original.bar.diameter_cm, 4200, 280,
                                                 excess_reinforcement_factor=1.0)
        assert face.development_length_m == pytest.approx(ld/100)
    # Rechecking creates separate demand/status records, not a mutation or an extra mesh.
    assert upper.provided_area_cm2_m == 999
    assert not upper.is_compliant
    assert upper.required_area_cm2_m == lower.required_area_cm2_m == 1.0


@pytest.mark.parametrize("label, spacing, area_status, anchorage_status", [
    ('3/8"', .30, "NO CUMPLE", "OK"),
    ('1"', .15, "OK", "NO CUMPLE"),
    ('5/8"', .15, "OK", "OK"),
])
def test_area_and_anchorage_are_independent_and_never_substitute_bars(
    label, spacing, area_status, anchorage_status,
):
    selected = _option(label, spacing)
    _, result, _ = _result(overhang=.95, barrier_outer_face=1.50,
                           slab_selected=_slab_selection(selected, selected))
    for face in result.interior_collision.faces:
        assert face.option.bar == selected.bar
        assert face.option.spacing_m == spacing
        assert face.steel_status == area_status
        assert face.development_status == anchorage_status
        assert face.development_available_m == pytest.approx(1.45)
        assert face.status == ("OK" if area_status == anchorage_status == "OK" else "NO CUMPLE")


@pytest.mark.parametrize("selection, missing", [
    ((), {"superior", "inferior"}),
    ((("A. Acero negativo", _option()),), {"inferior"}),
    ((("C. Temperatura", _option()), ("D. Distribucion", _option())), {"superior", "inferior"}),
])
def test_missing_main_slab_steel_is_pending_not_an_automatic_recommendation(selection, missing):
    _, result, _ = _result(slab_selected=selection)
    collision = result.interior_collision
    for face in collision.faces:
        if face.face in missing:
            assert face.option is None
            assert face.status == face.steel_status == face.development_status == "PENDIENTE"
            assert face.development_available_m == pytest.approx(1.65)
    assert collision.shear_status == "PENDIENTE"
    assert collision.status != "OK"
    assert "phi Vc=0.000" not in " ".join(collision.report_lines())


def test_changing_slab_bars_refreshes_shear_but_not_loads_or_connection():
    _, light, _ = _result(slab_selected=_slab_selection(_option('1/2"', .30), _option('1/2"', .30)))
    _, heavy, _ = _result(slab_selected=_slab_selection(_option('3/4"', .10), _option('3/4"', .10)))
    a, b = light.interior_collision, heavy.interior_collision
    assert a.cases == b.cases
    assert a.shear_demand_tn_m == b.shear_demand_tn_m
    assert a.connection_checks == b.connection_checks
    assert a.shear_capacity_tn_m < b.shear_capacity_tn_m
    for light_face, heavy_face in zip(a.faces, b.faces):
        assert light_face.option.required_area_cm2_m < heavy_face.option.required_area_cm2_m
        assert light_face.development_length_m < heavy_face.development_length_m


def test_collision_uses_interior_slab_cover_without_changing_overhang_cover():
    g, result, _ = _result(cover_cm=5.0, slab_cover_cm=6.0)
    assert result.parameters.concrete_cover_cm == 5.0
    for face in result.interior_collision.faces:
        d = 100*g.slab_thickness_m - 6.0 - face.option.bar.diameter_cm/2
        required = flexural_steel_area_cm2(face.moment_tn_m_m, 100, d, 280, 4200, phi=.9)
        required += face.tension_tn_m*1000/(.75*4200)
        assert face.option.required_area_cm2_m == pytest.approx(required)
        assert face.development_available_m == pytest.approx(1.64)


@pytest.mark.parametrize("overhang, outer_face, cover_cm, available", [
    (.95, 1.50, 5.0, 1.45),
    (.95, 1.65, 5.0, 1.60),
    (.80, 1.50, 5.0, 1.45),
    (.95, 1.50, 6.0, 1.44),
])
def test_local_development_uses_outer_barrier_face_not_girder_axis(
    overhang, outer_face, cover_cm, available,
):
    g, result, _ = _result(overhang=overhang, barrier_outer_face=outer_face, cover_cm=cover_cm)
    collision = result.interior_collision
    assert {face.face for face in collision.faces} == {"superior", "inferior"}
    # The right barrier is mirrored: its outer face lies at B - ab.
    right_outer_face = g.total_width_m - outer_face
    assert g.total_width_m - right_outer_face - cover_cm / 100 == pytest.approx(available)
    for face in collision.faces:
        assert face.development_available_m == pytest.approx(available)
        assert face.development_available_m != pytest.approx(overhang - cover_cm / 100)
        assert face.development_available_m != pytest.approx(outer_face + .375 / 2 - cover_cm / 100)


@pytest.mark.parametrize("delta, expected", [(-1e-6, "NO CUMPLE"), (0, "OK"), (1e-6, "OK")])
def test_local_development_still_requires_full_bar_development(delta, expected):
    selected = _option('1/2"')
    bar = selected.bar
    required_cm, _ = mtc_tension_development_length_cm(
        bar.diameter_cm, 4200, 280, excess_reinforcement_factor=1.0,
    )
    available = required_cm / 100 + delta
    result = interior_collision._review_face(
        "superior", [(1.0, 0.0, "caso", 1.50, 1.25, 1.50)],
        _geometry(), _materials(), CantileverSlabParameters(), selected,
        development_available_m=available,
    )
    assert result.option is not None
    assert result.option.provided_area_cm2_m > result.option.required_area_cm2_m
    assert result.development_length_m == pytest.approx(required_cm / 100)
    assert result.development_available_m == pytest.approx(available)
    assert result.status == expected


@pytest.mark.parametrize("demand_position", [.95, 1.50, 1.875, 10.64])
def test_demand_coordinate_is_separate_from_local_anchorage_reference(demand_position):
    result = interior_collision._review_face(
        "superior", [(1.0, 0.0, "caso", demand_position, 1.25, 1.50)],
        _geometry(), _materials(), CantileverSlabParameters(), _option(),
        development_available_m=1.45,
    )
    assert result.position_m == demand_position
    assert result.development_available_m == pytest.approx(1.45)


def test_unavailable_reinforcement_keeps_geometric_anchorage_length():
    result = interior_collision._review_face(
        "superior", [(1.0, 0.0, "caso", 1.50, 1.25, 1.50)],
        replace(_geometry(), slab_thickness_m=.05), _materials(), CantileverSlabParameters(),
        _option(),
        development_available_m=1.45,
    )
    assert result.option.bar == _option().bar
    assert result.option.required_area_cm2_m == float("inf")
    assert result.reason
    assert result.status == "NO CUMPLE"
    assert result.development_available_m == pytest.approx(1.45)


def test_report_contains_demands_steel_reactions_and_real_status():
    from bridge_design.cli.ascii_output import format_cantilever_slab_design_result
    _, result, _ = _result()
    text = format_cantilever_slab_design_result(result)
    for expected in (
        "TRANSFERENCIA DE COLISION A LOSA INTERIOR",
        "COLISION: VERIFICACION DEL ARMADO DE LOSA",
        "Reacciones verticales incrementales",
        "As req",
        "ld req",
        "Conexion barrera-losa",
        "Resultado local:",
        "desde la cara exterior de la base de cada barrera",
        "Ldisp = max(0, ab - c) = max(0, 1.700 - 0.050) = 1.650 m",
        "no sustituye las verificaciones de desarrollo en otras secciones",
    ):
        assert expected in text
    applicability = text.split("Notas de aplicabilidad:", 1)[1]
    assert "EEII horizontal" not in applicability


def test_unknown_level_is_rejected_instead_of_assuming_tl4():
    with pytest.raises(ValueError, match="Nivel de ensayo"):
        _result("unknown")


def test_capacity_anchorage_does_not_reuse_reduced_barrier_check():
    _, result, barrier = _result()
    hook = barrier.development
    full = max(hook.minimum_ldh_cm, hook.modified_ldh_cm/hook.excess_reinforcement_factor)
    if hook.available_length_cm < full:
        assert result.interior_collision.connection_status == "NO CUMPLE"
        assert result.interior_collision.status == "NO CUMPLE"


def test_word_and_pdf_sections_include_interior_transfer():
    from types import SimpleNamespace
    from docx import Document
    from bridge_design.reporting.deck_docx import _cantilever, _configure_document
    from bridge_design.reporting.deck_sections_components import cantilever_story
    from bridge_design.reporting import pdf_style
    g, result, _ = _result()
    data = SimpleNamespace(cantilever_result=result, cantilever_selected=(),
        cantilever_crack=result.crack_control, cantilever_development=result.development,
        project_inputs=SimpleNamespace(transverse_slab=SimpleNamespace(geometry=g), materials=_materials()))
    document = Document()
    _configure_document(document)
    _cantilever(document, data)
    text = document._element.xml
    assert "Transferencia local a la losa interior" in text
    assert "Control" in text
    assert "Acero transversal por colisión" in text
    # Use the production style factory rather than mocking report paragraphs.
    story = cantilever_story(data, pdf_style.report_styles())
    paragraphs = " ".join(getattr(item, "text", "") for item in story)
    assert "Transferencia de colision a la losa interior" in paragraphs
    assert "Verificacion por colision del acero transversal de losa adoptado" in paragraphs
    for expected in (
        "desde la cara exterior de la base de cada barrera",
        "Ldisp = max(0, ab - c) = max(0, 1.700 - 0.050) = 1.650 m",
        "no sustituye las verificaciones de desarrollo en otras secciones",
    ):
        assert expected in text
        assert expected in paragraphs
        assert expected in " ".join(local_line for local_line in result.interior_collision.report_lines())


@pytest.mark.parametrize("missing", [False, True])
def test_word_pdf_and_terminal_publish_actual_bars_and_separate_failures(missing):
    from docx import Document
    from bridge_design.cli.ascii_output import format_interior_collision_result
    from bridge_design.reporting.deck_docx import _configure_document, _interior_collision_docx
    from bridge_design.reporting.deck_sections_components import _interior_collision_story
    from bridge_design.reporting.pdf_style import report_styles

    selection = () if missing else _slab_selection(_option('3/8"', .30), _option('1"', .15))
    _, result, _ = _result(slab_selected=selection)
    collision = result.interior_collision
    document = Document()
    _configure_document(document)
    _interior_collision_docx(document, collision)
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    areas = next(rows for rows in tables if rows[0] == ["Cara", "As requerido", "Acero de losa adoptado", "As provisto", "Área"])
    anchorage = next(rows for rows in tables if rows[0] == ["Cara", "ld requerido", "ld disponible", "Anclaje", "Resultado de cara"])
    story = _interior_collision_story(collision, report_styles())
    pdf_tables = [item for item in story if hasattr(item, "_cellvalues")]
    def cell_text(cell):
        return getattr(cell, "text", str(cell))
    pdf_areas = next(table._cellvalues for table in pdf_tables
                     if any("Acero de losa adoptado" in cell_text(c) for c in table._cellvalues[0]))
    pdf_anchorage = next(table._cellvalues for table in pdf_tables
                         if any("Resultado de cara" in cell_text(c) for c in table._cellvalues[0]))
    terminal = format_interior_collision_result(collision)
    for index, face in enumerate(collision.faces, 1):
        bar_text = f'{face.option.bar.label} @ {face.option.spacing_m:.3f} m' if face.option else "Sin seleccionar"
        assert areas[index][2] == bar_text
        assert areas[index][-1] == face.steel_status
        assert anchorage[index][-2:] == [face.development_status, face.status]
        assert cell_text(pdf_areas[index][2]) == bar_text
        assert cell_text(pdf_areas[index][-1]) == face.steel_status
        assert [cell_text(c) for c in pdf_anchorage[index][-2:]] == [face.development_status, face.status]
        assert bar_text in terminal
    if missing:
        assert "seleccionar el acero de losa de ambas caras" in document._element.xml
        assert "φVc=0.000" not in document._element.xml
    else:
        assert areas[1][-1] == anchorage[2][-2] == "NO CUMPLE"
        assert areas[2][-1] == anchorage[1][-2] == "OK"
