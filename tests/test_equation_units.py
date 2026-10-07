from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from bridge_design.reporting.deck_docx import _add_native_equation


def test_strict_comparisons_remain_outside_fraction_denominators():
    document = Document()
    document.styles.add_style('Equation', WD_STYLE_TYPE.PARAGRAPH)
    for operator in ('<', '>'):
        paragraph = _add_native_equation(document, f'S^2/n {operator} 16')
        fraction = paragraph._p.xpath('.//m:f')[0]
        assert ''.join(fraction.find(qn('m:den')).itertext()) == 'n'
        assert any(operator in text for text in paragraph._p.xpath('.//m:oMath/m:r/m:t/text()'))


def test_general_shear_beta_preserves_two_fractional_factors():
    document = Document()
    document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    paragraph = _add_native_equation(document, 'β = (4.8/(1+750·εs))·(51/(39+sxe))')
    fractions = paragraph._p.xpath('.//m:f')
    assert len(fractions) == 2
    assert not paragraph._p.xpath('.//m:den//m:f')
    assert [''.join(f.find(qn('m:num')).itertext()) for f in fractions] == ['4.8', '51']
    assert [''.join(f.find(qn('m:den')).itertext()) for f in fractions] == ['(1 + 750 · εs)', '(39 + sxe)']


def _fraction_texts(paragraph):
    return [
        (
            ''.join(fraction.find(qn('m:num')).itertext()),
            ''.join(fraction.find(qn('m:den')).itertext()),
        )
        for fraction in paragraph._p.xpath('.//m:f')
    ]


def test_published_abutment_factors_stay_outside_the_first_denominator():
    document = Document()
    document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    pressure = _add_native_equation(document, "qmax = (66.944/6.000)·(1 + 6·0.841/6.000)/10 = 2.054")
    pressure_fractions = _fraction_texts(pressure)
    assert any(denominator.strip() == "10" for _, denominator in pressure_fractions)
    assert ("66.944", "6.000") in pressure_fractions
    assert any("0.841" in numerator and denominator.strip() == "6.000" for numerator, denominator in pressure_fractions)
    general = _add_native_equation(document, "qmax,min = (Vu/B)·(1 ± 6|e|/B)/10")
    general_fractions = _fraction_texts(general)
    assert ("Vu", "B") in general_fractions
    assert any(denominator.strip() == "10" for _, denominator in general_fractions)
    eccentricity = next(item for item in general_fractions if "|e|" in item[0] and "Vu" not in item[0])
    assert eccentricity[0].strip() == "6|e|"
    assert eccentricity[1].strip() == "B"
    general_text = ''.join(general._p.xpath('.//m:t/text()'))
    assert "1" in general_text and "±" in general_text
    meyerhof = _add_native_equation(document, "qM = (66.944/4.318)/10 = 1.550")
    assert _fraction_texts(meyerhof)[0][1] == '10'
    temperature = _add_native_equation(document, "As,temp = (7.65·b·h/(2·(b+h)·fy))·100")
    temperature_text = ''.join(temperature._p.xpath('.//m:t/text()'))
    assert '100' in temperature_text
    assert all(denominator != '100' and not denominator.endswith('·100') for _, denominator in _fraction_texts(temperature))
    beta = _add_native_equation(document, "β = (4.8/(1 + 750·0.002255))·(51/(39 + 32.438)) = 1.273233")
    beta_fractions = _fraction_texts(beta)
    assert [item[0] for item in beta_fractions] == ['4.8', '51']
    assert '32.438' in beta_fractions[1][1]
    assert '0.002255' in beta_fractions[0][1]


def test_report_units_stay_inline_and_calculation_divisions_remain_fractions():
    document = Document()
    document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    for text in ("Ec=2717585.754 tn/m²", "q=26.700 tn/m²", "fs=2520 kgf/cm²",
                 "M=12.5 tn·m/m", "γ=2.4 tn/m³"):
        paragraph = _add_native_equation(document, text)
        assert not paragraph._p.xpath(".//m:f")
        assert ''.join(paragraph._p.xpath(".//m:t/text()")) == text.replace('=', ' = ')
    paragraph = _add_native_equation(document, "q=R_i/A_i")
    assert len(paragraph._p.xpath(".//m:f")) == 1


def test_native_powers_bind_only_their_base_and_exponent():
    document = Document()
    document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    cases = (("I=b*t^3/12", "t", "3"),
             ("EH=Ka*gamma_s*H^2/2", "H", "2"),
             ("Mcr=1.1*sqrt(fc)*b*h^2/(6*100000)", "h", "2"),
             ("K=cos(a)^2*cos(b)^2", "(a)", "2"),
             ("Ec=x^0.33 kgf/cm²", "x", "0.33"),
             ("F=(a*b)^(-2)", "(a*b)", "(-2)"))
    for expression, base, exponent in cases:
        paragraph = _add_native_equation(document, expression)
        power = paragraph._p.xpath(".//m:sSup")[0]
        assert ''.join(power.find(qn("m:e")).itertext()) == base
        assert ''.join(power.find(qn("m:sup")).itertext()) == exponent
    assert len(_add_native_equation(document, "K=cos(a)^2*cos(b)^2")._p.xpath(".//m:sSup")) == 2
