"""The centre service peak must not disqualify valid continuous steel near walls."""

import csv
import json
from dataclasses import replace

import pytest
from docx import Document

from bridge_design.cli.connected_2_yaml import connected_2_inputs_from_yaml, connected_2_yaml_template
from bridge_design.cli.connected_selection import collect_connected_selection, format_connected_selection
from bridge_design.domain.connected_design import solve_connected_abutments
from bridge_design.domain.connected_foundation_zones import foundation_zone_choice
from bridge_design.domain.connected_options import (
    apply_connected_selections, check_region_choice, choice_from_steel, connected_reinforcement_options,
)
from bridge_design.domain.connected_reinforcement import ConnectedBarChoice, ConnectedSteelChoice
from bridge_design.reporting.connected_cut_report import foundation_cut_rows, write_foundation_cut
from bridge_design.reporting.connected_export import export_connected_results
from bridge_design.reporting.deck_docx import _configure_document

TOP = "Zapata combinada - longitudinal superior"
BOTTOM = "Zapata combinada - longitudinal inferior"


@pytest.fixture(scope="module")
def analysis():
    raw = connected_2_yaml_template(example=True)
    raw["estribo"]["geometria"]["altura_parapeto_m"] = 1.22
    raw["estribo"]["materiales"].update(fc_concreto_kg_cm2=280, peso_unitario_relleno_tn_m3=2.038)
    raw["estribo"]["suelo_sismo"].update(angulo_friccion_relleno_grados=36.12, pga=.5, fpga=1)
    raw["suelo_cimentacion"].update(modulo_balasto_vertical_tn_m3=620, qadm_tn_m2=13.5)
    raw["cimentacion"]["cantidad_nudos_cimentacion"] = 82
    return solve_connected_abutments(connected_2_inputs_from_yaml(raw))


@pytest.fixture(scope="module")
def zonal(analysis):
    choice = foundation_zone_choice(ConnectedBarChoice('1"', .20, True))
    return apply_connected_selections(analysis, {TOP: choice, BOTTOM: choice})


def test_local_continuous_steel_passes_despite_failing_at_centre(analysis, zonal):
    value = ConnectedBarChoice('1"', .20)
    uniform = check_region_choice(analysis, TOP, ConnectedSteelChoice(value, value))
    assert uniform.crack_utilization == pytest.approx(3008.7687565/2520)
    top = next(s for s in zonal.reinforcement if s.region == TOP)
    cut = top.foundation_reinforcement_cut
    assert cut.status == "APLICA" and top.status == "OK"
    assert top.spacing_m == .1 and cut.pattern.equivalent_spacing_m == .2
    assert cut.pattern.cycle_bars == 2  # Explicit base cannot silently become a denser pattern.
    left, centre, right = cut.zones
    assert [z.reinforcement for z in cut.zones] == ["Solo continuo", "Continuo + adicional", "Solo continuo"]
    assert [z.area_cm2_m for z in cut.zones] == [25, 50, 25]
    assert all(z.status == "OK" and z.crack_utilization <= 1 for z in cut.zones)
    assert centre.service_moment_tn_m_m == pytest.approx(-95.608892584)
    assert abs(left.service_moment_tn_m_m) < abs(centre.service_moment_tn_m_m)
    assert abs(right.service_moment_tn_m_m) < abs(centre.service_moment_tn_m_m)
    assert cut.cutoff_left_m < cut.theoretical_left_m
    assert zonal.results is analysis.results  # No load/FRAME changes.


def test_bottom_additional_steel_at_both_ends(zonal):
    steel = next(s for s in zonal.reinforcement if s.region == BOTTOM)
    cut = steel.foundation_reinforcement_cut
    assert cut.status == "APLICA"
    assert [z.area_cm2_m for z in cut.zones] == [50, 25, 50]
    assert len(cut.additional_intervals_m) == 2
    assert all(z.status == "OK" for z in cut.zones)


def test_user_model_stem_continues_one_in_two(zonal):
    stem = next(s for s in zonal.reinforcement if s.region == "Pantalla - vertical relleno")
    cut = stem.stem_reinforcement_cut
    assert cut.status == "OK"
    assert cut.lower_bar_label == cut.upper_bar_label == '1"'
    assert cut.lower_spacing_m == .1 and cut.upper_spacing_m == .2
    assert cut.continuous_every_n_bars == 2
    assert cut.constructive_cut_height_m == pytest.approx(
        cut.theoretical_cut_height_m + stem.required_straight_anchor_cm / 100)
    assert cut.constructive_cut_height_m < 6.896  # Earlier proposal retaining only one in three.
    assert stem.anchor_status == "SOLO GANCHO"
    assert cut.footing_anchorage_m == pytest.approx(stem.required_hook_anchor_cm / 100)
    assert cut.footing_anchorage_m < stem.available_anchor_cm / 100 < cut.development_extension_m
    assert cut.lower_cut_bar_length_m == pytest.approx(cut.constructive_cut_height_m+cut.footing_anchorage_m)
    assert cut.continuous_bar_length_m == pytest.approx(8.03+cut.footing_anchorage_m)


def test_word_stem_lengths_match_hooked_anchorage(zonal):
    from bridge_design.reporting.connected_audit import build_connected_audit
    from bridge_design.reporting.connected_distribution_report import write_distributed_design
    document = Document()
    _configure_document(document)
    write_distributed_design(document, zonal, zonal.reinforcement, build_connected_audit(zonal))
    rows = {row.cells[0].text: row.cells[1].text for table in document.tables for row in table.rows[1:]}
    stem = next(s for s in zonal.reinforcement if s.region == "Pantalla - vertical relleno")
    cut = stem.stem_reinforcement_cut
    assert rows["Anclaje en zapata para longitudes"] == f"{stem.required_hook_anchor_cm/100:.3f} m con gancho"
    assert rows["Prolongación ld"] == f"{stem.required_straight_anchor_cm/100:.3f} m"
    assert rows["Longitud barras cortadas"] == f"{cut.lower_cut_bar_length_m:.3f} m"
    assert rows["Longitud barras continuas"] == f"{cut.continuous_bar_length_m:.3f} m"


def test_word_summary_uses_local_footing_steel_and_states(zonal):
    from bridge_design.reporting.connected_docx_layout import connected_summary_and_references
    document = Document()
    _configure_document(document)
    connected_summary_and_references(document, zonal)
    rows = [tuple(cell.text for cell in row.cells) for row in document.tables[0].rows[1:]]
    for region in (TOP, BOTTOM):
        steel = next(s for s in zonal.reinforcement if s.region == region)
        selected = [row for row in rows if row[0].startswith(region)]
        assert len(selected) == 3
        for row, zone in zip(selected, steel.foundation_reinforcement_cut.zones):
            assert zone.name in row[0]
            assert f"x={zone.start_m:.3f}–{zone.end_m:.3f} m" in row[0]
            assert zone.reinforcement in row[1]
            spacing = .20 if zone.reinforcement == "Solo continuo" else .10
            assert f'1" @ {spacing:.3f} m' in row[1]
            assert row[2] == f"{zone.area_cm2_m:.3f}"
            assert row[3] == zone.status
    unchanged = next(s for s in zonal.reinforcement if s.region == "Pantalla - vertical relleno")
    assert (unchanged.region, '1" @ 0.100 m', '50.000', unchanged.status) in rows


def test_word_summary_rejected_cut_keeps_uniform_steel(analysis):
    from bridge_design.reporting.connected_docx_layout import reinforcement_summary_rows
    steel = check_region_choice(analysis, TOP, foundation_zone_choice(ConnectedBarChoice('3/4"', .2)))
    assert steel.foundation_reinforcement_cut.status == "NO APLICA"
    assert tuple(reinforcement_summary_rows(steel)) == ((TOP, '3/4" @ 0.100 m', '28.400', steel.status),)


def test_word_summary_nonuniform_continuous_pattern_is_not_uniform_spacing(zonal):
    from bridge_design.reporting.connected_docx_layout import reinforcement_summary_rows
    steel = next(s for s in zonal.reinforcement if s.region == TOP)
    cut = steel.foundation_reinforcement_cut
    pattern = replace(cut.pattern, cycle_bars=3, continuing_bars=2,
                      equivalent_spacing_m=.15, remaining_area_cm2_m=100/3)
    zones = tuple(replace(z, area_cm2_m=100/3) if z.reinforcement == "Solo continuo" else z for z in cut.zones)
    steel = replace(steel, foundation_reinforcement_cut=replace(cut, pattern=pattern, zones=zones))
    rows = tuple(reinforcement_summary_rows(steel))
    assert 'continúan 2 de cada 3' in rows[0][1]
    assert 'hueco máx. 0.200 m' in rows[0][1]
    assert '@ 0.150' not in rows[0][1]
    assert rows[0][2] == '33.333'
    assert rows[1][2] == '50.000'


def test_console_cut_table_and_enter_retain_families(analysis, zonal, monkeypatch):
    from bridge_design.cli.connected_selection import format_connected_options
    group = next(g for g in connected_reinforcement_options(analysis) if g.region == TOP)
    chosen = next(c for c in group.cuts if c.bar_label == '1"' and abs(c.light_spacing_m - .2) < 1e-9)
    text = format_connected_options((group,))
    assert "CORTES - " + TOP in text and chosen.location_text in text
    assert "Cortes definitivos" not in text
    assert not any(c.bar_label == '3/4"' and abs(c.light_spacing_m - .2) < 1e-9 for c in group.cuts)
    stem = next(g for g in connected_reinforcement_options(analysis) if g.region == "Pantalla - vertical relleno")
    assert stem.offer_cuts and all("sobre la base" in c.location_text for c in stem.cuts)
    assert all(abs(c.light_spacing_m - 2 * c.heavy_spacing_m) < 1e-9 for c in (*group.cuts, *stem.cuts))
    answers = iter(('s', chosen.code.lower()))
    monkeypatch.setattr("builtins.input", lambda prompt: next(answers))
    choice = collect_connected_selection(analysis, (group,))[TOP]
    assert choice.foundation_continuous.spacing_m == .20 and choice.principal.spacing_m == .10
    group = next(g for g in connected_reinforcement_options(zonal) if g.region == TOP)
    answers = iter(('s', ''))
    assert collect_connected_selection(zonal, (group,))[TOP] == choice_from_steel(group.adopted)


@pytest.mark.parametrize("spacing", [.05, .4, float('nan'), float('inf')])
def test_invalid_spacing_rejected(analysis, spacing):
    with pytest.raises(ValueError):
        check_region_choice(analysis, TOP, foundation_zone_choice(ConnectedBarChoice('1"', spacing)))


def test_mismatched_families_and_other_regions_rejected(analysis):
    choice = foundation_zone_choice(ConnectedBarChoice('1"', .2))
    with pytest.raises(ValueError, match="dos familias"):
        check_region_choice(analysis, TOP, replace(choice, foundation_continuous=ConnectedBarChoice('3/4"', .2)))
    with pytest.raises(ValueError, match="solo corresponde"):
        check_region_choice(analysis, "Pantalla - vertical relleno", choice)


def test_failed_full_family_does_not_produce_zones_or_cut(analysis):
    steel = check_region_choice(analysis, TOP, foundation_zone_choice(ConnectedBarChoice('3/4"', .2)))
    assert steel.status == "NO CUMPLE"
    assert steel.foundation_reinforcement_cut.status == "NO APLICA"
    assert not steel.foundation_reinforcement_cut.zones
    assert steel.foundation_reinforcement_cut.cutoff_left_m is None


@pytest.mark.parametrize("accepted", [False, True])
def test_word_only_calls_verified_cut_families_adopted(analysis, zonal, accepted):
    steel = (next(s for s in zonal.reinforcement if s.region == TOP) if accepted else
             check_region_choice(analysis, TOP, foundation_zone_choice(ConnectedBarChoice('3/4"', .2))))
    assert steel.foundation_reinforcement_cut.status == ("APLICA" if accepted else "NO APLICA")
    document = Document()
    _configure_document(document)
    write_foundation_cut(document, steel)
    rows = {row.cells[0].text: row.cells[1].text for table in document.tables for row in table.rows[1:]}
    suffix = "adoptada" if accepted else "solicitada"
    assert rows["Familia continua " + suffix] == f"{steel.bar_label} @ 0.200 m"
    assert rows["Familia adicional intercalada " + suffix] == f"{steel.bar_label} @ 0.200 m"
    if not accepted:
        assert not any("adoptada" in label for label in rows)
        assert "sin cortes" in " ".join(p.text for p in document.paragraphs)
        assert f"{steel.bar_label} @ {steel.spacing_m:.3f} m" in " ".join(p.text for p in document.paragraphs)


def test_selected_cut_summary_uses_actual_minor_and_major_spacing(zonal):
    text = format_connected_selection(zonal)
    for region in (TOP, BOTTOM):
        line = next(line for line in text.splitlines() if line.startswith(region + ":"))
        assert "s menor 0.100 m; s mayor 0.200 m" in line
        assert "s mayor 0.100" not in line
    assert not any(line.startswith("Pantalla - vertical relleno:") for line in text.splitlines())


def test_interleaved_families_must_respect_spacing_where_they_overlap(analysis):
    side = replace(analysis.inputs.left, reinforcement=replace(
        analysis.inputs.left.reinforcement, minimum_spacing_m=.025))
    result = replace(analysis, inputs=replace(analysis.inputs, left=side, right=side))
    # Each 0.10 m family fits alone; their 0.05 m combined spacing does not.
    with pytest.raises(ValueError, match="espacio libre del catalogo"):
        check_region_choice(result, TOP, foundation_zone_choice(ConnectedBarChoice('1"', .1)))


def test_reports_and_reselection_remove_old_zone_checks(zonal, tmp_path):
    export_connected_results(zonal, tmp_path)
    zones = list(csv.DictReader((tmp_path/'zonas_zapata.csv').open(encoding='utf-8-sig')))
    assert len(zones) == 6 and all(z['estado'] == 'OK' for z in zones)
    assert all(z['uso'] == 'ADOPTADO' for z in zones)
    payload = json.loads((tmp_path/'resultados.json').read_text(encoding='utf-8'))
    assert payload['selected_reinforcement'][TOP]['foundation_continuous']['spacing_m'] == .2
    for steel in zonal.reinforcement:
        if steel.region in (TOP, BOTTOM):
            text = str(foundation_cut_rows(steel))
            assert 'Familia continua solicitada' in text and 'Momento de servicio local' in text
    value = ConnectedBarChoice('1"', .2)
    updated = apply_connected_selections(zonal, {TOP: ConnectedSteelChoice(value, value)})
    cut = next(s for s in updated.reinforcement if s.region == TOP).foundation_reinforcement_cut
    assert not cut.zones and cut.requested_continuous_spacing_m is None
    export_connected_results(updated, tmp_path)
    zones = list(csv.DictReader((tmp_path/'zonas_zapata.csv').open(encoding='utf-8-sig')))
    assert not any(z['distribucion'] == TOP for z in zones)


@pytest.mark.parametrize("spacing", [.1, .125])
def test_dense_uniform_choice_does_not_adopt_automatic_cuts(analysis, spacing):
    from bridge_design.reporting.connected_docx_layout import reinforcement_summary_rows
    from bridge_design.domain.connected_distributions import reinforcement_demands
    from bridge_design.domain.connected_steel_audit import region_audit
    value = ConnectedBarChoice('1"', spacing, True)
    steel = check_region_choice(analysis, TOP, ConnectedSteelChoice(value, value))
    cut = steel.foundation_reinforcement_cut
    assert cut.status == 'APLICA' and cut.zones  # The useful proposal is retained.
    assert not cut.is_adopted
    assert tuple(reinforcement_summary_rows(steel)) == (
        (TOP, f'1" @ {spacing:.3f} m', f'{5/spacing:.3f}', steel.status),)
    result = replace(analysis, reinforcement=(steel,))
    assert not any(line.startswith(TOP + ':') for line in format_connected_selection(result).splitlines())
    trace = region_audit(result, steel, reinforcement_demands(result.inputs, result.mesh, result.results)[TOP])
    assert trace['service']['spacing'] == spacing and 'zone' not in trace['service']
    assert trace['service']['crack_ratio'] == pytest.approx(steel.crack_utilization)
    document = Document()
    _configure_document(document)
    write_foundation_cut(document, steel)
    cells = ' '.join(c.text for t in document.tables for row in t.rows for c in row.cells)
    assert 'PROPUESTA NO ADOPTADA' in cells and 'Cortes definitivos' not in cells
    assert 'sin cortes' in ' '.join(p.text for p in document.paragraphs)


def test_reselecting_same_dense_grid_clears_cut_adoption(zonal):
    from bridge_design.reporting.connected_docx_layout import reinforcement_summary_rows
    value = ConnectedBarChoice('1"', .1)
    steel = check_region_choice(zonal, TOP, ConnectedSteelChoice(value, value))
    assert steel.foundation_reinforcement_cut.status == 'APLICA'
    assert not steel.foundation_reinforcement_cut.is_adopted
    assert len(tuple(reinforcement_summary_rows(steel))) == 1
    assert choice_from_steel(steel).foundation_continuous is None


def test_stem_cut_choice_survives_enter_but_not_numeric_or_custom(analysis, monkeypatch):
    from bridge_design.reporting.connected_docx_layout import reinforcement_summary_rows
    region = 'Pantalla - vertical relleno'
    group = next(g for g in connected_reinforcement_options(analysis) if g.region == region)
    cut = next(c for c in group.cuts if c.bar_label == '1"' and c.heavy_spacing_m == .1)
    answers = iter(('s', cut.code))
    monkeypatch.setattr('builtins.input', lambda prompt: next(answers))
    choice = collect_connected_selection(analysis, (group,))[region]
    assert choice.stem_cut
    steel = check_region_choice(analysis, region, choice)
    assert steel.stem_cut_adopted and choice_from_steel(steel).stem_cut
    assert len(tuple(reinforcement_summary_rows(steel))) == 2
    current = replace(group, adopted=steel)
    answers = iter(('s', ''))
    assert collect_connected_selection(analysis, (current,))[region].stem_cut
    item = next(i for i, option in enumerate(group.principal, 1) if option.is_compliant)
    answers = iter(('s', str(item)))
    assert not collect_connected_selection(analysis, (current,))[region].stem_cut
    answers = iter(('s', 'p', '1"', '.1'))
    uniform = collect_connected_selection(analysis, (current,))[region]
    assert not uniform.stem_cut
    revised = check_region_choice(analysis, region, uniform)
    assert revised.stem_reinforcement_cut is not None and not revised.stem_cut_adopted
    assert len(tuple(reinforcement_summary_rows(revised))) == 1


def test_adopted_cut_service_and_word_use_the_local_governing_zone(analysis):
    from bridge_design.domain.connected_distributions import reinforcement_demands
    from bridge_design.domain.connected_steel_audit import region_audit
    from bridge_design.reporting.connected_distribution_report import write_distributed_design
    from bridge_design.reporting.connected_design_calculations import developed_steps
    # A controlled service envelope: smaller end moment, but half the steel there.
    # Retain every strength case, geometry and cutoff algorithm.
    def local_moment(row):
        element = analysis.mesh.frame.elements[row.element]
        a, b = (analysis.mesh.frame.nodes[n].x for n in (element.start, element.end))
        x = a + row.station*(b-a)
        return -40 if x < 2 or x > analysis.inputs.total_length_m-2 else -50
    result = replace(analysis, results=tuple(replace(case, sections=tuple(
        replace(row, moment=local_moment(row)) for row in case.sections))
        if case.limit_state == 'service' else case for case in analysis.results))
    steel = check_region_choice(result, TOP, foundation_zone_choice(ConnectedBarChoice('1"', .2)))
    assert steel.foundation_reinforcement_cut.is_adopted
    zones = steel.foundation_reinforcement_cut.zones
    worst = max(zones, key=lambda z: z.crack_utilization)
    assert worst.reinforcement == 'Solo continuo'
    demands = reinforcement_demands(result.inputs, result.mesh, result.results)[TOP]
    trace = region_audit(result, steel, demands)
    service = trace['service']
    assert service['crack_ratio'] > steel.crack_utilization
    assert service['zone'] == worst.name
    assert service['stress'] == pytest.approx(worst.service_stress_kg_cm2)
    assert service['area'] == 25 and service['maximum_gap'] == .2
    assert service['spacing_ratio'] == pytest.approx(.2/service['maximum_spacing'])
    before = dict(worst.service_check)
    region_audit(result, steel, demands)
    assert worst.service_check == before and 'x' not in worst.service_check
    step = next(s for s in developed_steps(steel, trace) if s.title.startswith('Fisuracion y'))
    assert f'{service["stress"]:.2f}' in step.substitution
    assert f'0.200/{service["maximum_spacing"]:.3f}' in step.substitution
    document = Document()
    _configure_document(document)
    write_distributed_design(document, result, (steel,), {'steel': {TOP: trace}})
    table = next(t for t in document.tables if 'Servicio gobernante' in [c.text for c in t.rows[0].cells])
    cells = [c.text for c in table.rows[1].cells]
    assert worst.name in cells[1]
    assert cells[2] == f'{worst.service_stress_kg_cm2:.2f}'
    assert cells[-1] == 'CUMPLE'


def test_automatic_nonuniform_proposal_keeps_real_gap_in_local_trace(analysis):
    value = ConnectedBarChoice('1"', .125)
    steel = check_region_choice(analysis, TOP, ConnectedSteelChoice(value, value))
    cut = steel.foundation_reinforcement_cut
    assert cut.status == 'APLICA' and cut.pattern.cycle_bars > 2
    for zone in cut.zones:
        check = zone.service_check
        assert check['area'] == pytest.approx(zone.area_cm2_m)
        assert check['maximum_gap'] == zone.maximum_gap_m
        assert check['crack_ratio'] == pytest.approx(zone.crack_utilization)
        assert check['spacing_ratio'] == pytest.approx(zone.maximum_gap_m/check['maximum_spacing'])
