"""Selection recalculates actual cuts and every output consumes those results."""

import csv
import json
from pathlib import Path

import pytest
from docx import Document

from bridge_design.cli.connected_2_yaml import connected_2_inputs_from_yaml, connected_2_yaml_template
from bridge_design.cli.connected_selection import format_connected_options
from bridge_design.domain.connected_design import solve_connected_abutments
from bridge_design.domain.connected_options import apply_connected_selections, connected_reinforcement_options
from bridge_design.domain.connected_reinforcement import ConnectedBarChoice, ConnectedSteelChoice
from bridge_design.reporting.connected_audit import build_connected_audit
from bridge_design.reporting.connected_distribution_report import write_distributed_design
from bridge_design.reporting.connected_export import export_connected_results
from bridge_design.reporting import connected_docx


def choice(bar, spacing):
    value = ConnectedBarChoice(bar, spacing)
    return ConnectedSteelChoice(value, value)


@pytest.fixture(scope="module")
def selected():
    result = solve_connected_abutments(connected_2_inputs_from_yaml(connected_2_yaml_template(example=True)))
    return apply_connected_selections(result, {
        "Zapata combinada - longitudinal superior": choice('3/4"', .125),
        "Zapata combinada - longitudinal inferior": choice('1"', .1),
    })


def test_real_model_has_both_pdf_arrangements_and_reports(selected, tmp_path):
    cuts = [s for s in selected.reinforcement if s.foundation_reinforcement_cut]
    assert len(cuts) == 2
    assert all(s.foundation_reinforcement_cut.status == "APLICA" for s in cuts)
    top, bottom = (s.foundation_reinforcement_cut for s in cuts)
    assert len(top.additional_intervals_m) == 1
    assert len(bottom.additional_intervals_m) == 2
    assert top.cutoff_left_m + top.cutoff_right_m == pytest.approx(selected.inputs.total_length_m)
    assert bottom.cutoff_left_m + bottom.cutoff_right_m == pytest.approx(selected.inputs.total_length_m)
    options = format_connected_options(connected_reinforcement_options(selected))
    assert options.count("CORTES DE ACERO - Zapata combinada") == 2
    assert "Cortes definitivos" in options and "Patron de corte" in options
    export_connected_results(selected, tmp_path)
    payload = json.loads((tmp_path/"resultados.json").read_text(encoding="utf-8"))
    serialized = [s["foundation_reinforcement_cut"] for s in payload["reinforcement"] if s["foundation_reinforcement_cut"]]
    assert serialized[0]["cutoff_left_m"] == top.cutoff_left_m
    assert serialized[1]["cutoff_left_m"] == bottom.cutoff_left_m
    rows = list(csv.DictReader((tmp_path/"cortes_zapata.csv").open(encoding="utf-8-sig")))
    assert len(rows) == 2 and all(r["estado"] == "APLICA" for r in rows)
    for name in ("resumen.txt", "auditoria.txt"):
        text = (tmp_path/name).read_text(encoding="utf-8")
        assert "CORTES DE ACERO - Zapata combinada - longitudinal superior" in text
        assert "CORTES DE ACERO - Zapata combinada - longitudinal inferior" in text
    document = Document(Path(connected_docx.__file__).parent / "templates" / "connected_reference.docx")
    write_distributed_design(document, selected, selected.reinforcement, build_connected_audit(selected))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Cortes del acero de zapata - superior" in text
    assert "Cortes del acero de zapata - inferior" in text
    cells = "\n".join(c.text for t in document.tables for row in t.rows for c in row.cells)
    assert f"{top.cutoff_left_m:.3f} / {top.cutoff_right_m:.3f} m" in cells
    assert f"{bottom.cutoff_left_m:.3f} / {bottom.cutoff_right_m:.3f} m" in cells


def test_reselection_discards_stale_cuts_and_exported_lengths(selected, tmp_path):
    region = "Zapata combinada - longitudinal superior"
    updated = apply_connected_selections(selected, {region: choice('1"', .2)})
    steel = next(s for s in updated.reinforcement if s.region == region)
    cut = steel.foundation_reinforcement_cut
    assert cut.status == "NO APLICA"
    assert "0.400" in cut.reason
    assert cut.cutoff_left_m is None and not cut.additional_intervals_m
    export_connected_results(updated, tmp_path)
    row = next(r for r in csv.DictReader((tmp_path/"cortes_zapata.csv").open(encoding="utf-8-sig")) if r["distribucion"] == region)
    assert row["estado"] == "NO APLICA"
    assert row["x_corte_izq_m"] == row["x_corte_der_m"] == ""
