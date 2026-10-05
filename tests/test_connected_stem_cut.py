from dataclasses import replace
from types import SimpleNamespace

import pytest

from bridge_design.domain.abutment import _stem_thickness_at_height_m
from bridge_design.domain.connected_inputs import ConnectedInputs, FoundationSoil
from bridge_design.domain.connected_reinforcement import (
    SectionDemand, connected_stem_reinforcement_cut,
)
from bridge_design.domain.connected_section_checks import section_check
from bridge_design.domain.rebar_catalog import reinforcing_bar_by_label


@pytest.fixture
def cut_model():
    data = ConnectedInputs(FoundationSoil(3000, .5, 26.7))
    nodes = tuple(SimpleNamespace(y=float(y)) for y in range(5))
    elements = tuple(SimpleNamespace(start=i, end=i+1, region="Pantalla " + side)
                     for side in ("izquierda", "derecha") for i in range(4))
    mesh = SimpleNamespace(frame=SimpleNamespace(nodes=nodes, elements=elements))
    steel = SimpleNamespace(region="Pantalla - vertical relleno", status="OK", bar_label='3/4"', spacing_m=.1,
        area_per_face_cm2_m=28.4, temperature_cm2_m=5.0, required_straight_anchor_cm=50.0)
    demands = tuple(SectionDemand("strength", "strength", 100, 0, 0, 0, e, .5)
                    for e in range(8))
    return data, mesh, steel, demands


def test_cut_reuses_pattern_and_development_and_checks_both_sides(cut_model):
    data, mesh, steel, demands = cut_model
    bar = reinforcing_bar_by_label(steel.bar_label)
    depth = 100 * _stem_thickness_at_height_m(data.right, 2-1e-8)
    capacity = section_check(replace(demands[5], depth_cm=depth), data.right,
                             data.right.reinforcement.stem_cover_cm, bar, .3)["capacity"]
    # Right abutment controls, despite zero demands on the left.
    demands = tuple(replace(d, moment=1.05*capacity) if d.element == 5 else d for d in demands)
    cut = connected_stem_reinforcement_cut(data, mesh, steel, demands)
    assert cut is not None
    assert cut.continuous_every_n_bars == 3
    assert cut.upper_spacing_m == pytest.approx(.3)
    assert cut.theoretical_cut_height_m == pytest.approx(2.0, abs=1e-7)
    assert cut.constructive_cut_height_m == pytest.approx(2.5, abs=1e-7)
    assert cut.lower_cut_bar_length_m == pytest.approx(3.0, abs=1e-7)
    assert cut.continuous_bar_length_m == pytest.approx(4.5)
    assert cut.status == "OK"
    assert "cajuela" in cut.notes


@pytest.mark.parametrize("field,limit_state,value", [
    ("moment", "strength", 500), ("shear", "strength", 500), ("moment", "service", 500),
])
def test_upper_peak_prevents_cut_even_when_base_is_unloaded(cut_model, field, limit_state, value):
    data, mesh, steel, demands = cut_model
    demands = tuple(replace(d, limit_state=limit_state, **{field: value})
                    if d.element == 7 else d for d in demands)
    assert connected_stem_reinforcement_cut(data, mesh, steel, demands) is None


def test_sparse_or_noncompliant_selected_steel_has_no_cut(cut_model):
    data, mesh, steel, demands = cut_model
    steel.spacing_m = .2
    assert connected_stem_reinforcement_cut(data, mesh, steel, demands) is None
    steel.spacing_m = .1
    steel.status = "NO CUMPLE"
    assert connected_stem_reinforcement_cut(data, mesh, steel, demands) is None


def test_development_reaching_top_is_not_worth_cutting(cut_model):
    data, mesh, steel, demands = cut_model
    steel.required_straight_anchor_cm = 400
    cut = connected_stem_reinforcement_cut(data, mesh, steel, demands)
    assert cut is not None
    assert cut.status == "NO CONVIENE"
    assert cut.lower_cut_bar_length_m == cut.continuous_bar_length_m


def test_exterior_vertical_steel_has_no_cut(cut_model):
    data, mesh, steel, demands = cut_model
    steel.region = "Pantalla - vertical exterior"
    assert connected_stem_reinforcement_cut(data, mesh, steel, demands) is None


def test_selection_recalculates_proposal_and_exports_it(tmp_path):
    import json
    from bridge_design.cli.connected_output import format_connected_result
    from bridge_design.cli.connected_selection import format_connected_options
    from bridge_design.domain.connected_design import solve_connected_abutments
    from bridge_design.domain.connected_options import apply_connected_selections, connected_reinforcement_options
    from bridge_design.domain.connected_reinforcement import ConnectedBarChoice, ConnectedSteelChoice
    from bridge_design.reporting.connected_export import export_connected_results

    analysis = solve_connected_abutments(ConnectedInputs(FoundationSoil(3000, .5, 26.7),
                                        mesh_size_m=1, include_without_bridge=False))
    # Isolate the selection/cut pipeline from the foundation's loaded checks.
    analysis = replace(analysis, results=tuple(replace(case, sections=tuple(
        replace(row, moment=0, shear=0) for row in case.sections)) for case in analysis.results))
    region = "Pantalla - vertical relleno"
    choice = ConnectedBarChoice('3/4"', .1)
    selected = apply_connected_selections(analysis, {region: ConnectedSteelChoice(choice, choice)})
    steel = next(s for s in selected.reinforcement if s.region == region)
    assert steel.stem_reinforcement_cut is not None
    assert steel.stem_reinforcement_cut.lower_spacing_m == .1
    options = format_connected_options(connected_reinforcement_options(selected))
    assert "OPCION DE CORTE DE ACERO PRINCIPAL DE PANTALLA" in options
    assert "Continua 1 de cada 3 barras inferiores" in options
    assert "Propuesta de corte para Pantalla - vertical exterior" not in options
    assert options.count("OPCION DE CORTE DE ACERO PRINCIPAL DE PANTALLA") == 1
    exterior = next(s for s in selected.reinforcement if s.region == "Pantalla - vertical exterior")
    assert exterior.stem_reinforcement_cut is None
    assert "Altura constructiva de corte" in format_connected_result(selected)
    assert "Propuesta de corte del acero elegido - Pantalla - vertical exterior" not in format_connected_result(selected)
    export_connected_results(selected, tmp_path)
    payload = json.loads((tmp_path / "resultados.json").read_text(encoding="utf-8"))
    row = next(s for s in payload["reinforcement"] if s["region"] == region)
    assert row["stem_reinforcement_cut"]["upper_spacing_m"] == pytest.approx(.3)
    from docx import Document
    from bridge_design.reporting.connected_audit import build_connected_audit
    from bridge_design.reporting.connected_distribution_report import write_distributed_design
    from pathlib import Path
    from bridge_design.reporting import connected_docx
    document = Document(Path(connected_docx.__file__).parent / "templates" / "connected_reference.docx")
    write_distributed_design(document, selected, selected.reinforcement, build_connected_audit(selected))
    text = "\n".join(p.text for p in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "Opción de corte de acero principal de pantalla" in text
    assert text.count("Opción de corte de acero principal de pantalla") == 1
    assert "Continúa 1 de cada 3 barras inferiores" in table_text
    assert "Longitud barras cortadas" in table_text
    choice = ConnectedBarChoice('3/4"', .2)
    changed = apply_connected_selections(selected, {region: ConnectedSteelChoice(choice, choice)})
    assert next(s for s in changed.reinforcement if s.region == region).stem_reinforcement_cut is None
