"""Independent geometry and output checks for beta > 0, vertical-back walls."""

from dataclasses import replace
from io import BytesIO
from math import radians, tan
import re

import pytest
from docx import Document

from bridge_design.domain import abutment as a
from bridge_design.domain.inclined_wall import stem_actions
from bridge_design.cli import abutment_ascii_output as terminal
from bridge_design.cli.yaml_inputs import cantilever_wall_inputs_from_yaml, cantilever_wall_yaml_template
from bridge_design.reporting import abutment_docx as report
from bridge_design.reporting import abutment_docx_detail as trace
from bridge_design.reporting.abutment_docx_charts import save_abutment_geometry


def inputs(beta=10., heel=3., delta=0.):
    base = a.pure_wall_inputs()
    return replace(
        base,
        geometry=replace(base.geometry, retained_height_m=8.3, footing_width_m=3.+heel,
                         footing_thickness_m=1., toe_length_m=2.,
                         lower_stem_thickness_m=1., upper_stem_thickness_m=.5, front_soil_depth_m=1.),
        soil=replace(base.soil, backfill_slope_deg=beta, wall_backface_angle_deg=90.,
                     wall_soil_friction_deg=delta, vehicular_surcharge_height_m=.6,
                     pedestrian_surcharge_tn_m2=.25),
        key=replace(base.key, enabled=False),
    )


def polygon_properties(points):
    pairs = list(zip(points, points[1:] + points[:1]))
    crosses = [x*v-u*y for (x,y),(u,v) in pairs]
    signed_area = sum(crosses)/2
    cx = sum((p[0]+q[0])*c for (p,q),c in zip(pairs,crosses))/(6*signed_area)
    cy = sum((p[1]+q[1])*c for (p,q),c in zip(pairs,crosses))/(6*signed_area)
    return abs(signed_area), cx, cy


@pytest.mark.parametrize("beta", [0., 5., 10., 19.])
@pytest.mark.parametrize("heel", [1., 3., 5.])
def test_soil_weight_and_centroids_match_full_sloping_polygon(beta, heel):
    data = inputs(beta, heel)
    rise = heel*tan(radians(beta))
    area, x, y = polygon_properties([(3.,1.), (3.+heel,1.), (3.+heel,8.3+rise), (3.,8.3)])
    actual = a._resultant(a._soil_components(data))
    assert actual == pytest.approx((area*1.925,x,y), rel=1e-12)
    assert a.virtual_retained_height_m(data) == pytest.approx(8.3+rise)
    slopes = [c for c in a._soil_components(data) if c.name == "Relleno por pendiente beta"]
    assert len(slopes) == (1 if beta else 0)


@pytest.mark.parametrize("beta", [0., 5., 10., 19.])
def test_global_pressures_arms_and_inertia_use_virtual_height(beta):
    result = a.solve_abutment_design(inputs(beta))
    height = 8.3+3*tan(radians(beta))
    p = result.pressures
    assert p.eh_tn_m == pytest.approx(.5*p.ka*1.925*height**2)
    assert p.pae_tn_m == pytest.approx(.5*p.k_ae*1.925*height**2)
    assert p.lsx_tn_m == pytest.approx(p.ka*(1.925*.6+.25)*height)
    assert p.pedestrian_lsx_tn_m == pytest.approx(p.ka*.25*height)
    assert p.lsy_tn_m == pytest.approx(3*.6*1.925)
    assert p.pedestrian_lsy_tn_m == pytest.approx(3*.25)
    assert p.pir_tn_m == pytest.approx(.18*(result.dc_self_weight_tn_m+result.ev_weight_tn_m))
    components = result.components.horizontal_without_bridge
    assert [c.arm_m for c in components[:3]] == pytest.approx([height/2,height/3,height/2])
    mass_first = sum(c.value_tn_m*c.vertical_arm_m for c in result.concrete_components+result.soil_components)
    assert components[3].moment_tn_m_m == pytest.approx(.5*.18*mass_first)
    combo_b = result.components.horizontal_without_bridge_seismic_b
    earth = max(.5*p.pae_tn_m,p.eh_tn_m)
    assert combo_b[1].value_tn_m == pytest.approx(earth)
    assert combo_b[1].arm_m == pytest.approx(height/3 if earth<=p.eh_tn_m+1e-12 else height/2)
    assert combo_b[2].moment_tn_m_m == pytest.approx(.18*mass_first)


@pytest.mark.parametrize("beta", [0., 10., 19.])
def test_heel_vertical_load_and_moment_integrate_sloping_overburden(beta):
    result = a.solve_abutment_design(inputs(beta))
    data = result.inputs
    # Independent integration of gamma_s*(Hp+x*tan(beta)) over the heel.
    length, hp, slope, gamma = 3.,7.3,tan(radians(beta)),1.925
    soil_force = gamma*(hp*length+slope*length**2/2)
    soil_moment = gamma*(hp*length**2/2+slope*length**3/3)
    weights = a._heel_weight_breakdown(data)
    assert weights.soil_weight_tn_m == pytest.approx(soil_force)
    assert weights.soil_weight_tn_m*weights.soil_arm_m == pytest.approx(soil_moment)
    assert weights.soil_slope_tn_m == pytest.approx(gamma*slope*length**2/2)
    states = result.without_bridge[:2]+result.extreme_seismic_without_bridge+result.service_without_bridge
    for state in states:
        f = state.load_factors
        actual = a._heel_state_demand(data,state)
        force = f.dc*2.4*length + f.ev*soil_force + f.ls_vertical*(.6*gamma+.25)*length
        moment = f.dc*2.4*length**2/2 + f.ev*soil_moment + f.ls_vertical*(.6*gamma+.25)*length**2/2
        assert actual.downward_shear_tn_m == pytest.approx(force)
        assert actual.downward_moment_tn_m_m == pytest.approx(moment)


@pytest.mark.parametrize("delta", [0., 10.])
def test_stem_real_height_does_not_grow_when_heel_length_changes(delta):
    short = a.solve_abutment_design(inputs(10.,1.,delta))
    long = a.solve_abutment_design(inputs(10.,5.,delta))
    assert long.pressures.eh_tn_m > short.pressures.eh_tn_m
    assert stem_actions(long.inputs,long.pressures) == pytest.approx(stem_actions(short.inputs,short.pressures))
    assert long.stem_design.controlling_moment_tn_m_m == pytest.approx(short.stem_design.controlling_moment_tn_m_m)


def test_yaml_preserves_positive_beta_with_vertical_backface():
    data = cantilever_wall_yaml_template()
    data['suelo_sismo'].update(beta_pendiente_relleno_grados=10.,theta_cara_posterior_desde_horizontal_grados=90.)
    parsed = cantilever_wall_inputs_from_yaml(data)
    assert parsed.soil.backfill_slope_deg == 10.
    assert parsed.soil.wall_backface_angle_deg == 90.
    assert a.virtual_retained_height_m(parsed) > parsed.geometry.retained_height_m


def test_combined_slopes_accepted_but_invalid_mononobe_okabe_rejected():
    assert replace(inputs().soil,wall_backface_angle_deg=88.).backfill_slope_deg == 10.
    with pytest.raises(ValueError,match='angulo sismico'):
        inputs(20.)


def test_beta_zero_limit_is_continuous_and_abutment_height_unchanged():
    flat = a.solve_abutment_design(inputs(0.))
    near = a.solve_abutment_design(inputs(1e-7))
    for l,r in zip(flat.without_bridge,near.without_bridge):
        assert l.vu_tn_m == pytest.approx(r.vu_tn_m,abs=1e-5)
        assert l.hu_tn_m == pytest.approx(r.hu_tn_m,abs=1e-5)
        assert l.eccentricity_m == pytest.approx(r.eccentricity_m,abs=1e-6)
    data = a.AbutmentInputs(soil=a.AbutmentSoilInputs(backfill_slope_deg=10.))
    assert a.virtual_retained_height_m(data) == data.geometry.retained_height_m
    assert a._sloping_backfill_component(data) is None


def test_word_and_both_terminal_tables_reproduce_global_pressure_numbers(monkeypatch):
    result = a.solve_abutment_design(inputs())
    blocks = {}
    monkeypatch.setattr(report,'_calc',lambda doc,title,*fields: blocks.update({title:fields}))
    document = Document()
    report._configure_document(document)
    report._earth_pressures(document,result)
    for title,prefix in [('Empuje estático del relleno','EH'),('Incremento sísmico de empuje','PAE')]:
        match = re.search(rf'{prefix} = 0.5·([\d.]+)·([\d.]+)·([\d.]+)² = ([\d.]+)',blocks[title][2])
        assert match
        coefficient,gamma,height,printed = map(float,match.groups())
        assert height == pytest.approx(8.3+3*tan(radians(10)),abs=.00000051)
        assert .5*coefficient*gamma*height**2 == pytest.approx(printed,abs=.001)
    assert 'Wβ' in ' '.join(blocks['Geometría del relleno y altura del plano virtual'])
    assert 'Hv' in trace.pressure_arms_trace(result)[0]
    for formatter in (terminal._format_soil_pressure_procedure,terminal._format_pressures):
        output = '\n'.join(formatter(result))
        assert 'Hv plano virtual' in output
        assert f'{a.virtual_retained_height_m(result.inputs):.6f}' in output
    heel_trace = trace.heel_toe_demand_trace(result,result.heel_design)
    assert 'WEV,β = 1.527' in heel_trace[2]
    assert '43.685 Tn/m' in heel_trace[2]


def test_serialized_word_contains_same_geometry_and_weight():
    result = a.solve_abutment_design(inputs())
    document = Document()
    report._configure_document(document)
    report._weights(document,result,report._labels(is_pure_wall=True))
    report._earth_pressures(document,result)
    report._structural_case(document,result,result.heel_design)
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    text = ' '.join(Document(stream).element.xpath('//*[local-name()="t"]/text()'))
    assert 'Relleno por pendiente beta' in text
    assert '8.828981' in text
    # Word encodes fractions as numerator/denominator nodes, not a slash.
    assert 'Wβ=1.925·3.000000·0.528981' in ''.join(text.split())
    assert '1.527432' in text
    assert 'WEV,β=1.527' in ''.join(text.split())


def test_geometry_picture_uses_sloping_surface(tmp_path):
    from PIL import Image
    result = a.solve_abutment_design(inputs())
    path = save_abutment_geometry(result,tmp_path/'wall_beta10.png')
    with Image.open(path) as picture:
        # With beta > 0, the surface at the heel is higher than at the stem.
        # At this elevation (between the two), the far heel is soil, near stem is air.
        assert picture.getpixel((1550,160)) == (243,232,210)
        assert picture.getpixel((950,160)) == (255,255,255)
